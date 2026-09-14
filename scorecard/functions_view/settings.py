from collections import defaultdict
from datetime import time as datetime_time
import re
from urllib.parse import urlencode

from django.contrib import messages
from django.contrib.auth.decorators import login_required, permission_required
from django.contrib.auth.models import Group, Permission
from django.core.cache import cache
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils.dateparse import parse_time

import openpyxl
from openpyxl.utils import get_column_letter

from Users.models import CustomUser
from scorecard.functions_view.audit import log_scorecard_audit, scorecard_audit_filter
from scorecard.models import (
    BankBranch,
    ScorecardUserBranchAccess,
    ScorecardWorkflowApprovalSetting,
)
from scorecard.default_role_seeder import ensure_default_role_seeder_synced
from scorecard.permission_catalog import (
    DEFAULT_ROLE_DEFINITIONS,
    MODULE_ACCESS_REGISTRY,
    SCORECARD_PERMISSION_DEFINITIONS,
)
from scorecard.workflow_approval import get_scorecard_workflow_approval_settings


ROLE_LEVEL_ORDER = {
    "view": 1,
    "work": 2,
    "approve": 3,
    "admin": 4,
}

REVIEW_ONLY_NAV_MODULE_KEYS = {"basel_templates", "ifrs9_templates", "basel_scores", "ifrs9_scores"}
REVIEW_ONLY_PERMISSION_CODES = {
    "basel_templates": {"review_basel_templates"},
    "ifrs9_templates": {"review_ifrs9_templates"},
    "basel_scores": {"review_basel_scores"},
    "ifrs9_scores": {"review_ifrs9_scores"},
}
OPENABLE_LEVEL_KEYS = {"view", "work", "admin"}

ROLE_LEVEL_BADGE_STYLES = {
    "view": "view",
    "work": "work",
    "approve": "approve",
    "admin": "admin",
}

PERMISSION_BADGE_LABELS = {
    "view": "View",
    "work": "Can Work",
    "approve": "Can Approve",
    "assign": "Can Assign",
    "admin": "Admin",
}

WORKFLOW_APPROVAL_SECTIONS = [
    {
        "key": "basel_scores",
        "label": "Basel Scores",
        "copy": "Control whether Basel score submissions skip checker review when the submitter already has elevated authority.",
        "fields": [
            {
                "name": "basel_score_superuser_auto_approve",
                "label": "Superusers auto-approve Basel scores",
                "copy": "When enabled, a superuser’s Basel score submission is approved immediately on submit.",
            },
            {
                "name": "basel_score_admin_auto_approve",
                "label": "Administrator-role users auto-approve Basel scores",
                "copy": "When enabled, users with Basel score administrator permission approve their own Basel score submission immediately.",
            },
            {
                "name": "basel_score_reviewer_auto_approve",
                "label": "Checker-role users auto-approve Basel scores",
                "copy": "When enabled, users with Basel score review permission approve their own Basel score submission immediately.",
            },
        ],
    },
    {
        "key": "ifrs9_scores",
        "label": "IFRS9 Scores",
        "copy": "Control whether IFRS9 score submissions can complete approval immediately when the submitter already has review authority.",
        "fields": [
            {
                "name": "ifrs9_score_superuser_auto_approve",
                "label": "Superusers auto-approve IFRS9 scores",
                "copy": "When enabled, a superuser’s IFRS9 score submission is approved immediately on submit.",
            },
            {
                "name": "ifrs9_score_admin_auto_approve",
                "label": "Administrator-role users auto-approve IFRS9 scores",
                "copy": "When enabled, users with IFRS9 score administrator permission approve their own IFRS9 submission immediately.",
            },
            {
                "name": "ifrs9_score_reviewer_auto_approve",
                "label": "Checker-role users auto-approve IFRS9 scores",
                "copy": "When enabled, users with IFRS9 score review permission approve their own IFRS9 submission immediately.",
            },
        ],
    },
    {
        "key": "basel_templates",
        "label": "Basel Templates",
        "copy": "Control whether Basel template submissions go straight to approved when the submitter can already review template workflow items.",
        "fields": [
            {
                "name": "basel_template_superuser_auto_approve",
                "label": "Superusers auto-approve Basel templates",
                "copy": "When enabled, a superuser’s Basel template submission is approved immediately on submit.",
            },
            {
                "name": "basel_template_admin_auto_approve",
                "label": "Administrator-role users auto-approve Basel templates",
                "copy": "When enabled, users with Basel template administrator permission approve their own template submission immediately.",
            },
            {
                "name": "basel_template_reviewer_auto_approve",
                "label": "Checker-role users auto-approve Basel templates",
                "copy": "When enabled, users with Basel template review permission approve their own template submission immediately.",
            },
        ],
    },
    {
        "key": "ifrs9_templates",
        "label": "IFRS9 Templates",
        "copy": "Control whether IFRS9 template submissions can complete approval immediately when the submitter already has review rights.",
        "fields": [
            {
                "name": "ifrs9_template_superuser_auto_approve",
                "label": "Superusers auto-approve IFRS9 templates",
                "copy": "When enabled, a superuser’s IFRS9 template submission is approved immediately on submit.",
            },
            {
                "name": "ifrs9_template_admin_auto_approve",
                "label": "Administrator-role users auto-approve IFRS9 templates",
                "copy": "When enabled, users with IFRS9 template administrator permission approve their own template submission immediately.",
            },
            {
                "name": "ifrs9_template_reviewer_auto_approve",
                "label": "Checker-role users auto-approve IFRS9 templates",
                "copy": "When enabled, users with IFRS9 template review permission approve their own template submission immediately.",
            },
        ],
    },
    {
        "key": "counterpart_completion",
        "label": "Score Pairing",
        "copy": "Control whether officers must complete the matching Basel/IFRS9 score for their own submitted customer before starting another customer.",
        "fields": [
            {
                "name": "enforce_score_counterpart_completion",
                "label": "Require officers to complete missing Basel/IFRS9 counterpart scores",
                "copy": "When enabled, an officer who submits Basel or IFRS9 for a customer must complete the missing counterpart score for that same customer before starting another customer.",
            },
        ],
    },
    {
        "key": "cross_branch_duplicate_scoring",
        "label": "Customer Re-Scoring",
        "title": "Cross-branch scored customer lock",
        "copy": "Control whether a customer already scored in one branch can be entered again for that same scorecard type in another branch.",
        "fields": [
            {
                "name": "prevent_cross_branch_duplicate_scoring",
                "label": "Prevent scoring a customer again once already scored in any branch",
                "copy": "When enabled, a customer who already has a completed Basel II score cannot start a new Basel II score in any branch, and the same applies separately to IFRS9. The user can only preview the saved scored information in read-only mode.",
            },
        ],
    },
    {
        "key": "without_score_sources",
        "label": "Without Score Lists",
        "title": "Source tables for customers without scores",
        "copy": "Control which staged customer sources feed the Without Basel II Scores and Without IFRS9 Scores lists. If both sources are disabled, those lists show no staged customers.",
        "fields": [
            {
                "name": "without_score_list_include_loans",
                "label": "Use CustomerLoan in without-score lists",
                "copy": "When enabled, customers from the CustomerLoan staging table are included in the missing Basel and IFRS9 score lists.",
            },
            {
                "name": "without_score_list_include_overdrafts",
                "label": "Use CustomerOverdraft in without-score lists",
                "copy": "When enabled, customers from the CustomerOverdraft staging table are included in the missing Basel and IFRS9 score lists.",
            },
        ],
    },
    {
        "key": "submission_notification_routing",
        "label": "Submission Emails",
        "title": "Branch-specific approval notification routing",
        "copy": "Control who receives Basel/IFRS9 score submission alerts when a maker submits a score for checker approval.",
        "fields": [
            {
                "name": "single_branch_submission_reviewer_notifications",
                "label": "Send score submission alerts only to one-branch reviewers",
                "copy": "When enabled, Basel/IFRS9 score submission emails and in-app notifications go only to reviewers, administrators, or superusers whose branch scope contains exactly one branch matching the submitted score branch. When disabled, all eligible reviewers for the branch are notified.",
            },
        ],
    },
    {
        "key": "without_score_email_routing",
        "label": "Without-Score Emails",
        "title": "Branch-specific missing-score email routing",
        "copy": "Control who receives scheduled customer lists for customers missing Basel II or IFRS9 scores.",
        "fields": [
            {
                "name": "single_branch_without_score_maker_notifications",
                "label": "Send missing-score lists only to one-branch makers",
                "copy": "When enabled, scheduled missing-score emails go only to Basel/IFRS9 score makers, administrators, or superusers whose branch scope contains exactly one matching branch. When disabled, all eligible score makers for the branch are notified.",
            },
        ],
    },
]

AUTO_REFRESH_FREQUENCY_CHOICES = [
    ("daily", "Daily"),
    ("weekly", "Weekly"),
    ("monthly", "Monthly"),
]

AUTO_REFRESH_WEEKDAY_CHOICES = [
    (0, "Monday"),
    (1, "Tuesday"),
    (2, "Wednesday"),
    (3, "Thursday"),
    (4, "Friday"),
    (5, "Saturday"),
    (6, "Sunday"),
]

AUTO_REFRESH_MONTH_DAY_CHOICES = list(range(1, 32))
AUTO_REFRESH_BATCH_DEFAULT = 1000
AUTO_REFRESH_BATCH_MIN = 1
AUTO_REFRESH_BATCH_MAX = 10000

AUDIT_MODEL_LABELS = {
    "ScorecardPermissionAssignment": "Permission Role Assignment",
    "ScorecardPermissionBranchAccess": "Permission Branch Access",
    "ScorecardWorkflowApprovalSetting": "Workflow Rules",
    "ScorecardBaselScore": "Basel Scores",
    "ScorecardIFRS9Score": "IFRS9 Scores",
    "ScorecardIFRS9SupportingData": "IFRS9 Supporting Data",
    "ScorecardBaselTemplate": "Basel Templates",
    "ScorecardIFRS9Template": "IFRS9 Templates",
    "ScorecardEvaluationDocument": "Evaluation Documents",
    "ScorecardDocumentLibrary": "Scorecard Documents",
    "ScorecardUpload": "External Imports",
    "ScorecardEmail": "Email Administration",
    "ScorecardApi": "API Operations",
    "ScorecardCheckerApprovals": "Checker My Approvals",
    "ScorecardIFRS9Results": "IFRS9 Results",
    "ScorecardHistoricalScore": "Historical Scores",
    "ScorecardManualOverdraftCustomer": "Manual Overdraft Customers",
}


def _post_checkbox_is_enabled(post_data, field_name):
    value = post_data.get(field_name)
    if value is None:
        return False
    return str(value).strip().lower() in {"on", "true", "1", "yes"}


def _post_choice_value(post_data, field_name, allowed_values, default):
    value = str(post_data.get(field_name, default) or default).strip().lower()
    return value if value in set(allowed_values) else default


def _post_int_range(post_data, field_name, default, minimum, maximum):
    try:
        value = int(post_data.get(field_name, default))
    except (TypeError, ValueError):
        return default
    return max(minimum, min(maximum, value))


def _post_time_value(post_data, field_name, default):
    parsed = parse_time(str(post_data.get(field_name, "") or "").strip())
    return parsed or default


def _permission_definition_map():
    return {item["codename"]: item for item in SCORECARD_PERMISSION_DEFINITIONS}


def _audit_model_label(model_name):
    return AUDIT_MODEL_LABELS.get(model_name, model_name)


def _audit_action_label(action_name):
    if not action_name:
        return ""
    action_map = {
        "create": "Create",
        "created": "Create",
        "update": "Update",
        "updated": "Update",
        "delete": "Delete",
        "deleted": "Delete",
        "submit": "Submit",
        "submitted": "Submit",
        "approve": "Approve",
        "approved": "Approve",
        "return": "Return",
        "returned": "Return",
        "assign": "Assign",
        "import": "Import",
        "bulk_upload": "Bulk Upload",
        "bulk_upload_failed": "Bulk Upload Failed",
        "retry": "Retry",
        "download": "Download",
        "run": "Run",
        "stop": "Stop",
        "view_my_approvals": "View My Approvals",
        "create_section": "Add Section",
        "update_section": "Edit Section",
        "delete_section": "Delete Section",
        "reorder_section": "Reorder Sections",
        "create_risk_driver": "Add Risk Driver",
        "update_risk_driver": "Edit Risk Driver",
        "delete_risk_driver": "Delete Risk Driver",
        "reorder_risk_driver": "Reorder Risk Drivers",
        "create_attribute": "Add Attribute",
        "update_attribute": "Edit Attribute",
        "delete_attribute": "Delete Attribute",
        "reorder_attribute": "Reorder Attributes",
        "create_option": "Add Option",
        "update_option": "Edit Option",
        "delete_option": "Delete Option",
        "reorder_option": "Reorder Options",
        "create_grade_band": "Add Grade Band",
        "update_grade_band": "Edit Grade Band",
        "delete_grade_band": "Delete Grade Band",
        "reorder_grade_band": "Reorder Grade Bands",
        "submit_for_review": "Submit for Review",
        "withdraw_submission": "Withdraw Submission",
    }
    return action_map.get(action_name.lower(), action_name.replace("_", " ").title())


def _extract_audit_branch_label(change_description: str) -> str:
    description = (change_description or "").strip()
    if not description:
        return "-"

    multi_branch_patterns = [
        r"New branch access:\s*(.+?)(?:\.|$)",
        r"New branches:\s*(.+?)(?:\.|$)",
        r"Branch access:\s*(.+?)(?:\.|$)",
    ]
    for pattern in multi_branch_patterns:
        match = re.search(pattern, description, flags=re.IGNORECASE)
        if match:
            value = (match.group(1) or "").strip()
            if not value:
                return "-"
            branch_parts = [item.strip() for item in value.split(",") if item.strip()]
            if len(branch_parts) <= 1:
                return branch_parts[0] if branch_parts else "-"
            return "Multiple branches"

    single_branch_patterns = [
        r"Branch:\s*([^;.\n]+)",
        r"branch\s+'([^']+)'",
    ]
    for pattern in single_branch_patterns:
        match = re.search(pattern, description, flags=re.IGNORECASE)
        if match:
            value = (match.group(1) or "").strip()
            return value or "-"

    return "-"


def _clean_page_size(value, default=20):
    allowed = {10, 20, 50, 100}
    try:
        parsed = int(value or default)
    except (TypeError, ValueError):
        parsed = default
    return parsed if parsed in allowed else default


def _user_display_label(user_obj):
    if not user_obj:
        return ""
    if getattr(user_obj, "email", ""):
        return user_obj.email
    full_name = f"{getattr(user_obj, 'name', '')} {getattr(user_obj, 'surname', '')}".strip()
    if full_name:
        return full_name
    return str(user_obj).strip()


def _build_settings_query_string(request, excluded_keys=None):
    excluded = set(excluded_keys or [])
    query_params = []
    for key in request.GET:
        if key in excluded:
            continue
        for value in request.GET.getlist(key):
            if value:
                query_params.append((key, value))
    return urlencode(query_params)


AUDIT_TRAIL_CACHE_TTL_SECONDS = 120


def _audit_summary_cache_key(suffix: str) -> str:
    return f"scorecard:audit_trail:{suffix}"


def _get_audit_static_filter_options():
    from scorecard.models import ScorecardUserAuditTrail

    branch_cache_key = _audit_summary_cache_key("branch_options")
    model_cache_key = _audit_summary_cache_key("model_options")
    action_cache_key = _audit_summary_cache_key("action_options")
    cached_branches = cache.get(branch_cache_key)
    cached_models = cache.get(model_cache_key)
    cached_actions = cache.get(action_cache_key)
    if cached_branches is not None and cached_models is not None and cached_actions is not None:
        return cached_branches, cached_models, cached_actions

    audit_entries = ScorecardUserAuditTrail.objects.filter(scorecard_audit_filter())
    direct_branch_options = set(
        audit_entries.exclude(branch_name="")
        .exclude(branch_name__isnull=True)
        .order_by()
        .values_list("branch_name", flat=True)
        .distinct()
    )
    legacy_branch_options: set[str] = set()
    blank_branch_entries = audit_entries.filter(Q(branch_name="") | Q(branch_name__isnull=True))
    for _, change_description in blank_branch_entries.values_list("id", "change_description").iterator(chunk_size=1000):
        branch_label = _extract_audit_branch_label(change_description)
        normalized_label = (branch_label or "").strip()
        if normalized_label and normalized_label != "-":
            legacy_branch_options.add(normalized_label)

    branch_options = sorted(direct_branch_options | legacy_branch_options)
    unique_model_values = list(
        audit_entries.order_by().values_list("model_name", flat=True).distinct()
    )
    unique_action_values = list(
        audit_entries.order_by().values_list("action", flat=True).distinct()
    )
    model_options = [
        {
            "value": item,
            "label": _audit_model_label(item),
        }
        for item in unique_model_values
        if item
    ]
    action_options = [
        {
            "value": item,
            "label": _audit_action_label(item),
        }
        for item in unique_action_values
        if item
    ]
    action_options.sort(key=lambda item: item["label"])

    cache.set(branch_cache_key, branch_options, AUDIT_TRAIL_CACHE_TTL_SECONDS)
    cache.set(model_cache_key, model_options, AUDIT_TRAIL_CACHE_TTL_SECONDS)
    cache.set(action_cache_key, action_options, AUDIT_TRAIL_CACHE_TTL_SECONDS)
    return branch_options, model_options, action_options


def _role_metadata_map():
    metadata = {}
    for definition in DEFAULT_ROLE_DEFINITIONS:
        names = [definition["name"], *definition.get("legacy_names", [])]
        for name in names:
            metadata[name] = definition
    return metadata


def _module_lookup():
    return {module["key"]: module for module in MODULE_ACCESS_REGISTRY}


def _module_label_by_permission():
    labels = {}
    for module in MODULE_ACCESS_REGISTRY:
        for code in module.get("permission_codes", []):
            labels[code] = module["label"]
    return labels


def _permission_badge_for_codename(codename):
    if codename == "assign_scorecard_permissions":
        return "assign"
    if codename.startswith("review_"):
        return "approve"
    if codename.startswith("view_") or codename == "access_scorecard_dashboard":
        return "view"
    if codename.startswith("manage_") or codename.startswith("run_") or codename.startswith("use_") or codename.startswith("reopen_"):
        return "work"
    return "view"


def _scorecard_role_names():
    names = set()
    for item in DEFAULT_ROLE_DEFINITIONS:
        names.add(item["name"])
        names.update(item.get("legacy_names", []))
    return names


def _scorecard_roles_queryset():
    ensure_default_role_seeder_synced()
    return (
        Group.objects.filter(
            Q(name__in=_scorecard_role_names()) | Q(permissions__content_type__app_label="scorecard")
        )
        .distinct()
        .annotate(
            member_count=Count("customuser_groups", distinct=True),
            permission_count=Count("permissions", distinct=True),
        )
        .order_by("name")
    )


def _scorecard_role_ids():
    return set(_scorecard_roles_queryset().values_list("id", flat=True))


def _module_is_openable(module_key, matching_roles, module_permission_hits):
    if module_key not in REVIEW_ONLY_NAV_MODULE_KEYS:
        return bool(matching_roles or module_permission_hits)
    role_levels = {role.get("level_key") for role in matching_roles}
    permission_levels = {_permission_badge_for_codename(code) for code in module_permission_hits}
    return bool(role_levels & OPENABLE_LEVEL_KEYS or permission_levels & OPENABLE_LEVEL_KEYS)


def _has_checker_module_access(roles, effective_permission_codes):
    review_codes = {
        "review_basel_templates",
        "review_ifrs9_templates",
        "review_basel_scores",
        "review_ifrs9_scores",
    }
    if any(role.get("level_key") in {"approve", "admin"} and role.get("module_key") in REVIEW_ONLY_NAV_MODULE_KEYS for role in roles):
        return True
    return any(code in effective_permission_codes for code in review_codes)


def _visible_module_labels_for_roles(roles, effective_permission_codes):
    labels = []
    for module in MODULE_ACCESS_REGISTRY:
        matching_roles = [role for role in roles if role.get("module_key") == module["key"]]
        module_permission_hits = [code for code in module.get("permission_codes", []) if code in effective_permission_codes]
        if not _module_is_openable(module["key"], matching_roles, module_permission_hits):
            continue
        labels.append(module["label"])
    if _has_checker_module_access(roles, effective_permission_codes):
        labels.append("Checker")
    deduped = []
    for label in labels:
        if label not in deduped:
            deduped.append(label)
    return deduped


def _serialize_role(role):
    permission_definitions = _permission_definition_map()
    role_metadata = _role_metadata_map().get(role.name, {})
    permission_entries = []
    permissions_qs = role.permissions.filter(content_type__app_label="scorecard").order_by("codename")
    for permission in permissions_qs:
        definition = permission_definitions.get(permission.codename)
        permission_entries.append(
            {
                "codename": permission.codename,
                "label": definition["label"] if definition else permission.codename.replace("_", " ").title(),
                "description": definition["description"] if definition else "",
                "badge_key": _permission_badge_for_codename(permission.codename),
                "badge_label": PERMISSION_BADGE_LABELS.get(_permission_badge_for_codename(permission.codename), "View"),
            }
        )

    return {
        "id": role.id,
        "name": role.name,
        "label": role_metadata.get("label", role.name),
        "description": role_metadata.get("description", "Scorecard access role."),
        "module_key": role_metadata.get("module_key", "other"),
        "module_label": role_metadata.get("module_label", "Other Roles"),
        "level_key": role_metadata.get("level_key", "view"),
        "level_label": role_metadata.get("level_label", "View"),
        "level_order": ROLE_LEVEL_ORDER.get(role_metadata.get("level_key", "view"), 99),
        "badge_style": ROLE_LEVEL_BADGE_STYLES.get(role_metadata.get("level_key", "view"), "view"),
        "access_badges": role_metadata.get("access_badges", []),
        "action_summary": role_metadata.get("action_summary", []),
        "member_count": getattr(role, "member_count", 0),
        "permission_count": getattr(role, "permission_count", 0),
        "permissions": permission_entries,
    }


def _build_module_access_summary():
    permissions_by_code = _permission_definition_map()
    role_definitions = defaultdict(list)
    for definition in DEFAULT_ROLE_DEFINITIONS:
        role_definitions[definition.get("module_key", "other")].append(definition)

    summary = []
    for module in MODULE_ACCESS_REGISTRY:
        permission_records = []
        for code in module.get("permission_codes", []):
            definition = permissions_by_code.get(code)
            if not definition:
                continue
            badge_key = _permission_badge_for_codename(code)
            permission_records.append(
                {
                    "codename": definition["codename"],
                    "label": definition["label"],
                    "description": definition["description"],
                    "badge_key": badge_key,
                    "badge_label": PERMISSION_BADGE_LABELS.get(badge_key, "View"),
                }
            )

        role_records = []
        for definition in sorted(role_definitions.get(module["key"], []), key=lambda item: (ROLE_LEVEL_ORDER.get(item.get("level_key", "view"), 99), item["name"])):
            role_records.append(
                {
                    "label": definition.get("label", definition["name"]),
                    "description": definition.get("description", ""),
                    "level_label": definition.get("level_label", "View"),
                    "badge_style": ROLE_LEVEL_BADGE_STYLES.get(definition.get("level_key", "view"), "view"),
                    "action_summary": definition.get("action_summary", []),
                }
            )

        summary.append(
            {
                "key": module["key"],
                "label": module["label"],
                "views_file": module["views_file"],
                "template_folder": module["template_folder"],
                "notes": module.get("notes", ""),
                "permissions": permission_records,
                "roles": role_records,
                "permission_count": len(permission_records),
                "route_count": len(module.get("route_names", [])),
            }
        )
    return summary


def _group_roles_for_assignment(roles):
    grouped = {}
    module_order = [module["label"] for module in MODULE_ACCESS_REGISTRY]

    for role in sorted((_serialize_role(role) for role in roles), key=lambda item: (item["module_label"], item["level_order"], item["label"])):
        grouped.setdefault(role["module_label"], []).append(role)

    ordered_groups = []
    seen = set()
    for module_label in module_order:
        if module_label in grouped:
            ordered_groups.append({"module_label": module_label, "roles": grouped[module_label]})
            seen.add(module_label)

    for module_label, items in grouped.items():
        if module_label not in seen:
            ordered_groups.append({"module_label": module_label, "roles": items})
    return ordered_groups


def _user_effective_scorecard_permissions(user_obj):
    direct_permissions = {
        permission.split(".", 1)[1]
        for permission in user_obj.get_all_permissions()
        if permission.startswith("scorecard.")
    }
    return direct_permissions


def _serialize_branch(branch):
    return {
        "id": branch.id,
        "bank_name": branch.bank_name,
        "branch_name": branch.branch_name,
        "branch_code": branch.branch_code,
        "label": f"{branch.bank_name} - {branch.branch_name}",
    }


def _get_explicit_branch_assignments(user_obj):
    return list(
        BankBranch.objects.filter(scorecard_user_access_entries__user=user_obj)
        .order_by("bank_name", "branch_name")
        .distinct()
    )


def _build_user_access_summary(user_obj):
    permission_definitions = _permission_definition_map()
    module_labels = _module_label_by_permission()
    scorecard_role_ids = _scorecard_role_ids()
    selected_roles = [
        _serialize_role(group)
        for group in user_obj.groups.filter(id__in=scorecard_role_ids).prefetch_related("permissions")
    ]
    selected_roles.sort(key=lambda item: (item["module_label"], item["level_order"], item["label"]))

    effective_permission_codes = _user_effective_scorecard_permissions(user_obj)
    grouped_permissions = defaultdict(list)
    for codename in sorted(effective_permission_codes):
        definition = permission_definitions.get(codename)
        if not definition:
            continue
        badge_key = _permission_badge_for_codename(codename)
        grouped_permissions[module_labels.get(codename, "Other")].append(
            {
                "codename": codename,
                "label": definition["label"],
                "description": definition["description"],
                "badge_key": badge_key,
                "badge_label": PERMISSION_BADGE_LABELS.get(badge_key, "View"),
            }
        )

    grouped_permissions_list = [
        {"module_label": module_label, "permissions": permissions}
        for module_label, permissions in grouped_permissions.items()
    ]
    grouped_permissions_list.sort(key=lambda item: item["module_label"])

    assigned_branches = [_serialize_branch(branch) for branch in _get_explicit_branch_assignments(user_obj)]
    if assigned_branches:
        branch_scope_label = f"{len(assigned_branches)} assigned branch{'es' if len(assigned_branches) != 1 else ''}"
        branch_scope_style = "work"
        branch_scope_copy = "This user is restricted to the branches selected in the scorecard permission workspace."
    else:
        branch_scope_label = "No branches assigned"
        branch_scope_style = "none"
        branch_scope_copy = "No scorecard branches are assigned to this user, so they will not see branch-based scorecard information."

    roles_by_module = defaultdict(list)
    for role in selected_roles:
        roles_by_module[role["module_key"]].append(role)

    module_access = []
    for module in MODULE_ACCESS_REGISTRY:
        matching_roles = roles_by_module.get(module["key"], [])
        module_permission_hits = [code for code in module.get("permission_codes", []) if code in effective_permission_codes]
        if not matching_roles and not module_permission_hits:
            continue
        if not _module_is_openable(module["key"], matching_roles, module_permission_hits):
            continue

        highest_role = max((ROLE_LEVEL_ORDER.get(role["level_key"], 0) for role in matching_roles), default=0)
        highest_permission_level = max((ROLE_LEVEL_ORDER.get(_permission_badge_for_codename(code), 0) for code in module_permission_hits), default=0)
        highest_level = max(highest_role, highest_permission_level)

        access_label = "View Only"
        access_style = "view"
        if highest_level >= ROLE_LEVEL_ORDER["admin"]:
            access_label = "Admin"
            access_style = "admin"
        elif highest_level >= ROLE_LEVEL_ORDER["approve"]:
            access_label = "Can Approve"
            access_style = "approve"
        elif highest_level >= ROLE_LEVEL_ORDER["work"]:
            access_label = "Can Work"
            access_style = "work"

        action_summary = []
        for role in matching_roles:
            action_summary.extend(role.get("action_summary", []))
        if not action_summary:
            for code in module_permission_hits:
                definition = permission_definitions.get(code)
                if definition:
                    action_summary.append(definition["description"])

        seen_actions = []
        for item in action_summary:
            if item and item not in seen_actions:
                seen_actions.append(item)

        module_access.append(
            {
                "label": module["label"],
                "notes": module.get("notes", ""),
                "route_count": len(module.get("route_names", [])),
                "access_label": access_label,
                "access_style": access_style,
                "role_labels": [role["label"] for role in matching_roles],
                "actions": seen_actions[:4],
            }
        )

    if _has_checker_module_access(selected_roles, effective_permission_codes):
        checker_roles = [
            role
            for role in selected_roles
            if role.get("module_key") in REVIEW_ONLY_NAV_MODULE_KEYS and role.get("level_key") in {"approve", "admin"}
        ]
        checker_actions = []
        for role in checker_roles:
            checker_actions.extend(role.get("action_summary", []))
        seen_checker_actions = []
        for action in checker_actions:
            if action not in seen_checker_actions:
                seen_checker_actions.append(action)
        module_access.append(
            {
                "label": "Checker",
                "notes": "Review, approve, and return submitted score and template workflow items.",
                "route_count": 4,
                "access_label": "Can Approve",
                "access_style": "approve",
                "role_labels": [role["label"] for role in checker_roles],
                "actions": seen_checker_actions[:4],
            }
        )

    return {
        "roles": selected_roles,
        "effective_permissions": sorted(effective_permission_codes),
        "grouped_permissions": grouped_permissions_list,
        "module_access": module_access,
        "assigned_branches": assigned_branches,
        "branch_scope_label": branch_scope_label,
        "branch_scope_style": branch_scope_style,
        "branch_scope_copy": branch_scope_copy,
        "branch_total_count": len(assigned_branches),
        "has_explicit_branch_access": bool(assigned_branches),
    }


def _build_permission_matrix():
    matrix_rows = []
    role_buckets = defaultdict(lambda: defaultdict(list))
    module_lookup = _module_lookup()

    for definition in DEFAULT_ROLE_DEFINITIONS:
        module_key = definition.get("module_key", "other")
        role_buckets[module_key][definition.get("level_key", "view")].append(definition.get("label", definition["name"]))

    modules_for_matrix = [*MODULE_ACCESS_REGISTRY]
    modules_for_matrix.append(
        {
            "key": "platform",
            "label": "All Scorecard Areas",
            "notes": "Full-platform emergency role that bundles every scorecard permission.",
            "route_names": [],
        }
    )

    for module in modules_for_matrix:
        view_roles = role_buckets[module["key"]].get("view", [])
        work_roles = role_buckets[module["key"]].get("work", [])
        approve_roles = role_buckets[module["key"]].get("approve", [])
        admin_roles = role_buckets[module["key"]].get("admin", [])
        if not any([view_roles, work_roles, approve_roles, admin_roles]):
            continue
        matrix_rows.append(
            {
                "label": module["label"],
                "notes": module.get("notes", ""),
                "route_count": len(module.get("route_names", [])),
                "view_roles": view_roles,
                "work_roles": work_roles,
                "approve_roles": approve_roles,
                "admin_roles": admin_roles,
            }
        )
    return matrix_rows


def _serialize_workflow_approval_settings(settings_obj):
    section_rows = []
    enabled_count = 0
    total_count = 0

    for section in WORKFLOW_APPROVAL_SECTIONS:
        rows = []
        for field in section["fields"]:
            is_enabled = bool(getattr(settings_obj, field["name"], False))
            rows.append(
                {
                    "name": field["name"],
                    "label": field["label"],
                    "copy": field["copy"],
                    "is_enabled": is_enabled,
                }
            )
            enabled_count += 1 if is_enabled else 0
            total_count += 1
        section_rows.append(
            {
                "key": section["key"],
                "label": section["label"],
                "title": section.get("title") or f"{section['label']} auto-approval",
                "copy": section["copy"],
                "rows": rows,
                "enabled_count": sum(1 for row in rows if row["is_enabled"]),
                "total_count": len(rows),
            }
        )

    auto_refresh_enabled = bool(getattr(settings_obj, "auto_refresh_autofilled_scores_enabled", False))
    auto_refresh_frequency = (
        getattr(settings_obj, "auto_refresh_autofilled_scores_frequency", "daily") or "daily"
    ).lower()
    valid_frequencies = {value for value, _ in AUTO_REFRESH_FREQUENCY_CHOICES}
    if auto_refresh_frequency not in valid_frequencies:
        auto_refresh_frequency = "daily"
    try:
        auto_refresh_weekday = int(getattr(settings_obj, "auto_refresh_autofilled_scores_weekday", 0) or 0)
    except (TypeError, ValueError):
        auto_refresh_weekday = 0
    auto_refresh_weekday = max(0, min(6, auto_refresh_weekday))
    try:
        auto_refresh_month_day = int(getattr(settings_obj, "auto_refresh_autofilled_scores_month_day", 1) or 1)
    except (TypeError, ValueError):
        auto_refresh_month_day = 1
    auto_refresh_month_day = max(1, min(31, auto_refresh_month_day))
    auto_refresh_time = getattr(settings_obj, "auto_refresh_autofilled_scores_time", None) or datetime_time(2, 0)
    auto_refresh_batch_size = _post_int_range(
        {
            "auto_refresh_autofilled_scores_batch_size": getattr(
                settings_obj,
                "auto_refresh_autofilled_scores_batch_size",
                AUTO_REFRESH_BATCH_DEFAULT,
            )
        },
        "auto_refresh_autofilled_scores_batch_size",
        AUTO_REFRESH_BATCH_DEFAULT,
        AUTO_REFRESH_BATCH_MIN,
        AUTO_REFRESH_BATCH_MAX,
    )
    auto_frequency_labels = dict(AUTO_REFRESH_FREQUENCY_CHOICES)
    auto_weekday_labels = dict(AUTO_REFRESH_WEEKDAY_CHOICES)
    auto_refresh = {
        "is_enabled": auto_refresh_enabled,
        "frequency": auto_refresh_frequency,
        "frequency_label": auto_frequency_labels.get(auto_refresh_frequency, "Daily"),
        "time_value": auto_refresh_time.strftime("%H:%M"),
        "weekday": auto_refresh_weekday,
        "weekday_label": auto_weekday_labels.get(auto_refresh_weekday, "Monday"),
        "month_day": auto_refresh_month_day,
        "batch_size": auto_refresh_batch_size,
        "basel_cursor_id": getattr(settings_obj, "auto_refresh_autofilled_scores_basel_cursor_id", 0) or 0,
        "ifrs9_cursor_id": getattr(settings_obj, "auto_refresh_autofilled_scores_ifrs9_cursor_id", 0) or 0,
        "pending_update_count": len(
            getattr(settings_obj, "auto_refresh_autofilled_scores_pending_updates", []) or []
        ),
        "last_run_at": getattr(settings_obj, "auto_refresh_autofilled_scores_last_run_at", None),
        "frequency_choices": [
            {
                "value": value,
                "label": label,
                "selected": value == auto_refresh_frequency,
            }
            for value, label in AUTO_REFRESH_FREQUENCY_CHOICES
        ],
        "weekday_choices": [
            {
                "value": value,
                "label": label,
                "selected": value == auto_refresh_weekday,
            }
            for value, label in AUTO_REFRESH_WEEKDAY_CHOICES
        ],
        "month_day_choices": [
            {
                "value": value,
                "label": value,
                "selected": value == auto_refresh_month_day,
            }
            for value in AUTO_REFRESH_MONTH_DAY_CHOICES
        ],
    }
    enabled_count += 1 if auto_refresh_enabled else 0
    total_count += 1

    return {
        "sections": section_rows,
        "auto_refresh": auto_refresh,
        "enabled_count": enabled_count,
        "total_count": total_count,
    }


@login_required
@permission_required("scorecard.view_scorecard_permissions", raise_exception=True)
def settings_dashboard(request):
    module_summary = _build_module_access_summary()
    
    roles = list(_scorecard_roles_queryset().prefetch_related("permissions"))
    serialized_roles = [_serialize_role(role) for role in roles]
    role_library = defaultdict(list)
    for role in sorted(serialized_roles, key=lambda item: (item["module_label"], item["level_order"], item["label"])):
        role_library[role["module_label"]].append(role)

    context = {
        "total_users": CustomUser.objects.count(),
        "total_scorecard_roles": len(roles),
        "total_permissions": len(SCORECARD_PERMISSION_DEFINITIONS),
        "secured_routes": sum(item["route_count"] for item in module_summary),
        "module_summary": module_summary,
        "role_library": [{"module_label": label, "roles": items} for label, items in role_library.items()],
        "matrix_preview_count": len(_build_permission_matrix()),
    }
    return render(request, "settings/scorecard_permission_dashboard.html", context)


@login_required
@permission_required("scorecard.view_scorecard_permissions", raise_exception=True)
def settings_user_roles(request):
    search_query = request.GET.get("q", "").strip()
    try:
        page_size = int(request.GET.get("page_size", 20))
    except (TypeError, ValueError):
        page_size = 20
    if page_size not in {10, 20, 50, 100}:
        page_size = 20

    users_qs = (
        CustomUser.objects.prefetch_related(
            "groups",
            "groups__permissions",
            "scorecard_branch_access_entries__branch",
        )
        .order_by("email", "name", "surname")
    )
    if search_query:
        users_qs = users_qs.filter(
            Q(name__icontains=search_query)
            | Q(surname__icontains=search_query)
            | Q(email__icontains=search_query)
            | Q(department__icontains=search_query)
        )

    paginator = Paginator(users_qs, page_size)
    page_obj = paginator.get_page(request.GET.get("page"))
    users = list(page_obj.object_list)
    scorecard_role_ids = _scorecard_role_ids()
    user_rows = []
    for user in users:
        scorecard_groups = [group for group in user.groups.all() if group.id in scorecard_role_ids]
        serialized_roles = [_serialize_role(group) for group in scorecard_groups]
        explicit_branches = [entry.branch for entry in user.scorecard_branch_access_entries.all()]
        effective_permission_codes = _user_effective_scorecard_permissions(user)
        module_labels = _visible_module_labels_for_roles(serialized_roles, effective_permission_codes)
        highest_level = max((role["level_order"] for role in serialized_roles), default=0)
        access_label = "No Role"
        access_style = "none"
        if highest_level >= ROLE_LEVEL_ORDER["admin"]:
            access_label = "Admin"
            access_style = "admin"
        elif highest_level >= ROLE_LEVEL_ORDER["approve"]:
            access_label = "Can Approve"
            access_style = "approve"
        elif highest_level >= ROLE_LEVEL_ORDER["work"]:
            access_label = "Can Work"
            access_style = "work"
        elif highest_level >= ROLE_LEVEL_ORDER["view"]:
            access_label = "View Only"
            access_style = "view"

        user_rows.append(
            {
                "user": user,
                "scorecard_roles": serialized_roles,
                "module_labels": module_labels[:4],
                "module_count": len(module_labels),
                "extra_module_count": max(len(module_labels) - 4, 0),
                "access_label": access_label,
                "access_style": access_style,
                "branch_count": len(explicit_branches),
                "branch_scope_label": f"{len(explicit_branches)} assigned" if explicit_branches else "No branches",
            }
        )

    context = {
        "user_rows": user_rows,
        "page_obj": page_obj,
        "page_size": page_size,
        "search_query": search_query,
        "filtered_total_users": paginator.count,
        "total_users": CustomUser.objects.count(),
        "visible_users": len(user_rows),
        "total_roles": len(scorecard_role_ids),
    }
    return render(request, "settings/scorecard_permission_user_roles.html", context)


@login_required
@permission_required("scorecard.view_scorecard_permissions", raise_exception=True)
def settings_user_access_summary(request, user_id):
    user_obj = get_object_or_404(CustomUser.objects.prefetch_related("groups", "groups__permissions"), pk=user_id)
    access_summary = _build_user_access_summary(user_obj)
    context = {
        "user_obj": user_obj,
        "assigned_roles": access_summary["roles"],
        "module_access": access_summary["module_access"],
        "grouped_permissions": access_summary["grouped_permissions"],
        "total_effective_permissions": len(access_summary["effective_permissions"]),
        "assigned_branches": access_summary["assigned_branches"],
        "branch_scope_label": access_summary["branch_scope_label"],
        "branch_scope_style": access_summary["branch_scope_style"],
        "branch_scope_copy": access_summary["branch_scope_copy"],
        "branch_total_count": access_summary["branch_total_count"],
        "has_explicit_branch_access": access_summary["has_explicit_branch_access"],
    }
    return render(request, "settings/scorecard_permission_user_access_summary.html", context)


@login_required
@permission_required("scorecard.view_scorecard_permissions", raise_exception=True)
def settings_permission_matrix(request):
    matrix_rows = _build_permission_matrix()
    context = {
        "matrix_rows": matrix_rows,
        "total_modules": len(matrix_rows),
        "total_roles": len(DEFAULT_ROLE_DEFINITIONS),
    }
    return render(request, "settings/scorecard_permission_matrix.html", context)


@login_required
@permission_required("scorecard.view_scorecard_workflow", raise_exception=True)
def settings_workflow_approvals(request):
    workflow_settings = get_scorecard_workflow_approval_settings()
    can_manage = request.user.is_superuser or request.user.has_perm("scorecard.manage_scorecard_workflow")
    has_saved_configuration = workflow_settings.updated_by_id is not None
    is_edit_mode = can_manage and (request.GET.get("edit") == "1" or not has_saved_configuration)

    if request.method == "POST":
        if not can_manage:
            messages.error(request, "You do not have permission to change workflow approval rules.")
            return redirect("scorecard:settings_workflow_approvals")

        previous_summary = _serialize_workflow_approval_settings(workflow_settings)
        previous_enabled = previous_summary["enabled_count"]

        for section in WORKFLOW_APPROVAL_SECTIONS:
            for field in section["fields"]:
                setattr(
                    workflow_settings,
                    field["name"],
                    _post_checkbox_is_enabled(request.POST, field["name"]),
                )

        workflow_settings.auto_refresh_autofilled_scores_enabled = _post_checkbox_is_enabled(
            request.POST, "auto_refresh_autofilled_scores_enabled"
        )
        workflow_settings.auto_refresh_autofilled_scores_frequency = _post_choice_value(
            request.POST,
            "auto_refresh_autofilled_scores_frequency",
            [value for value, _ in AUTO_REFRESH_FREQUENCY_CHOICES],
            "daily",
        )
        workflow_settings.auto_refresh_autofilled_scores_time = _post_time_value(
            request.POST,
            "auto_refresh_autofilled_scores_time",
            datetime_time(2, 0),
        )
        workflow_settings.auto_refresh_autofilled_scores_weekday = _post_int_range(
            request.POST,
            "auto_refresh_autofilled_scores_weekday",
            0,
            0,
            6,
        )
        workflow_settings.auto_refresh_autofilled_scores_month_day = _post_int_range(
            request.POST,
            "auto_refresh_autofilled_scores_month_day",
            1,
            1,
            31,
        )
        workflow_settings.auto_refresh_autofilled_scores_batch_size = _post_int_range(
            request.POST,
            "auto_refresh_autofilled_scores_batch_size",
            AUTO_REFRESH_BATCH_DEFAULT,
            AUTO_REFRESH_BATCH_MIN,
            AUTO_REFRESH_BATCH_MAX,
        )

        workflow_settings.updated_by = request.user
        workflow_settings.save()

        current_summary = _serialize_workflow_approval_settings(workflow_settings)
        changed_fields = []
        changed_rule_names = []
        for section in current_summary["sections"]:
            for row in section["rows"]:
                previous_row = next(
                    previous
                    for previous_section in previous_summary["sections"]
                    if previous_section["key"] == section["key"]
                    for previous in previous_section["rows"]
                    if previous["name"] == row["name"]
                )
                if previous_row["is_enabled"] != row["is_enabled"]:
                    changed_fields.append(
                        f"{row['label']}: {'Enabled' if row['is_enabled'] else 'Disabled'}"
                    )
                    changed_rule_names.append(row["name"])

        previous_auto_refresh = previous_summary.get("auto_refresh", {})
        current_auto_refresh = current_summary.get("auto_refresh", {})
        auto_refresh_changes = [
            (
                "Auto-update scored forms",
                "Enabled" if current_auto_refresh.get("is_enabled") else "Disabled",
                previous_auto_refresh.get("is_enabled"),
                current_auto_refresh.get("is_enabled"),
            ),
            (
                "Auto-update frequency",
                current_auto_refresh.get("frequency_label"),
                previous_auto_refresh.get("frequency"),
                current_auto_refresh.get("frequency"),
            ),
            (
                "Auto-update run time",
                current_auto_refresh.get("time_value"),
                previous_auto_refresh.get("time_value"),
                current_auto_refresh.get("time_value"),
            ),
            (
                "Auto-update weekly day",
                current_auto_refresh.get("weekday_label"),
                previous_auto_refresh.get("weekday"),
                current_auto_refresh.get("weekday"),
            ),
            (
                "Auto-update monthly day",
                current_auto_refresh.get("month_day"),
                previous_auto_refresh.get("month_day"),
                current_auto_refresh.get("month_day"),
            ),
            (
                "Auto-update batch size",
                current_auto_refresh.get("batch_size"),
                previous_auto_refresh.get("batch_size"),
                current_auto_refresh.get("batch_size"),
            ),
        ]
        for label, display_value, previous_value, current_value in auto_refresh_changes:
            if previous_value != current_value:
                changed_fields.append(f"{label}: {display_value}")
                changed_rule_names.append("auto_refresh_autofilled_scores")

        if {
            "without_score_list_include_loans",
            "without_score_list_include_overdrafts",
        } & set(changed_rule_names):
            from scorecard.functions_view.customers import bump_customer_list_summary_version

            bump_customer_list_summary_version()

        log_scorecard_audit(
            request.user,
            "ScorecardWorkflowApprovalSetting",
            "update",
            object_id=workflow_settings.id,
            change_description=(
                f"Updated scorecard workflow rules. "
                f"Enabled rules changed from {previous_enabled} to {current_summary['enabled_count']}. "
                f"Changes: {', '.join(changed_fields) if changed_fields else 'No field state changed.'}"
            ),
        )
        messages.success(request, "Scorecard workflow rules saved.")
        return redirect("scorecard:settings_workflow_approvals")

    context = {
        "workflow_settings": workflow_settings,
        "workflow_summary": _serialize_workflow_approval_settings(workflow_settings),
        "can_manage": can_manage,
        "has_saved_configuration": has_saved_configuration,
        "is_edit_mode": is_edit_mode,
        "workflow_updated_by_label": _user_display_label(workflow_settings.updated_by),
    }
    return render(request, "settings/scorecard_permission_workflow_approvals.html", context)


@login_required
@permission_required("scorecard.view_scorecard_permissions", raise_exception=True)
def settings_user_role_assignment(request, user_id):
    user_obj = get_object_or_404(CustomUser.objects.prefetch_related("groups", "groups__permissions"), pk=user_id)
    roles = list(_scorecard_roles_queryset().prefetch_related("permissions"))
    scorecard_role_ids = {role.id for role in roles}
    available_branches = list(BankBranch.objects.all().order_by("bank_name", "branch_name"))
    available_branch_ids = {branch.id for branch in available_branches}

    if request.method == "POST":
        if not request.user.is_superuser and not request.user.has_perm("scorecard.assign_scorecard_permissions"):
            messages.error(request, "You do not have permission to change scorecard access assignments.")
            return redirect("scorecard:settings_user_roles")

        previous_scorecard_group_names = list(
            user_obj.groups.filter(id__in=scorecard_role_ids).order_by("name").values_list("name", flat=True)
        )
        previous_branch_labels = [
            branch["label"] for branch in [_serialize_branch(branch) for branch in _get_explicit_branch_assignments(user_obj)]
        ]
        selected_role_ids = {
            int(role_id)
            for role_id in request.POST.getlist("role_ids")
            if role_id.isdigit() and int(role_id) in scorecard_role_ids
        }
        selected_branch_ids = {
            int(branch_id)
            for branch_id in request.POST.getlist("branch_ids")
            if branch_id.isdigit() and int(branch_id) in available_branch_ids
        }
        preserved_groups = list(user_obj.groups.exclude(id__in=scorecard_role_ids))
        selected_scorecard_groups = list(Group.objects.filter(id__in=selected_role_ids).order_by("name"))
        selected_branch_objects = [branch for branch in available_branches if branch.id in selected_branch_ids]

        with transaction.atomic():
            user_obj.groups.set(preserved_groups + selected_scorecard_groups)
            ScorecardUserBranchAccess.objects.filter(user=user_obj).exclude(branch_id__in=selected_branch_ids).delete()
            existing_branch_ids = set(
                ScorecardUserBranchAccess.objects.filter(user=user_obj).values_list("branch_id", flat=True)
            )
            ScorecardUserBranchAccess.objects.bulk_create(
                [
                    ScorecardUserBranchAccess(user=user_obj, branch=branch)
                    for branch in selected_branch_objects
                    if branch.id not in existing_branch_ids
                ]
            )

        updated_scorecard_group_names = [group.name for group in selected_scorecard_groups]
        updated_branch_labels = [_serialize_branch(branch)["label"] for branch in selected_branch_objects]
        log_scorecard_audit(
            request.user,
            "ScorecardPermissionAssignment",
            "update",
            object_id=user_obj.id,
            change_description=(
                f"Updated scorecard access assignments for {user_obj.email or user_obj.name}. "
                f"Previous scorecard roles: {', '.join(previous_scorecard_group_names) if previous_scorecard_group_names else 'None'}. "
                f"New scorecard roles: {', '.join(updated_scorecard_group_names) if updated_scorecard_group_names else 'None'}. "
                f"Previous branch access: {', '.join(previous_branch_labels) if previous_branch_labels else 'No branches assigned'}. "
                f"New branch access: {', '.join(updated_branch_labels) if updated_branch_labels else 'No branches assigned'}."
            ),
        )
        if previous_branch_labels != updated_branch_labels:
            log_scorecard_audit(
                request.user,
                "ScorecardPermissionBranchAccess",
                "update",
                object_id=user_obj.id,
                change_description=(
                    f"Updated branch access for {user_obj.email or user_obj.name}. "
                    f"Previous branches: {', '.join(previous_branch_labels) if previous_branch_labels else 'No branches assigned'}. "
                    f"New branches: {', '.join(updated_branch_labels) if updated_branch_labels else 'No branches assigned'}."
                ),
            )
        messages.success(request, f"Scorecard roles and branch access saved for {user_obj.email or user_obj.name}.")
        return redirect("scorecard:settings_user_roles")

    grouped_roles = _group_roles_for_assignment(roles)
    selected_user_role_ids = set(user_obj.groups.filter(id__in=scorecard_role_ids).values_list("id", flat=True))
    access_summary = _build_user_access_summary(user_obj)
    selected_branch_ids = {branch["id"] for branch in access_summary["assigned_branches"]}
    context = {
        "user_obj": user_obj,
        "grouped_roles": grouped_roles,
        "selected_user_role_ids": selected_user_role_ids,
        "available_branches": [_serialize_branch(branch) for branch in available_branches],
        "selected_branch_ids": selected_branch_ids,
        "total_roles": len(roles),
        "selected_roles": access_summary["roles"],
        "module_access": access_summary["module_access"],
        "total_effective_permissions": len(access_summary["effective_permissions"]),
        "assigned_branches": access_summary["assigned_branches"],
        "branch_scope_label": access_summary["branch_scope_label"],
        "branch_scope_style": access_summary["branch_scope_style"],
        "branch_scope_copy": access_summary["branch_scope_copy"],
    }
    return render(request, "settings/scorecard_permission_user_role_assignment.html", context)


@login_required
@permission_required("scorecard.view_scorecard_audit_trail", raise_exception=True)
def settings_audit_trail(request):
    from scorecard.models import ScorecardUserAuditTrail

    audit_entries = (
        ScorecardUserAuditTrail.objects.select_related("user")
        .only(
            "id",
            "user__email",
            "user__name",
            "user__surname",
            "model_name",
            "action",
            "object_id",
            "branch_name",
            "timestamp",
            "change_description",
        )
        .filter(scorecard_audit_filter())
    )

    search_query = request.GET.get("q", "").strip()
    model_name = request.GET.get("model_name", "").strip()
    action = request.GET.get("action", "").strip()
    branch_name = request.GET.get("branch_name", "").strip()
    date_from = request.GET.get("date_from", "").strip()
    date_to = request.GET.get("date_to", "").strip()
    page_size = _clean_page_size(request.GET.get("page_size"), default=20)
    list_query_string = _build_settings_query_string(request, excluded_keys={"page", "download"})

    if search_query:
        audit_entries = audit_entries.filter(
            Q(user__email__icontains=search_query)
            | Q(user__name__icontains=search_query)
            | Q(user__surname__icontains=search_query)
            | Q(model_name__icontains=search_query)
            | Q(object_id__icontains=search_query)
            | Q(change_description__icontains=search_query)
        )
    if model_name:
        audit_entries = audit_entries.filter(model_name=model_name)
    if action:
        audit_entries = audit_entries.filter(action=action)
    if date_from:
        audit_entries = audit_entries.filter(timestamp__date__gte=date_from)
    if date_to:
        audit_entries = audit_entries.filter(timestamp__date__lte=date_to)

    branch_options, model_options, action_options = _get_audit_static_filter_options()
    entry_branch_labels: dict[int, str] = {}
    matching_branch_ids: list[int] = []
    blank_branch_entries = audit_entries.filter(Q(branch_name="") | Q(branch_name__isnull=True))
    for entry_id, change_description in blank_branch_entries.values_list("id", "change_description").iterator(chunk_size=1000):
        branch_label = _extract_audit_branch_label(change_description)
        entry_branch_labels[entry_id] = branch_label
        normalized_label = (branch_label or "").strip()
        if branch_name and normalized_label == branch_name:
            matching_branch_ids.append(entry_id)

    ordered_entries = audit_entries.order_by("-timestamp")
    if branch_name:
        ordered_entries = ordered_entries.filter(Q(branch_name=branch_name) | Q(id__in=matching_branch_ids or [-1]))

    if request.GET.get("download") == "excel":
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Scorecard Audit"
        headers = ["User", "Area", "Stored Model", "Action", "Object ID", "Branch", "Timestamp", "Change Description"]
        sheet.append(headers)

        for entry in ordered_entries:
            entry.audit_branch_label = (entry.branch_name or "").strip() or entry_branch_labels.get(entry.id, "-")
            user_label = entry.user.email if entry.user and entry.user.email else (
                f"{entry.user.name} {entry.user.surname}".strip() if entry.user else "System"
            )
            sheet.append(
                [
                    user_label or "System",
                    _audit_model_label(entry.model_name),
                    entry.model_name,
                    entry.action,
                    entry.object_id or "",
                    entry.audit_branch_label,
                    entry.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
                    entry.change_description,
                ]
            )

        for column_index, _ in enumerate(headers, start=1):
            sheet.column_dimensions[get_column_letter(column_index)].width = 28

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="scorecard_audit.xlsx"'
        workbook.save(response)
        return response

    paginator = Paginator(ordered_entries, page_size)
    page_obj = paginator.get_page(request.GET.get("page"))
    for entry in page_obj.object_list:
        entry.audit_branch_label = (entry.branch_name or "").strip() or entry_branch_labels.get(entry.id, "-")
        entry.audit_model_label = _audit_model_label(entry.model_name)
        entry.audit_action_label = _audit_action_label(entry.action)

    context = {
        "page_obj": page_obj,
        "total_records": paginator.count,
        "model_options": model_options,
        "action_options": action_options,
        "search_query": search_query,
        "selected_model_name": model_name,
        "selected_action": action,
        "selected_branch_name": branch_name,
        "branch_options": branch_options,
        "date_from": date_from,
        "date_to": date_to,
        "page_size": page_size,
        "list_query_string": list_query_string,
    }
    return render(request, "settings/scorecard_permission_audit_trail.html", context)
