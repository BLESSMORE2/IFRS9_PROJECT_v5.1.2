from __future__ import annotations

from collections import Counter, defaultdict, deque
from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any

import openpyxl
from django.contrib import messages
from django.core.paginator import Paginator
from django.http import HttpRequest, HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from scorecard.functions_view.audit import log_ifrs9_results_audit
from scorecard.functions_view.main_customer_lookup import (
    current_branch_display_name,
    get_request_branch_names,
    get_request_branch_scope,
    is_all_branches_selected,
)
from scorecard.functions_view.validation_export_cache import (
    cache_validation_payload,
    cache_validation_export,
    get_cached_validation_payload,
    get_cached_validation_export,
    validation_payload_cache_key,
    validation_export_cache_key,
)
from scorecard.models import HistoricalScore


GRADE_ORDER = ("A1", "A2", "A3", "A4", "B1", "B2", "B3", "B4", "C", "D", "E")
GRADE_RANK = {grade: index for index, grade in enumerate(GRADE_ORDER)}
NPL_GRADE_SET = {"C", "D", "E"}
SUPERSCRIPT_TO_DIGIT = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹", "0123456789")
MOVEMENT_OPTIONS = (
    ("", "All movements"),
    ("upgrade", "Upgraded"),
    ("downgrade", "Downgraded"),
    ("unchanged", "Grade unchanged"),
    ("changed", "Grade changed"),
    ("incomplete", "Incomplete grade"),
    ("new", "New this month"),
    ("exited", "Exited since prior month"),
)
VALID_MOVEMENT_FILTERS = {value for value, _label in MOVEMENT_OPTIONS}
PAGE_SIZE_OPTIONS = (10, 25, 50, 100)


def _clean(value: Any) -> str:
    return str(value or "").strip()


def _parse_date(value: str | None) -> date | None:
    try:
        return date.fromisoformat(_clean(value)) if _clean(value) else None
    except ValueError:
        return None


def _decimal(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value))
    except Exception:
        return Decimal("0")


def _optional_decimal(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value))
    except Exception:
        return None


def _percent(numerator: Any, denominator: Any) -> float:
    denominator_value = _decimal(denominator)
    if denominator_value == 0:
        return 0.0
    return float((_decimal(numerator) / denominator_value) * Decimal("100"))


def _plain_grade(value: Any) -> str:
    return _clean(value).upper().replace(" ", "").translate(SUPERSCRIPT_TO_DIGIT)


def _effective_grade(row: dict[str, Any] | None) -> str:
    if not row:
        return ""
    return _clean(row.get("basel_override_grade")) or _clean(row.get("basel_ii_grade"))


def _grade_sort_key(value: str) -> tuple[int, str]:
    plain_grade = _plain_grade(value)
    return GRADE_RANK.get(plain_grade, len(GRADE_RANK)), plain_grade


def _classify_grade_movement(previous_grade: str, current_grade: str) -> tuple[str, str, int | None]:
    previous_plain = _plain_grade(previous_grade)
    current_plain = _plain_grade(current_grade)
    if not previous_plain or not current_plain:
        return "incomplete", "Incomplete grade", None
    if previous_plain == current_plain:
        return "unchanged", "Grade unchanged", 0

    previous_rank = GRADE_RANK.get(previous_plain)
    current_rank = GRADE_RANK.get(current_plain)
    if previous_rank is None or current_rank is None:
        return "changed", "Grade changed", None

    notch_change = previous_rank - current_rank
    if notch_change > 0:
        return "upgrade", "Upgraded", notch_change
    return "downgrade", "Downgraded", notch_change


def _score_direction(previous_score: Decimal | None, current_score: Decimal | None) -> str:
    if previous_score is None or current_score is None:
        return "Not comparable"
    if current_score > previous_score:
        return "Improved"
    if current_score < previous_score:
        return "Declined"
    return "Unchanged"


def _historical_scope_queryset(request: HttpRequest, selected_branch: str = ""):
    branch_names = get_request_branch_names(request)
    queryset = HistoricalScore.objects.filter(branch_name__in=branch_names) if branch_names else HistoricalScore.objects.none()
    if selected_branch:
        queryset = queryset.filter(branch_name__iexact=selected_branch)
    return queryset, branch_names


def _available_historical_dates(queryset) -> list[date]:
    return list(queryset.order_by("-reporting_date").values_list("reporting_date", flat=True).distinct())


def _resolve_comparison_dates(
    available_dates: list[date],
    current_value: str | None,
    previous_value: str | None,
) -> tuple[date | None, date | None]:
    current_date = _parse_date(current_value)
    if current_date not in available_dates:
        current_date = available_dates[0] if available_dates else None

    previous_date = _parse_date(previous_value)
    if previous_date not in available_dates or current_date is None or previous_date >= current_date:
        previous_date = next((item for item in available_dates if current_date and item < current_date), None)
    return current_date, previous_date


def _historical_values_for_dates(
    queryset,
    previous_date: date,
    current_date: date,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    rows_by_date = {previous_date: [], current_date: []}
    rows = (
        queryset.filter(reporting_date__in=(previous_date, current_date))
        .values(
            "reporting_date",
            "branch_name",
            "customer_id",
            "customer_name",
            "basel_ii_score",
            "basel_ii_grade",
            "basel_override_grade",
            "ifrs_9_score",
            "has_active_loan",
            "has_active_overdraft",
        )
        .order_by("reporting_date", "branch_name", "customer_id")
    )
    for row in rows:
        reporting_date = row.pop("reporting_date")
        rows_by_date[reporting_date].append(row)
    return rows_by_date[previous_date], rows_by_date[current_date]


def _build_detail_row(
    previous_row: dict[str, Any] | None,
    current_row: dict[str, Any] | None,
) -> dict[str, Any]:
    previous_score = _optional_decimal(previous_row.get("basel_ii_score")) if previous_row else None
    current_score = _optional_decimal(current_row.get("basel_ii_score")) if current_row else None
    previous_grade = _effective_grade(previous_row)
    current_grade = _effective_grade(current_row)

    if previous_row is None:
        movement_key, movement_label, grade_notches = "new", "New this month", None
    elif current_row is None:
        movement_key, movement_label, grade_notches = "exited", "Exited since prior month", None
    else:
        movement_key, movement_label, grade_notches = _classify_grade_movement(previous_grade, current_grade)

    score_delta = None
    if previous_score is not None and current_score is not None:
        score_delta = current_score - previous_score

    previous_branch = _clean(previous_row.get("branch_name")) if previous_row else ""
    current_branch = _clean(current_row.get("branch_name")) if current_row else ""
    customer_id = _clean((current_row or previous_row or {}).get("customer_id"))
    grade_change = f"{previous_grade or 'Not scored'} -> {current_grade or 'Not scored'}"
    return {
        "customer_id": customer_id,
        "customer_name": _clean((current_row or previous_row or {}).get("customer_name")),
        "branch_name": current_branch or previous_branch,
        "previous_branch": previous_branch,
        "current_branch": current_branch,
        "branch_transfer": bool(previous_branch and current_branch and previous_branch.casefold() != current_branch.casefold()),
        "previous_score": previous_score,
        "current_score": current_score,
        "score_delta": score_delta,
        "score_direction": _score_direction(previous_score, current_score),
        "previous_basel_grade": _clean(previous_row.get("basel_ii_grade")) if previous_row else "",
        "previous_override_grade": _clean(previous_row.get("basel_override_grade")) if previous_row else "",
        "previous_effective_grade": previous_grade,
        "current_basel_grade": _clean(current_row.get("basel_ii_grade")) if current_row else "",
        "current_override_grade": _clean(current_row.get("basel_override_grade")) if current_row else "",
        "current_effective_grade": current_grade,
        "grade_change": grade_change,
        "grade_notches": grade_notches,
        "movement_key": movement_key,
        "movement_label": movement_label,
        "has_previous": previous_row is not None,
        "has_current": current_row is not None,
    }


def _pair_historical_rows(
    previous_rows: list[dict[str, Any]],
    current_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    previous_by_customer: dict[str, list[dict[str, Any]]] = defaultdict(list)
    current_by_customer: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in previous_rows:
        previous_by_customer[_clean(row.get("customer_id"))].append(row)
    for row in current_rows:
        current_by_customer[_clean(row.get("customer_id"))].append(row)

    details: list[dict[str, Any]] = []
    for customer_id in sorted(set(previous_by_customer) | set(current_by_customer)):
        previous_remaining = list(previous_by_customer.get(customer_id, []))
        current_remaining = list(current_by_customer.get(customer_id, []))
        previous_by_branch: dict[str, deque] = defaultdict(deque)
        for row in previous_remaining:
            previous_by_branch[_clean(row.get("branch_name")).casefold()].append(row)

        unmatched_current = []
        matched_previous_ids: set[int] = set()
        for current_row in current_remaining:
            branch_key = _clean(current_row.get("branch_name")).casefold()
            if previous_by_branch[branch_key]:
                previous_row = previous_by_branch[branch_key].popleft()
                matched_previous_ids.add(id(previous_row))
                details.append(_build_detail_row(previous_row, current_row))
            else:
                unmatched_current.append(current_row)

        unmatched_previous = [row for row in previous_remaining if id(row) not in matched_previous_ids]
        pair_count = min(len(unmatched_previous), len(unmatched_current))
        for index in range(pair_count):
            details.append(_build_detail_row(unmatched_previous[index], unmatched_current[index]))
        for current_row in unmatched_current[pair_count:]:
            details.append(_build_detail_row(None, current_row))
        for previous_row in unmatched_previous[pair_count:]:
            details.append(_build_detail_row(previous_row, None))

    return sorted(details, key=lambda row: (row["branch_name"].casefold(), row["customer_name"].casefold(), row["customer_id"]))


def _movement_summary(details: list[dict[str, Any]]) -> dict[str, Any]:
    counts = Counter(row["movement_key"] for row in details)
    current_scores = [row["current_score"] for row in details if row["current_score"] is not None]
    previous_scores = [row["previous_score"] for row in details if row["previous_score"] is not None]
    comparable_deltas = [row["score_delta"] for row in details if row["score_delta"] is not None]
    return {
        "current_customers": sum(1 for row in details if row["has_current"]),
        "previous_customers": sum(1 for row in details if row["has_previous"]),
        "matched_customers": sum(1 for row in details if row["has_current"] and row["has_previous"]),
        "upgrades": counts["upgrade"],
        "downgrades": counts["downgrade"],
        "unchanged": counts["unchanged"],
        "changed": counts["changed"],
        "incomplete": counts["incomplete"],
        "new": counts["new"],
        "exited": counts["exited"],
        "branch_transfers": sum(1 for row in details if row["branch_transfer"]),
        "average_current_score": float(sum(current_scores) / len(current_scores)) if current_scores else None,
        "average_previous_score": float(sum(previous_scores) / len(previous_scores)) if previous_scores else None,
        "average_score_delta": float(sum(comparable_deltas) / len(comparable_deltas)) if comparable_deltas else None,
    }


def _grade_migration(details: list[dict[str, Any]]) -> dict[str, Any]:
    comparable = [
        row for row in details
        if row["has_previous"] and row["has_current"]
        and row["previous_effective_grade"] and row["current_effective_grade"]
    ]
    previous_grades = sorted({row["previous_effective_grade"] for row in comparable}, key=_grade_sort_key)
    current_grades = sorted({row["current_effective_grade"] for row in comparable}, key=_grade_sort_key)
    counts = Counter((row["previous_effective_grade"], row["current_effective_grade"]) for row in comparable)
    rows = []
    for previous_grade in previous_grades:
        cells = []
        for current_grade in current_grades:
            if previous_grade == current_grade:
                movement_key = "unchanged"
            elif GRADE_RANK.get(current_grade, 999) < GRADE_RANK.get(previous_grade, 999):
                movement_key = "upgrade"
            else:
                movement_key = "downgrade"
            cells.append(
                {
                    "current_grade": current_grade,
                    "count": counts[(previous_grade, current_grade)],
                    "movement_key": movement_key,
                }
            )
        row_counts = [cell["count"] for cell in cells]
        rows.append({"grade": previous_grade, "cells": cells, "counts": row_counts, "total": sum(row_counts)})
    return {
        "current_grades": current_grades,
        "rows": rows,
        "column_totals": [sum(counts[(previous_grade, grade)] for previous_grade in previous_grades) for grade in current_grades],
        "total": len(comparable),
    }


def _grade_distribution(details: list[dict[str, Any]]) -> list[dict[str, Any]]:
    previous_counts = Counter(row["previous_effective_grade"] for row in details if row["previous_effective_grade"])
    current_counts = Counter(row["current_effective_grade"] for row in details if row["current_effective_grade"])
    grades = sorted(set(previous_counts) | set(current_counts), key=_grade_sort_key)
    return [
        {
            "grade": grade,
            "previous_count": previous_counts[grade],
            "current_count": current_counts[grade],
            "change": current_counts[grade] - previous_counts[grade],
        }
        for grade in grades
    ]


def _npl_row(historical_customer_codes: set[str], npl_customer_codes: set[str]) -> dict[str, Any]:
    historical_customers = len(historical_customer_codes)
    npl_customers = len(npl_customer_codes)
    return {
        "historical_customers": historical_customers,
        "npl_customers": npl_customers,
        "customer_ratio": _percent(npl_customers, historical_customers),
    }


def _npl_payload(
    reporting_date: date | None,
    historical_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    empty_payload = {
        "available": False,
        "reporting_date": reporting_date,
        "rows": [],
        "overall": _npl_row(set(), set()),
        "source_mode": "unavailable",
        "source_reporting_date": None,
    }
    if reporting_date is None:
        return empty_payload

    historical_codes_by_branch: dict[str, set[str]] = defaultdict(set)
    npl_codes_by_branch: dict[str, set[str]] = defaultdict(set)
    overall_historical_codes: set[str] = set()
    overall_npl_codes: set[str] = set()
    for row in historical_rows:
        customer_code = _clean(row.get("customer_id")).casefold()
        branch_name = _clean(row.get("branch_name")) or "Unassigned"
        if not customer_code:
            continue
        has_basel_score = row.get("basel_ii_score") is not None or bool(_effective_grade(row))
        if not has_basel_score:
            continue
        historical_codes_by_branch[branch_name].add(customer_code)
        overall_historical_codes.add(customer_code)
        if _plain_grade(_effective_grade(row)) in NPL_GRADE_SET:
            npl_codes_by_branch[branch_name].add(customer_code)
            overall_npl_codes.add(customer_code)

    if not overall_historical_codes:
        return empty_payload
    rows = [
        {"branch_name": branch_name}
        | _npl_row(historical_codes, npl_codes_by_branch.get(branch_name, set()))
        for branch_name, historical_codes in sorted(
            historical_codes_by_branch.items(), key=lambda item: item[0].casefold()
        )
    ]
    return {
        "available": True,
        "reporting_date": reporting_date,
        "rows": rows,
        "overall": _npl_row(overall_historical_codes, overall_npl_codes),
        "source_mode": "historical_basel_population",
        "source_reporting_date": reporting_date,
    }


def _duplicate_customer_rows(
    historical_rows: list[dict[str, Any]],
    reporting_date: date,
    period_label: str,
) -> list[dict[str, Any]]:
    rows_by_customer: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in historical_rows:
        customer_id = _clean(row.get("customer_id"))
        if customer_id:
            rows_by_customer[customer_id.casefold()].append(row)

    duplicate_rows = []
    for customer_rows in rows_by_customer.values():
        if len(customer_rows) <= 1:
            continue
        ordered_rows = sorted(
            customer_rows,
            key=lambda row: (
                _clean(row.get("branch_name")).casefold(),
                _clean(row.get("customer_name")).casefold(),
            ),
        )
        occurrence_count = len(ordered_rows)
        for occurrence_number, row in enumerate(ordered_rows, start=1):
            duplicate_rows.append(
                {
                    "period_label": period_label,
                    "reporting_date": reporting_date,
                    "customer_id": _clean(row.get("customer_id")),
                    "customer_name": _clean(row.get("customer_name")),
                    "branch_name": _clean(row.get("branch_name")) or "Unassigned",
                    "occurrence_number": occurrence_number,
                    "occurrence_count": occurrence_count,
                    "basel_ii_score": _optional_decimal(row.get("basel_ii_score")),
                    "basel_ii_grade": _clean(row.get("basel_ii_grade")),
                    "basel_override_grade": _clean(row.get("basel_override_grade")),
                    "effective_grade": _effective_grade(row),
                }
            )
    return sorted(
        duplicate_rows,
        key=lambda row: (
            row["reporting_date"],
            row["customer_id"].casefold(),
            row["branch_name"].casefold(),
        ),
    )


def _branch_summary(
    details: list[dict[str, Any]],
    previous_npl: dict[str, Any],
    current_npl: dict[str, Any],
) -> list[dict[str, Any]]:
    buckets: dict[str, dict[str, Any]] = {}

    def bucket_for(branch_name: str) -> dict[str, Any]:
        key = _clean(branch_name).casefold()
        if key not in buckets:
            buckets[key] = {
                "branch_name": _clean(branch_name) or "Unassigned",
                "previous_customers": 0,
                "current_customers": 0,
                "matched_customers": 0,
                "upgrades": 0,
                "downgrades": 0,
                "unchanged": 0,
                "new": 0,
                "exited": 0,
                "score_deltas": [],
            }
        return buckets[key]

    for detail in details:
        if detail["has_previous"]:
            bucket_for(detail["previous_branch"])["previous_customers"] += 1
        if detail["has_current"]:
            bucket_for(detail["current_branch"])["current_customers"] += 1

        # Movements belong to the current branch; exits stay with their previous branch.
        bucket = bucket_for(detail["current_branch"] or detail["previous_branch"])
        bucket["matched_customers"] += int(detail["has_previous"] and detail["has_current"])
        movement_bucket = {
            "upgrade": "upgrades",
            "downgrade": "downgrades",
            "unchanged": "unchanged",
            "new": "new",
            "exited": "exited",
        }.get(detail["movement_key"])
        if movement_bucket:
            bucket[movement_bucket] += 1
        if detail["score_delta"] is not None:
            bucket["score_deltas"].append(detail["score_delta"])

    previous_npl_map = {row["branch_name"].casefold(): row for row in previous_npl.get("rows", [])}
    current_npl_map = {row["branch_name"].casefold(): row for row in current_npl.get("rows", [])}
    for source in (previous_npl_map, current_npl_map):
        for row in source.values():
            bucket_for(row["branch_name"])

    rows = []
    for key, bucket in sorted(buckets.items(), key=lambda item: item[1]["branch_name"].casefold()):
        previous_row = previous_npl_map.get(key, {})
        current_row = current_npl_map.get(key, {})
        deltas = bucket.pop("score_deltas")
        bucket["average_score_delta"] = float(sum(deltas) / len(deltas)) if deltas else None
        bucket["previous_npl"] = previous_row
        bucket["current_npl"] = current_row
        bucket["previous_npl_ratio"] = previous_row.get("customer_ratio") if previous_row else None
        bucket["current_npl_ratio"] = current_row.get("customer_ratio") if current_row else None
        if bucket["previous_npl_ratio"] is not None and bucket["current_npl_ratio"] is not None:
            bucket["npl_ratio_delta"] = bucket["current_npl_ratio"] - bucket["previous_npl_ratio"]
        else:
            bucket["npl_ratio_delta"] = None
        rows.append(bucket)
    return rows


def _filter_details(details: list[dict[str, Any]], search: str, movement_filter: str) -> list[dict[str, Any]]:
    filtered = details
    if movement_filter:
        filtered = [row for row in filtered if row["movement_key"] == movement_filter]
    search_key = _clean(search).casefold()
    if search_key:
        filtered = [
            row for row in filtered
            if search_key in " ".join(
                [
                    row["customer_id"],
                    row["customer_name"],
                    row["branch_name"],
                    row["previous_effective_grade"],
                    row["current_effective_grade"],
                    row["movement_label"],
                ]
            ).casefold()
        ]
    return filtered


def _build_validation_payload(
    historical_queryset,
    previous_date: date,
    current_date: date,
) -> dict[str, Any]:
    previous_rows, current_rows = _historical_values_for_dates(
        historical_queryset,
        previous_date,
        current_date,
    )
    details = _pair_historical_rows(previous_rows, current_rows)
    previous_npl = _npl_payload(previous_date, previous_rows)
    current_npl = _npl_payload(current_date, current_rows)
    previous_duplicate_rows = _duplicate_customer_rows(previous_rows, previous_date, "Previous")
    current_duplicate_rows = _duplicate_customer_rows(current_rows, current_date, "Current")
    previous_overall_npl_ratio = previous_npl["overall"].get("customer_ratio") if previous_npl["available"] else None
    current_overall_npl_ratio = current_npl["overall"].get("customer_ratio") if current_npl["available"] else None
    overall_npl_ratio_delta = (
        current_overall_npl_ratio - previous_overall_npl_ratio
        if previous_overall_npl_ratio is not None and current_overall_npl_ratio is not None
        else None
    )
    previous_duplicates = len({row["customer_id"].casefold() for row in previous_duplicate_rows})
    current_duplicates = len({row["customer_id"].casefold() for row in current_duplicate_rows})
    grade_distribution = _grade_distribution(details)
    return {
        "previous_date": previous_date,
        "current_date": current_date,
        "details": details,
        "summary": _movement_summary(details),
        "grade_migration": _grade_migration(details),
        "grade_distribution": grade_distribution,
        "grade_distribution_totals": {
            "previous_count": sum(row["previous_count"] for row in grade_distribution),
            "current_count": sum(row["current_count"] for row in grade_distribution),
            "change": sum(row["change"] for row in grade_distribution),
        },
        "previous_npl": previous_npl,
        "current_npl": current_npl,
        "duplicate_customers": previous_duplicate_rows + current_duplicate_rows,
        "overall_npl_ratio_delta": overall_npl_ratio_delta,
        "branch_summary": _branch_summary(details, previous_npl, current_npl),
        "data_quality": {
            "previous_duplicate_customer_codes": previous_duplicates,
            "current_duplicate_customer_codes": current_duplicates,
            "missing_previous_scores": sum(1 for row in details if row["has_previous"] and row["previous_score"] is None),
            "missing_current_scores": sum(1 for row in details if row["has_current"] and row["current_score"] is None),
            "incomplete_grades": sum(1 for row in details if row["movement_key"] == "incomplete"),
        },
    }


def _get_validation_payload(
    historical_queryset,
    previous_date: date,
    current_date: date,
    branch_names: list[str],
    selected_branch: str,
) -> dict[str, Any]:
    cache_key = validation_payload_cache_key(
        report_name="basel",
        queryset=historical_queryset,
        previous_date=previous_date,
        current_date=current_date,
        branch_names=branch_names,
        selected_branch=selected_branch,
    )
    cached_payload = get_cached_validation_payload(cache_key)
    if cached_payload is not None:
        return cached_payload

    payload = _build_validation_payload(historical_queryset, previous_date, current_date)
    cache_validation_payload(cache_key, payload)
    return payload


def _style_title(worksheet, title: str, total_columns: int) -> None:
    dark_fill = PatternFill("solid", fgColor="0B2D52")
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(total_columns, 8))
    worksheet["A1"] = title
    worksheet["A1"].fill = dark_fill
    worksheet["A1"].font = Font(name="Arial", color="FFFFFF", bold=True, size=15)
    worksheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[1].height = 27
    for column in range(2, max(total_columns, 8) + 1):
        worksheet.cell(1, column).fill = dark_fill


def _style_header_row(worksheet, row_number: int, column_count: int) -> None:
    fill = PatternFill("solid", fgColor="0E4377")
    border = Border(bottom=Side(style="thin", color="CAD8E6"))
    for column in range(1, column_count + 1):
        cell = worksheet.cell(row_number, column)
        cell.fill = fill
        cell.font = Font(name="Arial", color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    worksheet.row_dimensions[row_number].height = 28


def _fit_columns(worksheet, max_width: int = 34) -> None:
    for column in range(1, worksheet.max_column + 1):
        width = 12
        for row in range(1, min(worksheet.max_row, 250) + 1):
            value = worksheet.cell(row, column).value
            if value not in (None, ""):
                width = max(width, min(len(str(value)) + 2, max_width))
        worksheet.column_dimensions[get_column_letter(column)].width = width


def _add_table(worksheet, name: str, header_row: int) -> None:
    if worksheet.max_row <= header_row:
        return
    header_columns = [
        column
        for column in range(1, worksheet.max_column + 1)
        if worksheet.cell(header_row, column).value not in (None, "")
    ]
    if not header_columns:
        return
    last_header_column = max(header_columns)
    reference = f"A{header_row}:{get_column_letter(last_header_column)}{worksheet.max_row}"
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def _build_validation_workbook(
    payload: dict[str, Any],
    detail_rows: list[dict[str, Any]],
    branch_scope_label: str,
    selected_branch: str,
    movement_filter: str,
    search: str,
):
    workbook = openpyxl.Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Validation Summary"
    detail_sheet = workbook.create_sheet("Individual Customers")
    duplicate_sheet = workbook.create_sheet("Duplicate Customers")
    migration_sheet = workbook.create_sheet("Grade Migration")
    distribution_sheet = workbook.create_sheet("Grade Population by Date")
    npl_sheet = workbook.create_sheet("NPL by Branch")
    definitions_sheet = workbook.create_sheet("Definitions")
    for worksheet in workbook.worksheets:
        worksheet.sheet_view.showGridLines = False

    summary = payload["summary"]
    _style_title(summary_sheet, "NPL Migration Report", 12)
    summary_sheet.append([
        "Previous Month-End", payload["previous_date"], "Current Month-End", payload["current_date"],
        "Branch Scope", selected_branch or branch_scope_label, "Generated", timezone.localtime().replace(tzinfo=None),
    ])
    summary_sheet.append([])
    summary_sheet.append(["Validation Metric", "Value"])
    _style_header_row(summary_sheet, 4, 2)
    metrics = [
        ("Previous customers", summary["previous_customers"]),
        ("Current customers", summary["current_customers"]),
        ("Matched customers", summary["matched_customers"]),
        ("Upgraded", summary["upgrades"]),
        ("Downgraded", summary["downgrades"]),
        ("Grade unchanged", summary["unchanged"]),
        ("New this month", summary["new"]),
        ("Exited since prior month", summary["exited"]),
        ("Branch transfers", summary["branch_transfers"]),
        ("Average previous weighted score", summary["average_previous_score"]),
        ("Average current weighted score", summary["average_current_score"]),
        ("Average weighted score change (pp)", summary["average_score_delta"]),
        ("Previous overall NPL ratio", payload["previous_npl"]["overall"]["customer_ratio"] if payload["previous_npl"]["available"] else None),
        ("Current overall NPL ratio", payload["current_npl"]["overall"]["customer_ratio"] if payload["current_npl"]["available"] else None),
    ]
    for metric, value in metrics:
        summary_sheet.append([metric, value])
        value_cell = summary_sheet.cell(summary_sheet.max_row, 2)
        if "(pp)" in metric:
            value_cell.number_format = '+0.00" pp";-0.00" pp";0.00" pp"'
        elif "weighted score" in metric.casefold() or "npl ratio" in metric.casefold():
            value_cell.number_format = '0.00"%"'

    branch_header_row = summary_sheet.max_row + 2
    summary_sheet.cell(branch_header_row, 1, "Branch Validation Summary")
    summary_sheet.cell(branch_header_row, 1).font = Font(name="Arial", bold=True, size=12, color="0B2D52")
    branch_columns = [
        "Branch", "Previous Customers", "Current Customers", "Matched", "Upgrades", "Downgrades",
        "Unchanged", "New", "Exited", "Average Score Change (pp)", "Previous NPL Ratio", "Current NPL Ratio", "NPL Change (pp)",
    ]
    summary_sheet.append(branch_columns)
    branch_table_header = summary_sheet.max_row
    _style_header_row(summary_sheet, branch_table_header, len(branch_columns))
    for row in payload["branch_summary"]:
        summary_sheet.append([
            row["branch_name"], row["previous_customers"], row["current_customers"], row["matched_customers"],
            row["upgrades"], row["downgrades"], row["unchanged"], row["new"], row["exited"],
            row["average_score_delta"], row["previous_npl_ratio"], row["current_npl_ratio"], row["npl_ratio_delta"],
        ])
    for row in summary_sheet.iter_rows(min_row=branch_table_header + 1):
        row[9].number_format = '+0.00" pp";-0.00" pp";0.00" pp"'
        row[10].number_format = '0.00"%"'
        row[11].number_format = '0.00"%"'
        row[12].number_format = '+0.00" pp";-0.00" pp";0.00" pp"'
    summary_sheet.freeze_panes = "A5"
    _add_table(summary_sheet, "BaselValidationBranchSummary", branch_table_header)

    detail_columns = [
        "Branch", "Customer Code", "Customer Name", "Previous Effective Grade", "Current Effective Grade",
        "Grade Change", "Grade Notch Change", "Grade Movement", "Previous Weighted Score",
        "Current Weighted Score", "Score Change (pp)", "Score Direction", "Previous Branch",
        "Branch Transfer", "Previous Basel Grade", "Previous Override Grade", "Current Basel Grade",
        "Current Override Grade",
    ]
    _style_title(detail_sheet, "Individual Customer Basel II Grade and Weighted Score Changes", len(detail_columns))
    detail_sheet.append([
        "Previous Month-End", payload["previous_date"], "Current Month-End", payload["current_date"],
        "Branch Scope", selected_branch or branch_scope_label, "Movement Filter", movement_filter or "All",
        "Search", search or "All customers",
    ])
    detail_sheet.append([])
    detail_sheet.append(detail_columns)
    _style_header_row(detail_sheet, 4, len(detail_columns))
    for row in detail_rows:
        detail_sheet.append([
            row["branch_name"], row["customer_id"], row["customer_name"], row["previous_effective_grade"],
            row["current_effective_grade"], row["grade_change"], row["grade_notches"], row["movement_label"],
            row["previous_score"], row["current_score"], row["score_delta"], row["score_direction"],
            row["previous_branch"], "Yes" if row["branch_transfer"] else "No", row["previous_basel_grade"],
            row["previous_override_grade"], row["current_basel_grade"], row["current_override_grade"],
        ])
    for row in detail_sheet.iter_rows(min_row=5, min_col=9, max_col=11):
        for cell in row:
            cell.number_format = '0.00"%"' if cell.column in (9, 10) else '+0.00" pp";-0.00" pp";0.00" pp"'
    if detail_sheet.max_row >= 5:
        movement_range = f"H5:H{detail_sheet.max_row}"
        detail_sheet.conditional_formatting.add(
            movement_range,
            FormulaRule(
                formula=['$H5="Upgraded"'],
                fill=PatternFill("solid", fgColor="E6F7EF"),
                font=Font(color="087443", bold=True),
            ),
        )
        detail_sheet.conditional_formatting.add(
            movement_range,
            FormulaRule(
                formula=['$H5="Downgraded"'],
                fill=PatternFill("solid", fgColor="FFF0F1"),
                font=Font(color="B4232C", bold=True),
            ),
        )
        detail_sheet.conditional_formatting.add(
            movement_range,
            FormulaRule(
                formula=['OR($H5="New this month",$H5="Exited since prior month")'],
                fill=PatternFill("solid", fgColor="FFF4DF"),
                font=Font(color="8C4B00", bold=True),
            ),
        )
    detail_sheet.freeze_panes = "A5"
    _add_table(detail_sheet, "BaselCustomerMovements", 4)
    detail_sheet.sheet_properties.tabColor = "176BB3"

    duplicate_columns = [
        "Comparison Period", "Reporting Date", "Customer Code", "Customer Name", "Branch",
        "Occurrence", "Duplicate Count", "Basel II Score", "Basel II Grade", "Basel Override Grade",
        "Effective Basel Grade",
    ]
    _style_title(duplicate_sheet, "Historical Basel II Duplicate Customer Details", len(duplicate_columns))
    duplicate_sheet.append([
        "Definition",
        "A duplicate is a customer code appearing in more than one historical row for the same reporting date.",
        "Previous Month-End",
        payload["previous_date"],
        "Current Month-End",
        payload["current_date"],
    ])
    duplicate_sheet.append([])
    duplicate_sheet.append(duplicate_columns)
    _style_header_row(duplicate_sheet, 4, len(duplicate_columns))
    for row in payload["duplicate_customers"]:
        duplicate_sheet.append([
            row["period_label"], row["reporting_date"], row["customer_id"], row["customer_name"],
            row["branch_name"], f"{row['occurrence_number']} of {row['occurrence_count']}",
            row["occurrence_count"], row["basel_ii_score"], row["basel_ii_grade"],
            row["basel_override_grade"], row["effective_grade"],
        ])
    if payload["duplicate_customers"]:
        for cell in duplicate_sheet.iter_cols(min_col=8, max_col=8, min_row=5):
            for item in cell:
                item.number_format = '0.00"%"'
        duplicate_sheet.conditional_formatting.add(
            f"G5:G{duplicate_sheet.max_row}",
            FormulaRule(
                formula=["$G5>1"],
                fill=PatternFill("solid", fgColor="FFF0F1"),
                font=Font(color="B4232C", bold=True),
            ),
        )
        _add_table(duplicate_sheet, "BaselDuplicateCustomers", 4)
    else:
        duplicate_sheet.append(["No duplicate customer codes were found for the selected comparison dates."])
    duplicate_sheet.freeze_panes = "A5"
    duplicate_sheet.sheet_properties.tabColor = "E5484D"

    migration = payload["grade_migration"]
    migration_columns = [f"Previous Grade ({payload['previous_date']})"] + migration["current_grades"] + ["Matched Previous Total"]
    _style_title(migration_sheet, "Basel II Grade Migration Matrix", len(migration_columns))
    migration_sheet.append(["Each cell is one matched-customer movement count from the previous row grade to the current column grade. Separate date populations are in Grade Population by Date."])
    migration_sheet.cell(3, 1, f"Previous Grade ↓ ({payload['previous_date']})")
    migration_sheet.merge_cells(start_row=3, start_column=2, end_row=3, end_column=1 + len(migration["current_grades"]))
    migration_sheet.cell(3, 2, f"Current Grade → ({payload['current_date']})")
    migration_sheet.cell(3, len(migration_columns), "Matched Previous Total")
    for column in range(1, len(migration_columns) + 1):
        cell = migration_sheet.cell(3, column)
        cell.fill = PatternFill("solid", fgColor="176BB3" if 1 < column < len(migration_columns) else "082F58")
        cell.font = Font(name="Arial", color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    migration_sheet.row_dimensions[3].height = 24
    migration_sheet.append(migration_columns)
    _style_header_row(migration_sheet, 4, len(migration_columns))
    for row in migration["rows"]:
        migration_sheet.append([row["grade"]] + row["counts"] + [row["total"]])
        row_number = migration_sheet.max_row
        migration_sheet.cell(row_number, 1).font = Font(name="Arial", bold=True, color="0B2D52")
        migration_sheet.cell(row_number, 1).fill = PatternFill("solid", fgColor="F3F8FD")
        for column_number, cell_payload in enumerate(row["cells"], start=2):
            if not cell_payload["count"]:
                migration_sheet.cell(row_number, column_number).font = Font(name="Arial", color="9AAABA")
                continue
            movement_styles = {
                "unchanged": ("DCEEFF", "0B4F87"),
                "upgrade": ("E6F7EF", "087443"),
                "downgrade": ("FFF0F1", "B4232C"),
            }
            fill_color, font_color = movement_styles[cell_payload["movement_key"]]
            migration_sheet.cell(row_number, column_number).fill = PatternFill("solid", fgColor=fill_color)
            migration_sheet.cell(row_number, column_number).font = Font(name="Arial", color=font_color, bold=True)
        migration_sheet.cell(row_number, len(migration_columns)).fill = PatternFill("solid", fgColor="EDF4FA")
        migration_sheet.cell(row_number, len(migration_columns)).font = Font(name="Arial", color="0B2D52", bold=True)
    migration_sheet.append(["Matched Current Grade Totals →"] + migration["column_totals"] + [migration["total"]])
    for cell in migration_sheet[migration_sheet.max_row]:
        cell.fill = PatternFill("solid", fgColor="0B355E")
        cell.font = Font(name="Arial", color="FFFFFF", bold=True)
    migration_sheet.freeze_panes = "B5"

    distribution_columns = [
        "Grade",
        f"Previous Customers ({payload['previous_date']})",
        f"Current Customers ({payload['current_date']})",
        "Change",
    ]
    _style_title(distribution_sheet, "Grade Population by Date", len(distribution_columns))
    distribution_sheet.append(["Previous Month-End", payload["previous_date"], "Current Month-End", payload["current_date"]])
    distribution_sheet.append([])
    distribution_sheet.append(distribution_columns)
    _style_header_row(distribution_sheet, 4, len(distribution_columns))
    for row in payload["grade_distribution"]:
        distribution_sheet.append([row["grade"], row["previous_count"], row["current_count"], row["change"]])
    distribution_totals = payload["grade_distribution_totals"]
    distribution_sheet.append([
        "Total graded customers",
        distribution_totals["previous_count"],
        distribution_totals["current_count"],
        distribution_totals["change"],
    ])
    for cell in distribution_sheet[distribution_sheet.max_row]:
        cell.fill = PatternFill("solid", fgColor="EAF3FB")
        cell.font = Font(name="Arial", bold=True, color="0B2D52")
    distribution_sheet.freeze_panes = "A5"
    _add_table(distribution_sheet, "BaselGradeDistribution", 4)

    npl_columns = [
        "Branch", "Previous Historical Basel Customers", "Previous C/D/E Customers", "Previous NPL Ratio",
        "Current Historical Basel Customers", "Current C/D/E Customers", "Current NPL Ratio", "NPL Ratio Change (pp)",
    ]
    _style_title(npl_sheet, "Dashboard-Formula NPL Ratios by Branch", len(npl_columns))
    npl_sheet.append([
        "Previous Month-End", payload["previous_date"],
        "Current Month-End", payload["current_date"],
        "Formula", "C/D/E historical Basel customers divided by all historical Basel customers",
        "Previous NPL Source", f"{payload['previous_npl'].get('source_mode')} ({payload['previous_npl'].get('source_reporting_date') or payload['previous_date']})",
        "Current NPL Source", f"{payload['current_npl'].get('source_mode')} ({payload['current_npl'].get('source_reporting_date') or payload['current_date']})",
    ])
    npl_sheet.append([])
    npl_sheet.append(npl_columns)
    _style_header_row(npl_sheet, 4, len(npl_columns))
    for row in payload["branch_summary"]:
        previous_npl = row["previous_npl"]
        current_npl = row["current_npl"]
        npl_sheet.append([
            row["branch_name"], previous_npl.get("historical_customers"), previous_npl.get("npl_customers"),
            row["previous_npl_ratio"], current_npl.get("historical_customers"), current_npl.get("npl_customers"),
            row["current_npl_ratio"], row["npl_ratio_delta"],
        ])
    previous_overall = payload["previous_npl"]["overall"] if payload["previous_npl"]["available"] else {}
    current_overall = payload["current_npl"]["overall"] if payload["current_npl"]["available"] else {}
    npl_sheet.append([
        "Overall selected scope",
        previous_overall.get("historical_customers"), previous_overall.get("npl_customers"),
        previous_overall.get("customer_ratio"), current_overall.get("historical_customers"),
        current_overall.get("npl_customers"), current_overall.get("customer_ratio"),
        payload["overall_npl_ratio_delta"],
    ])
    overall_row_number = npl_sheet.max_row
    for cell in npl_sheet[overall_row_number]:
        cell.font = Font(name="Arial", bold=True, color="0B2D52")
        cell.fill = PatternFill("solid", fgColor="EAF3FB")
    for row in npl_sheet.iter_rows(min_row=5):
        for column in (4, 7):
            row[column - 1].number_format = '0.00"%"'
        row[7].number_format = '+0.00" pp";-0.00" pp";0.00" pp"'
    npl_sheet.freeze_panes = "A5"
    _add_table(npl_sheet, "BaselValidationNPLByBranch", 4)

    _style_title(definitions_sheet, "Validation Definitions and Sources", 4)
    definitions_sheet.append(["Item", "Definition", "Source", "Notes"])
    _style_header_row(definitions_sheet, 2, 4)
    definitions = [
        ("Weighted score movement", "Current Basel II weighted score minus the previous month-end score.", "SCORECARD_HISTORICAL_SCORES", "Shown in percentage points."),
        ("Effective grade", "Basel override grade when present; otherwise the captured Basel II grade.", "SCORECARD_HISTORICAL_SCORES", "Preserves approved override decisions."),
        ("Upgrade / downgrade", "Movement through A1, A2, A3, A4, B1, B2, B3, B4, C, D, E.", "SCORECARD_HISTORICAL_SCORES", "A lower rank is a stronger grade."),
        ("NPL population", "Unique customers in the selected historical snapshot whose effective Basel grade is C, D, or E.", "SCORECARD_HISTORICAL_SCORES", "The effective grade uses the captured override grade when present; otherwise it uses the captured Basel II grade."),
        ("NPL ratio", "C/D/E historical Basel customers divided by all unique historical Basel customers, multiplied by 100.", "SCORECARD_HISTORICAL_SCORES", "Calculated separately by branch and for the complete selected scope; no current loan-table recheck is performed."),
        ("Branch scope", "Only branches assigned to the user and selected in the scorecard workspace.", "Scorecard branch access", "A selected branch narrows every workbook sheet."),
    ]
    for row in definitions:
        definitions_sheet.append(row)
    definitions_sheet.freeze_panes = "A3"
    _add_table(definitions_sheet, "BaselValidationDefinitions", 2)

    for worksheet in workbook.worksheets:
        _fit_columns(worksheet)
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_margins.left = 0.25
        worksheet.page_margins.right = 0.25
        worksheet.page_margins.top = 0.5
        worksheet.page_margins.bottom = 0.5

    return workbook


def _validation_query_string(request: HttpRequest) -> str:
    keys = ("current_date", "previous_date", "branch", "movement", "search", "page_size")
    parts = []
    for key in keys:
        value = _clean(request.GET.get(key))
        if value:
            from urllib.parse import quote_plus

            parts.append(f"{key}={quote_plus(value)}")
    return "&".join(parts)


def basel_validations_view(request: HttpRequest):
    branch_scope = get_request_branch_scope(request)
    branch_names = [name for name in get_request_branch_names(request) if name]
    selected_branch = _clean(request.GET.get("branch"))
    if selected_branch and selected_branch.casefold() not in {name.casefold() for name in branch_names}:
        messages.error(request, "Choose a branch within your assigned branch scope.")
        selected_branch = ""

    historical_queryset, _branch_names = _historical_scope_queryset(request, selected_branch)
    available_dates = _available_historical_dates(historical_queryset)
    current_date, previous_date = _resolve_comparison_dates(
        available_dates,
        request.GET.get("current_date"),
        request.GET.get("previous_date"),
    )
    movement_filter = _clean(request.GET.get("movement"))
    if movement_filter not in VALID_MOVEMENT_FILTERS:
        movement_filter = ""
    search = _clean(request.GET.get("search"))
    try:
        page_size = int(request.GET.get("page_size") or 10)
    except (TypeError, ValueError):
        page_size = 10
    if page_size not in PAGE_SIZE_OPTIONS:
        page_size = 10

    payload = None
    filtered_details = []
    page_obj = None
    if current_date and previous_date:
        payload = _get_validation_payload(
            historical_queryset,
            previous_date,
            current_date,
            branch_names,
            selected_branch,
        )
        filtered_details = _filter_details(payload["details"], search, movement_filter)
        page_obj = Paginator(filtered_details, page_size).get_page(request.GET.get("page") or 1)

    return render(
        request,
        "credit_scoreshifts/basel_validations.html",
        {
            "results_page_key": "basel_validations",
            "results_branch_scope_label": current_branch_display_name(request) or "No branch selected",
            "results_branch_mode_is_all": is_all_branches_selected(request),
            "results_branch_scope_count": len(branch_scope),
            "validation_branch_names": branch_names,
            "validation_selected_branch": selected_branch,
            "validation_available_dates": available_dates,
            "validation_current_date": current_date,
            "validation_previous_date": previous_date,
            "validation_payload": payload,
            "validation_movement_options": MOVEMENT_OPTIONS,
            "validation_movement_filter": movement_filter,
            "validation_search": search,
            "validation_page_size": page_size,
            "validation_page_size_options": PAGE_SIZE_OPTIONS,
            "validation_filtered_count": len(filtered_details),
            "validation_page_obj": page_obj,
            "validation_query_string": _validation_query_string(request),
        },
    )


def basel_validations_download_view(request: HttpRequest):
    selected_branch = _clean(request.GET.get("branch"))
    branch_names = get_request_branch_names(request)
    if selected_branch and selected_branch.casefold() not in {name.casefold() for name in branch_names}:
        messages.error(request, "Choose a branch within your assigned branch scope before downloading.")
        return redirect("scorecard:ifrs9_results_basel_validations")

    historical_queryset, _branch_names = _historical_scope_queryset(request, selected_branch)
    available_dates = _available_historical_dates(historical_queryset)
    current_date, previous_date = _resolve_comparison_dates(
        available_dates,
        request.GET.get("current_date"),
        request.GET.get("previous_date"),
    )
    if current_date is None or previous_date is None:
        messages.error(request, "At least two historical month-end dates are required for the NPL migration report.")
        return redirect("scorecard:ifrs9_results_basel_validations")

    movement_filter = _clean(request.GET.get("movement"))
    if movement_filter not in VALID_MOVEMENT_FILTERS:
        movement_filter = ""
    search = _clean(request.GET.get("search"))
    branch_scope_label = current_branch_display_name(request) or "No branch selected"
    export_cache_key = validation_export_cache_key(
        report_name="basel",
        queryset=historical_queryset,
        previous_date=previous_date,
        current_date=current_date,
        branch_names=branch_names,
        selected_branch=selected_branch,
        branch_scope_label=branch_scope_label,
        movement_filter=movement_filter,
        search=search,
    )
    cached_export = get_cached_validation_export(export_cache_key)
    cache_status = "HIT" if cached_export else "MISS"
    if cached_export:
        workbook_bytes, customer_row_count = cached_export
    else:
        payload = _get_validation_payload(
            historical_queryset,
            previous_date,
            current_date,
            branch_names,
            selected_branch,
        )
        detail_rows = _filter_details(payload["details"], search, movement_filter)
        workbook = _build_validation_workbook(
            payload,
            detail_rows,
            branch_scope_label,
            selected_branch,
            movement_filter,
            search,
        )
        output = BytesIO()
        workbook.save(output)
        workbook_bytes = output.getvalue()
        customer_row_count = len(detail_rows)
        cache_validation_export(export_cache_key, workbook_bytes, customer_row_count)

    log_ifrs9_results_audit(
        request.user,
        "download_basel_validations",
        details=(
            f"Downloaded NPL migration report. Previous month-end: {previous_date.isoformat()}; "
            f"Current month-end: {current_date.isoformat()}; Branch scope: "
            f"{selected_branch or current_branch_display_name(request) or 'No branch selected'}; "
            f"Customer rows: {customer_row_count}"
        ),
        object_id=f"{previous_date.isoformat()}:{current_date.isoformat()}",
        branch_name=selected_branch or current_branch_display_name(request) or "",
    )
    response = HttpResponse(
        workbook_bytes,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="npl_migration_report_{previous_date.isoformat()}_to_{current_date.isoformat()}.xlsx"'
    )
    response["X-Validation-Export-Cache"] = cache_status
    return response
