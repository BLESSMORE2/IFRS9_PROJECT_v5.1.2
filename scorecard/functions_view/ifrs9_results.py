from __future__ import annotations

from datetime import date
from io import BytesIO
from decimal import Decimal

import openpyxl

from django.apps import apps
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Count, Max, Q, Sum
from django.http import HttpRequest, HttpResponse
from django.shortcuts import redirect, render
from django.template.loader import render_to_string
from django.urls import reverse
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from scorecard.functions_view.audit import log_ifrs9_results_audit
from scorecard.functions_view.main_customer_lookup import (
    current_branch_display_name,
    get_request_branch_scope,
    is_all_branches_selected,
)


DEFAULT_RESULTS_EXTRACT_COLUMNS = (
    "fic_mis_date",
    "n_run_skey",
    "v_branch_name",
    "v_branch_code",
    "v_loan_id",
    "v_cust_name",
    "v_portfolio_name",
    "n_prod_segment",
    "v_ccy_code",
    "n_stage_descr",
    "n_principal_balance_rcy",
    "n_exposure_at_default_rcy",
    "n_12m_ecl_rcy",
    "n_lifetime_ecl_rcy",
    "n_reporting_ecl_rcy",
    "n_weighted_reporting_ecl_rcy",
)
RESULTS_EXTRACT_PAGE_SIZE = 25
ECL_SUMMARY_GROUP_BY_OPTIONS = (
    ("n_stage_descr", "Stage Description"),
    ("v_portfolio_name", "Portfolio Name"),
    ("n_risk_group", "Risk Group"),
    ("v_branch_name", "Branch Name"),
)
DEFAULT_ECL_SUMMARY_GROUP_BY_FIELD = "n_stage_descr"


def _load_ifrs9_results_models():
    if not apps.is_installed("IFRS9"):
        return None, None, None, (
            "The IFRS9 module is not installed in this environment, so IFRS9 Results "
            "cannot be opened right now."
        )

    try:
        reporting_model = apps.get_model("IFRS9", "fct_reporting_table")
    except LookupError:
        return None, None, None, (
            "The IFRS9 reporting table model is not available in this environment, so "
            "IFRS9 Results cannot be opened right now."
        )

    try:
        report_config_model = apps.get_model("IFRS9", "ReportColumnConfig")
    except LookupError:
        report_config_model = None

    try:
        reporting_currency_model = apps.get_model("IFRS9", "ReportingCurrency")
    except LookupError:
        reporting_currency_model = None

    return reporting_model, report_config_model, reporting_currency_model, None


def _build_branch_scope_filter(branch_scope) -> Q:
    branch_filter = Q(pk__in=[])
    for branch in branch_scope:
        branch_code = (getattr(branch, "branch_code", "") or "").strip()
        branch_name = (getattr(branch, "branch_name", "") or "").strip()
        branch_match = Q()
        if branch_code:
            branch_match |= Q(v_branch_code__iexact=branch_code)
        if branch_name:
            branch_match |= Q(v_branch_name__iexact=branch_name)
        if branch_match.children:
            branch_filter |= branch_match
    return branch_filter


def _get_results_scope_queryset(request: HttpRequest, reporting_model):
    branch_scope = get_request_branch_scope(request)
    branch_filter = _build_branch_scope_filter(branch_scope)
    if not branch_scope or not branch_filter.children:
        return reporting_model.objects.none(), branch_scope
    return reporting_model.objects.filter(branch_filter), branch_scope


def _parse_reporting_date(raw_value: str | None) -> date | None:
    cleaned_value = (raw_value or "").strip()
    if not cleaned_value:
        return None
    try:
        return date.fromisoformat(cleaned_value)
    except ValueError:
        return None


def _get_reporting_currency_code(reporting_currency_model) -> str:
    if reporting_currency_model is None:
        return ""
    try:
        row = reporting_currency_model.objects.select_related("currency_code").first()
        if row and getattr(row, "currency_code", None):
            return getattr(row.currency_code, "code", "") or ""
    except Exception:
        return ""
    return ""


def _sanitize_reporting_columns(reporting_model, columns) -> list[str]:
    valid_fields = {field.name for field in reporting_model._meta.concrete_fields}
    return [column for column in (columns or []) if column in valid_fields]


def _resolve_results_extract_columns(reporting_model, report_config_model) -> tuple[list[str], str | None]:
    default_columns = _sanitize_reporting_columns(reporting_model, DEFAULT_RESULTS_EXTRACT_COLUMNS)
    if not default_columns:
        default_columns = [field.name for field in reporting_model._meta.concrete_fields[:10]]

    if report_config_model is None:
        return default_columns, "Using default results extract columns because IFRS9 report column settings are unavailable."

    report_config = (
        report_config_model.objects.filter(report_name="default_report")
        .order_by("-created_at", "-id")
        .first()
    )
    if not report_config:
        return default_columns, "Using default results extract columns because no saved IFRS9 report column configuration was found."

    selected_columns = _sanitize_reporting_columns(reporting_model, getattr(report_config, "selected_columns", None))
    if selected_columns:
        return selected_columns, None
    return default_columns, "Using default results extract columns because the saved IFRS9 column configuration is invalid."


def _display_column_label(column_name: str) -> str:
    known_labels = {
        "fic_mis_date": "Reporting Date",
        "n_run_skey": "Run Key",
        "v_branch_name": "Branch Name",
        "v_branch_code": "Branch Code",
        "v_loan_id": "Loan ID",
        "v_cust_name": "Customer Name",
        "v_portfolio_name": "Portfolio Name",
        "n_prod_segment": "Product Segment",
        "v_ccy_code": "Currency",
        "n_stage_descr": "Stage Description",
        "n_principal_balance_rcy": "Total Exposure",
        "n_exposure_at_default_rcy": "Total EAD",
        "n_12m_ecl_rcy": "12 Month ECL",
        "n_lifetime_ecl_rcy": "Lifetime ECL",
        "n_reporting_ecl_rcy": "Reporting ECL",
        "n_weighted_reporting_ecl_rcy": "Weighted Reporting ECL",
    }
    if column_name in known_labels:
        return known_labels[column_name]
    cleaned = str(column_name).strip()
    for prefix in ("v_", "n_", "d_", "f_"):
        if cleaned.lower().startswith(prefix):
            cleaned = cleaned[2:]
            break
    return cleaned.replace("_", " ").title()


def _resolve_ecl_summary_group_by_field(raw_value: str | None) -> str:
    valid_fields = {field for field, _label in ECL_SUMMARY_GROUP_BY_OPTIONS}
    cleaned_value = (raw_value or "").strip()
    if cleaned_value in valid_fields:
        return cleaned_value
    return DEFAULT_ECL_SUMMARY_GROUP_BY_FIELD


def _ecl_summary_group_by_label_map() -> dict[str, str]:
    return {field: label for field, label in ECL_SUMMARY_GROUP_BY_OPTIONS}


def _get_available_reporting_dates(scope_queryset) -> list[date]:
    return list(
        scope_queryset.order_by("-fic_mis_date")
        .values_list("fic_mis_date", flat=True)
        .distinct()
    )


def _get_latest_run_key_for_date(scope_queryset, reporting_date: date | None):
    if reporting_date is None:
        return None
    return (
        scope_queryset.filter(fic_mis_date=reporting_date)
        .aggregate(latest_run_key=Max("n_run_skey"))
        .get("latest_run_key")
    )


def _coerce_decimal(value) -> Decimal:
    if value is None:
        return Decimal("0")
    if isinstance(value, Decimal):
        return value
    try:
        return Decimal(str(value))
    except Exception:
        return Decimal("0")


def _decimal_to_float(value) -> float:
    return float(_coerce_decimal(value))


def _decimal_ratio_percent(numerator, denominator) -> Decimal:
    denominator_value = _coerce_decimal(denominator)
    if denominator_value == 0:
        return Decimal("0")
    return (_coerce_decimal(numerator) / denominator_value) * Decimal("100")


def _build_results_common_context(
    request: HttpRequest,
    *,
    page_key: str,
    available_dates: list[date],
    selected_date: date | None,
    latest_run_key,
    branch_scope,
    module_available: bool,
    module_error: str | None = None,
):
    return {
        "results_page_key": page_key,
        "results_available_dates": available_dates,
        "results_selected_date": selected_date,
        "results_selected_date_value": selected_date.isoformat() if selected_date else "",
        "results_latest_run_key": latest_run_key,
        "results_branch_scope_label": current_branch_display_name(request) or "No branch selected",
        "results_branch_mode_is_all": is_all_branches_selected(request),
        "results_branch_scope_count": len(branch_scope),
        "results_module_available": module_available,
        "results_module_error": module_error,
    }


def _build_extract_export_workbook(
    *,
    selected_columns: list[str],
    selected_column_labels: dict[str, str],
    report_rows,
    reporting_date_value: str,
    latest_run_key,
    branch_scope_label: str,
):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Results Extract"
    worksheet.sheet_view.showGridLines = False

    title_fill = PatternFill(start_color="082b4b", end_color="082b4b", fill_type="solid")
    subtitle_fill = PatternFill(start_color="eaf3fb", end_color="eaf3fb", fill_type="solid")
    header_fill = PatternFill(start_color="0e4377", end_color="0e4377", fill_type="solid")
    accent_fill = PatternFill(start_color="00a99d", end_color="00a99d", fill_type="solid")
    thin_border = Border(bottom=Side(style="thin", color="d9e4ef"))
    title_font = Font(color="FFFFFF", bold=True, size=16)
    white_font = Font(color="FFFFFF", bold=True)
    small_muted_font = Font(color="5d7189", bold=True, size=9)
    header_font = Font(color="FFFFFF", bold=True, size=10)

    headers = [selected_column_labels.get(column, column) for column in selected_columns]
    total_columns = max(len(headers), 8)
    last_column_letter = get_column_letter(len(headers))

    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_columns)
    worksheet["A1"] = "Scorecard IFRS9 Results Extract"
    worksheet["A1"].fill = title_fill
    worksheet["A1"].font = title_font
    worksheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[1].height = 28

    worksheet.append(
        [
            "Reporting Date",
            reporting_date_value or "Not selected",
            "Latest Run Key",
            latest_run_key or "Not available",
            "Rows",
            len(report_rows),
            "Branch Scope",
            branch_scope_label,
        ]
    )
    worksheet.append(
        [
            "Window",
            "Full export",
            "Columns",
            len(selected_columns),
            "Source",
            "fct_reporting_table",
            "",
            "",
        ]
    )
    worksheet.append([])
    worksheet.append(headers)

    for row_index in (2, 3):
        for cell in worksheet[row_index]:
            cell.fill = subtitle_fill
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
            if cell.column % 2 == 1:
                cell.font = small_muted_font
            else:
                cell.font = Font(color="0f2236", bold=True, size=9)

    header_row = 5
    for cell in worksheet[header_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    worksheet.row_dimensions[header_row].height = 24

    width_tracker = {index: min(max(len(header), 12), 34) for index, header in enumerate(headers, start=1)}
    for row_number, report_row in enumerate(report_rows, start=header_row + 1):
        row_values = [report_row.get(column) for column in selected_columns]
        worksheet.append(row_values)
        if row_number <= header_row + 200:
            for column_index, value in enumerate(row_values, start=1):
                if value in (None, ""):
                    continue
                width_tracker[column_index] = min(max(width_tracker[column_index], len(str(value)) + 2), 42)

    worksheet.freeze_panes = "A6"
    if report_rows:
        worksheet.auto_filter.ref = f"A{header_row}:{last_column_letter}{worksheet.max_row}"
        table = Table(displayName="ScorecardResultsExtractTable", ref=f"A{header_row}:{last_column_letter}{worksheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

    for column_index in range(1, len(headers) + 1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width_tracker.get(column_index, 16)

    for column_index in range(1, total_columns + 1):
        worksheet.cell(row=1, column=column_index).fill = title_fill
        worksheet.cell(row=1, column=column_index).font = title_font if column_index == 1 else white_font
        worksheet.cell(row=4, column=column_index).fill = accent_fill
    worksheet.row_dimensions[4].height = 5

    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_margins.left = 0.25
    worksheet.page_margins.right = 0.25
    worksheet.page_margins.top = 0.5
    worksheet.page_margins.bottom = 0.5

    return workbook


def _build_ecl_summary_payload(scope_queryset, selected_date: date | None, latest_run_key, group_by_fields: tuple[str, ...]):
    grouped_data = []
    grouped_percentages = []
    grand_totals = {
        "n_principal_balance_rcy": 0.0,
        "n_exposure_at_default_rcy": 0.0,
        "n_collateral_considered": 0.0,
        "n_reporting_ecl_rcy": 0.0,
        "n_weighted_reporting_ecl_rcy": 0.0,
        "n_account_number": 0,
        "ecl_ead_ratio": 0.0,
        "ecl_principal_ratio": 0.0,
    }

    if selected_date is None or latest_run_key is None:
        return grouped_data, grouped_percentages, grand_totals

    filtered_queryset = scope_queryset.filter(fic_mis_date=selected_date, n_run_skey=latest_run_key)
    grouped_rows = list(
        filtered_queryset.values(*group_by_fields)
        .annotate(
            sum_n_principal_balance_rcy=Sum("n_principal_balance_rcy"),
            sum_n_exposure_at_default_rcy=Sum("n_exposure_at_default_rcy"),
            sum_n_collateral_considered=Sum("n_collateral_considered"),
            sum_n_reporting_ecl_rcy=Sum("n_reporting_ecl_rcy"),
            sum_n_weighted_reporting_ecl_rcy=Sum("n_weighted_reporting_ecl_rcy"),
            n_account_number=Count("v_loan_id", distinct=True),
        )
        .order_by(*group_by_fields)
    )

    for row in grouped_rows:
        normalized_row = {}
        for field in group_by_fields:
            normalized_row[field] = row.get(field) or "N/A"
        normalized_row["n_principal_balance_rcy"] = _decimal_to_float(row.get("sum_n_principal_balance_rcy"))
        normalized_row["n_exposure_at_default_rcy"] = _decimal_to_float(row.get("sum_n_exposure_at_default_rcy"))
        normalized_row["n_collateral_considered"] = _decimal_to_float(row.get("sum_n_collateral_considered"))
        normalized_row["n_reporting_ecl_rcy"] = _decimal_to_float(row.get("sum_n_reporting_ecl_rcy"))
        normalized_row["n_weighted_reporting_ecl_rcy"] = _decimal_to_float(row.get("sum_n_weighted_reporting_ecl_rcy"))
        normalized_row["n_account_number"] = row.get("n_account_number") or 0
        normalized_row["ecl_ead_ratio"] = _decimal_to_float(
            _decimal_ratio_percent(
                normalized_row["n_reporting_ecl_rcy"],
                normalized_row["n_exposure_at_default_rcy"],
            )
        )
        normalized_row["ecl_principal_ratio"] = _decimal_to_float(
            _decimal_ratio_percent(
                normalized_row["n_reporting_ecl_rcy"],
                normalized_row["n_principal_balance_rcy"],
            )
        )
        grouped_data.append(normalized_row)

    totals = filtered_queryset.aggregate(
        sum_n_principal_balance_rcy=Sum("n_principal_balance_rcy"),
        sum_n_exposure_at_default_rcy=Sum("n_exposure_at_default_rcy"),
        sum_n_collateral_considered=Sum("n_collateral_considered"),
        sum_n_reporting_ecl_rcy=Sum("n_reporting_ecl_rcy"),
        sum_n_weighted_reporting_ecl_rcy=Sum("n_weighted_reporting_ecl_rcy"),
        n_account_number=Count("v_loan_id", distinct=True),
    )
    grand_totals = {
        "n_principal_balance_rcy": _decimal_to_float(totals.get("sum_n_principal_balance_rcy")),
        "n_exposure_at_default_rcy": _decimal_to_float(totals.get("sum_n_exposure_at_default_rcy")),
        "n_collateral_considered": _decimal_to_float(totals.get("sum_n_collateral_considered")),
        "n_reporting_ecl_rcy": _decimal_to_float(totals.get("sum_n_reporting_ecl_rcy")),
        "n_weighted_reporting_ecl_rcy": _decimal_to_float(totals.get("sum_n_weighted_reporting_ecl_rcy")),
        "n_account_number": totals.get("n_account_number") or 0,
    }
    total_reporting_ecl = grand_totals["n_reporting_ecl_rcy"]
    grand_totals["ecl_ead_ratio"] = _decimal_to_float(
        _decimal_ratio_percent(total_reporting_ecl, grand_totals["n_exposure_at_default_rcy"])
    )
    grand_totals["ecl_principal_ratio"] = _decimal_to_float(
        _decimal_ratio_percent(total_reporting_ecl, grand_totals["n_principal_balance_rcy"])
    )

    for row in grouped_data:
        grouped_percentages.append(
            {field: row[field] for field in group_by_fields}
            | {
                "percent_reporting_ecl_rcy": _decimal_to_float(
                    _decimal_ratio_percent(row.get("n_reporting_ecl_rcy"), total_reporting_ecl)
                ),
                "percent_accounts": (
                    float(Decimal(row["n_account_number"]) / Decimal(grand_totals["n_account_number"]) * Decimal("100"))
                    if grand_totals["n_account_number"]
                    else 0.0
                ),
            }
        )

    return grouped_data, grouped_percentages, grand_totals


def _build_ecl_summary_export_workbook(
    *,
    group_by_fields: tuple[str, ...],
    group_by_labels: dict[str, str],
    grouped_data,
    grouped_percentages,
    grand_totals,
    reporting_currency_code: str,
    reporting_date_value: str,
    latest_run_key,
    branch_scope_label: str,
):
    workbook = openpyxl.Workbook()
    ws1 = workbook.active
    ws1.title = "ECL Summary"
    ws2 = workbook.create_sheet(title="ECL Percentages")

    title_fill = PatternFill(start_color="082b4b", end_color="082b4b", fill_type="solid")
    subtitle_fill = PatternFill(start_color="eaf3fb", end_color="eaf3fb", fill_type="solid")
    header_fill = PatternFill(start_color="0e4377", end_color="0e4377", fill_type="solid")
    thin_border = Border(bottom=Side(style="thin", color="d9e4ef"))
    title_font = Font(color="FFFFFF", bold=True, size=15)
    meta_label_font = Font(color="5d7189", bold=True, size=9)
    meta_value_font = Font(color="0f2236", bold=True, size=9)
    header_font = Font(color="FFFFFF", bold=True, size=10)
    currency_suffix = f" ({reporting_currency_code})" if reporting_currency_code else ""

    group_headers = [group_by_labels.get(field, field.replace("_", " ").title()) for field in group_by_fields]
    absolute_headers = group_headers + [
        f"Total Exposure{currency_suffix}",
        f"Total EAD{currency_suffix}",
        f"Collateral Amount{currency_suffix}",
        f"Reporting ECL{currency_suffix}",
        f"Weighted Reporting ECL{currency_suffix}",
        "ECL / EAD Ratio",
        "ECL / Exposure Ratio",
        "Contracts",
    ]
    percentage_headers = group_headers + [
        "% of Reporting ECL",
        "% of Contracts",
    ]

    for worksheet, title, headers in (
        (ws1, "Scorecard IFRS9 ECL Summary", absolute_headers),
        (ws2, "Scorecard IFRS9 ECL Summary Percentages", percentage_headers),
    ):
        total_columns = max(len(headers), 8)
        worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_columns)
        worksheet["A1"] = title
        worksheet["A1"].fill = title_fill
        worksheet["A1"].font = title_font
        worksheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
        worksheet.row_dimensions[1].height = 26

        worksheet.append(
            [
                "Reporting Date",
                reporting_date_value or "Not selected",
                "Latest Run Key",
                latest_run_key or "Not available",
                "Branch Scope",
                branch_scope_label,
            ]
        )
        worksheet.append([])
        worksheet.append(headers)
        for cell in worksheet[2]:
            cell.fill = subtitle_fill
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
            if cell.column % 2 == 1:
                cell.font = meta_label_font
            else:
                cell.font = meta_value_font
        for cell in worksheet[4]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        worksheet.freeze_panes = "A5"

    for row in grouped_data:
        ws1.append(
            [row.get(field, "") for field in group_by_fields]
            + [
                row.get("n_principal_balance_rcy", 0),
                row.get("n_exposure_at_default_rcy", 0),
                row.get("n_collateral_considered", 0),
                row.get("n_reporting_ecl_rcy", 0),
                row.get("n_weighted_reporting_ecl_rcy", 0),
                row.get("ecl_ead_ratio", 0),
                row.get("ecl_principal_ratio", 0),
                row.get("n_account_number", 0),
            ]
        )
    ws1.append(
        ["Grand Total"] + ([""] * (len(group_by_fields) - 1)) + [
            grand_totals.get("n_principal_balance_rcy", 0),
            grand_totals.get("n_exposure_at_default_rcy", 0),
            grand_totals.get("n_collateral_considered", 0),
            grand_totals.get("n_reporting_ecl_rcy", 0),
            grand_totals.get("n_weighted_reporting_ecl_rcy", 0),
            grand_totals.get("ecl_ead_ratio", 0),
            grand_totals.get("ecl_principal_ratio", 0),
            grand_totals.get("n_account_number", 0),
        ]
    )

    for row in grouped_percentages:
        ws2.append(
            [row.get(field, "") for field in group_by_fields]
            + [
                row.get("percent_reporting_ecl_rcy", 0),
                row.get("percent_accounts", 0),
            ]
        )
    ws2.append(
        ["Grand Total"] + ([""] * (len(group_by_fields) - 1)) + [
            100 if grand_totals.get("n_reporting_ecl_rcy") else 0,
            100 if grand_totals.get("n_account_number") else 0,
        ]
    )

    for worksheet in (ws1, ws2):
        for row in worksheet.iter_rows(min_row=5, max_row=worksheet.max_row):
            for cell in row:
                cell.border = thin_border
        width_tracker = {}
        for row in worksheet.iter_rows(min_row=4, max_row=min(worksheet.max_row, 204)):
            for cell in row:
                width_tracker[cell.column] = min(max(width_tracker.get(cell.column, 12), len(str(cell.value or "")) + 2), 32)
        for column_index, width in width_tracker.items():
            worksheet.column_dimensions[get_column_letter(column_index)].width = width
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1

    return workbook


def _build_results_query_params(*, selected_date: date | None, group_by_field: str | None = None) -> str:
    params = []
    if selected_date is not None:
        params.append(f"reporting_date={selected_date.isoformat()}")
    if group_by_field:
        params.append(f"group_by_field={group_by_field}")
    return "&".join(params)


def _build_ecl_summary_pdf_rows(grouped_data, grouped_percentages, group_by_fields: tuple[str, ...]):
    summary_rows = []
    for row in grouped_data:
        summary_rows.append(
            {
                "group_values": [row.get(field, "N/A") for field in group_by_fields],
                "n_principal_balance_rcy": row.get("n_principal_balance_rcy", 0),
                "n_exposure_at_default_rcy": row.get("n_exposure_at_default_rcy", 0),
                "n_collateral_considered": row.get("n_collateral_considered", 0),
                "n_reporting_ecl_rcy": row.get("n_reporting_ecl_rcy", 0),
                "n_weighted_reporting_ecl_rcy": row.get("n_weighted_reporting_ecl_rcy", 0),
                "ecl_ead_ratio": row.get("ecl_ead_ratio", 0),
                "ecl_principal_ratio": row.get("ecl_principal_ratio", 0),
                "n_account_number": row.get("n_account_number", 0),
            }
        )

    percentage_rows = []
    for row in grouped_percentages:
        percentage_rows.append(
            {
                "group_values": [row.get(field, "N/A") for field in group_by_fields],
                "percent_reporting_ecl_rcy": row.get("percent_reporting_ecl_rcy", 0),
                "percent_accounts": row.get("percent_accounts", 0),
            }
        )

    return summary_rows, percentage_rows


def ifrs9_results_home_view(request: HttpRequest):
    reporting_model, _report_config_model, _reporting_currency_model, module_error = _load_ifrs9_results_models()
    available_dates: list[date] = []
    branch_scope = get_request_branch_scope(request)

    if reporting_model is not None:
        scope_queryset, branch_scope = _get_results_scope_queryset(request, reporting_model)
        available_dates = _get_available_reporting_dates(scope_queryset)

    context = _build_results_common_context(
        request,
        page_key="home",
        available_dates=available_dates[:8],
        selected_date=None,
        latest_run_key=None,
        branch_scope=branch_scope,
        module_available=reporting_model is not None,
        module_error=module_error,
    )
    context.update(
        {
            "results_workspace_title": "IFRS9 Results",
            "results_workspace_copy": (
                "Open branch-scoped IFRS9 reporting results using reporting date only. "
                "The latest run key for the chosen date is selected automatically."
            ),
        }
    )
    return render(request, "credit_scoreshifts/ifrs9_results_home.html", context)


def ifrs9_results_extract_view(request: HttpRequest):
    reporting_model, report_config_model, reporting_currency_model, module_error = _load_ifrs9_results_models()
    branch_scope = get_request_branch_scope(request)
    if reporting_model is None:
        context = _build_results_common_context(
            request,
            page_key="extract",
            available_dates=[],
            selected_date=None,
            latest_run_key=None,
            branch_scope=branch_scope,
            module_available=False,
            module_error=module_error,
        )
        return render(request, "credit_scoreshifts/ifrs9_results_extract.html", context)

    scope_queryset, branch_scope = _get_results_scope_queryset(request, reporting_model)
    available_dates = _get_available_reporting_dates(scope_queryset)
    raw_reporting_date = request.GET.get("reporting_date")
    selected_date = _parse_reporting_date(raw_reporting_date)
    if raw_reporting_date and selected_date is None:
        messages.error(request, "Choose a valid Reporting Date before loading IFRS9 results.")

    latest_run_key = _get_latest_run_key_for_date(scope_queryset, selected_date)
    if selected_date and latest_run_key is None:
        messages.error(
            request,
            "No IFRS9 results were found for the selected Reporting Date in the current branch scope.",
        )

    selected_columns, config_warning = _resolve_results_extract_columns(reporting_model, report_config_model)
    if config_warning:
        messages.warning(request, config_warning)
    selected_column_labels = {column: _display_column_label(column) for column in selected_columns}

    report_rows = []
    paginator_page = None
    total_rows = 0
    if selected_date is not None and latest_run_key is not None:
        result_queryset = (
            scope_queryset.filter(fic_mis_date=selected_date, n_run_skey=latest_run_key)
            .values(*selected_columns)
        )
        total_rows = result_queryset.count()
        paginator = Paginator(result_queryset, RESULTS_EXTRACT_PAGE_SIZE)
        paginator_page = paginator.get_page(request.GET.get("page") or 1)
        report_rows = list(paginator_page.object_list)

    context = _build_results_common_context(
        request,
        page_key="extract",
        available_dates=available_dates,
        selected_date=selected_date,
        latest_run_key=latest_run_key,
        branch_scope=branch_scope,
        module_available=True,
        module_error=None,
    )
    context.update(
        {
            "reporting_currency_code": _get_reporting_currency_code(reporting_currency_model),
            "results_extract_columns": selected_columns,
            "results_extract_column_labels": selected_column_labels,
            "results_extract_page": paginator_page,
            "results_extract_rows": report_rows,
            "results_extract_total_rows": total_rows,
        }
    )
    return render(request, "credit_scoreshifts/ifrs9_results_extract.html", context)


def ifrs9_results_extract_download_view(request: HttpRequest):
    reporting_model, report_config_model, _reporting_currency_model, module_error = _load_ifrs9_results_models()
    if reporting_model is None:
        messages.error(request, module_error or "IFRS9 Results are not available in this environment.")
        return redirect("scorecard:ifrs9_results_home")

    scope_queryset, branch_scope = _get_results_scope_queryset(request, reporting_model)
    selected_date = _parse_reporting_date(request.GET.get("reporting_date"))
    if selected_date is None:
        messages.error(request, "Choose a Reporting Date before downloading the IFRS9 results extract.")
        return redirect("scorecard:ifrs9_results_extract")

    latest_run_key = _get_latest_run_key_for_date(scope_queryset, selected_date)
    if latest_run_key is None:
        messages.error(request, "No IFRS9 results were found for the selected Reporting Date.")
        return redirect(
            f"{reverse('scorecard:ifrs9_results_extract')}?reporting_date={selected_date.isoformat()}"
        )

    selected_columns, _config_warning = _resolve_results_extract_columns(reporting_model, report_config_model)
    selected_column_labels = {column: _display_column_label(column) for column in selected_columns}
    report_rows = list(
        scope_queryset.filter(fic_mis_date=selected_date, n_run_skey=latest_run_key).values(*selected_columns)
    )
    workbook = _build_extract_export_workbook(
        selected_columns=selected_columns,
        selected_column_labels=selected_column_labels,
        report_rows=report_rows,
        reporting_date_value=selected_date.isoformat(),
        latest_run_key=latest_run_key,
        branch_scope_label=current_branch_display_name(request) or "No branch selected",
    )
    output_stream = BytesIO()
    workbook.save(output_stream)
    output_stream.seek(0)

    response = HttpResponse(
        output_stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="scorecard_ifrs9_results_{selected_date.isoformat()}_run_{latest_run_key}.xlsx"'
    )
    log_ifrs9_results_audit(
        request.user,
        "download_results_extract",
        details=(
            f"Downloaded IFRS9 Results Extract. Reporting Date: {selected_date.isoformat()}; "
            f"Run Key: {latest_run_key}; Branch scope: {current_branch_display_name(request) or 'No branch selected'}; "
            f"Rows: {len(report_rows)}"
        ),
        object_id=f"{selected_date.isoformat()}:{latest_run_key}",
        branch_name=current_branch_display_name(request) or "",
    )
    return response


def ifrs9_results_ecl_summary_view(request: HttpRequest):
    reporting_model, _report_config_model, reporting_currency_model, module_error = _load_ifrs9_results_models()
    branch_scope = get_request_branch_scope(request)
    if reporting_model is None:
        context = _build_results_common_context(
            request,
            page_key="summary",
            available_dates=[],
            selected_date=None,
            latest_run_key=None,
            branch_scope=branch_scope,
            module_available=False,
            module_error=module_error,
        )
        return render(request, "credit_scoreshifts/ifrs9_results_ecl_summary.html", context)

    scope_queryset, branch_scope = _get_results_scope_queryset(request, reporting_model)
    available_dates = _get_available_reporting_dates(scope_queryset)
    raw_reporting_date = request.GET.get("reporting_date")
    selected_date = _parse_reporting_date(raw_reporting_date)
    selected_group_by_field = _resolve_ecl_summary_group_by_field(request.GET.get("group_by_field"))
    group_by_fields = (selected_group_by_field,)
    if raw_reporting_date and selected_date is None:
        messages.error(request, "Choose a valid Reporting Date before loading the ECL Summary Report.")

    latest_run_key = _get_latest_run_key_for_date(scope_queryset, selected_date)
    if selected_date and latest_run_key is None:
        messages.error(
            request,
            "No IFRS9 summary rows were found for the selected Reporting Date in the current branch scope.",
        )

    grouped_data, grouped_percentages, grand_totals = _build_ecl_summary_payload(
        scope_queryset,
        selected_date,
        latest_run_key,
        group_by_fields,
    )

    context = _build_results_common_context(
        request,
        page_key="summary",
        available_dates=available_dates,
        selected_date=selected_date,
        latest_run_key=latest_run_key,
        branch_scope=branch_scope,
        module_available=True,
        module_error=None,
    )
    context.update(
        {
            "reporting_currency_code": _get_reporting_currency_code(reporting_currency_model),
            "ecl_summary_group_by_fields": group_by_fields,
            "ecl_summary_group_by_field": selected_group_by_field,
            "ecl_summary_group_by_options": ECL_SUMMARY_GROUP_BY_OPTIONS,
            "ecl_summary_group_by_labels": _ecl_summary_group_by_label_map(),
            "ecl_summary_rows": grouped_data,
            "ecl_summary_percentages": grouped_percentages,
            "ecl_summary_grand_totals": grand_totals,
        }
    )
    return render(request, "credit_scoreshifts/ifrs9_results_ecl_summary.html", context)


def ifrs9_results_ecl_summary_download_excel_view(request: HttpRequest):
    reporting_model, _report_config_model, reporting_currency_model, module_error = _load_ifrs9_results_models()
    if reporting_model is None:
        messages.error(request, module_error or "IFRS9 Results are not available in this environment.")
        return redirect("scorecard:ifrs9_results_home")

    scope_queryset, _branch_scope = _get_results_scope_queryset(request, reporting_model)
    selected_date = _parse_reporting_date(request.GET.get("reporting_date"))
    selected_group_by_field = _resolve_ecl_summary_group_by_field(request.GET.get("group_by_field"))
    group_by_fields = (selected_group_by_field,)
    if selected_date is None:
        messages.error(request, "Choose a Reporting Date before downloading the ECL Summary Report.")
        return redirect("scorecard:ifrs9_results_ecl_summary")

    latest_run_key = _get_latest_run_key_for_date(scope_queryset, selected_date)
    if latest_run_key is None:
        messages.error(request, "No IFRS9 summary rows were found for the selected Reporting Date.")
        query_string = _build_results_query_params(
            selected_date=selected_date,
            group_by_field=selected_group_by_field,
        )
        return redirect(f"{reverse('scorecard:ifrs9_results_ecl_summary')}?{query_string}")

    grouped_data, grouped_percentages, grand_totals = _build_ecl_summary_payload(
        scope_queryset,
        selected_date,
        latest_run_key,
        group_by_fields,
    )
    workbook = _build_ecl_summary_export_workbook(
        group_by_fields=group_by_fields,
        group_by_labels=_ecl_summary_group_by_label_map(),
        grouped_data=grouped_data,
        grouped_percentages=grouped_percentages,
        grand_totals=grand_totals,
        reporting_currency_code=_get_reporting_currency_code(reporting_currency_model),
        reporting_date_value=selected_date.isoformat(),
        latest_run_key=latest_run_key,
        branch_scope_label=current_branch_display_name(request) or "No branch selected",
    )
    output_stream = BytesIO()
    workbook.save(output_stream)
    output_stream.seek(0)

    response = HttpResponse(
        output_stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        "attachment; filename="
        f"\"scorecard_ifrs9_ecl_summary_{selected_date.isoformat()}_run_{latest_run_key}.xlsx\""
    )
    log_ifrs9_results_audit(
        request.user,
        "download_ecl_summary_excel",
        details=(
            f"Downloaded IFRS9 ECL Summary Excel. Reporting Date: {selected_date.isoformat()}; "
            f"Run Key: {latest_run_key}; Group by: {_ecl_summary_group_by_label_map().get(selected_group_by_field, selected_group_by_field)}; "
            f"Branch scope: {current_branch_display_name(request) or 'No branch selected'}; Rows: {len(grouped_data)}"
        ),
        object_id=f"{selected_date.isoformat()}:{latest_run_key}",
        branch_name=current_branch_display_name(request) or "",
    )
    return response


def ifrs9_results_ecl_summary_download_pdf_view(request: HttpRequest):
    reporting_model, _report_config_model, reporting_currency_model, module_error = _load_ifrs9_results_models()
    if reporting_model is None:
        messages.error(request, module_error or "IFRS9 Results are not available in this environment.")
        return redirect("scorecard:ifrs9_results_home")

    try:
        from xhtml2pdf import pisa
    except Exception:
        messages.error(request, "PDF export is not available because the PDF engine is not installed.")
        query_string = _build_results_query_params(
            selected_date=_parse_reporting_date(request.GET.get("reporting_date")),
            group_by_field=_resolve_ecl_summary_group_by_field(request.GET.get("group_by_field")),
        )
        redirect_url = reverse("scorecard:ifrs9_results_ecl_summary")
        if query_string:
            redirect_url = f"{redirect_url}?{query_string}"
        return redirect(redirect_url)

    scope_queryset, _branch_scope = _get_results_scope_queryset(request, reporting_model)
    selected_date = _parse_reporting_date(request.GET.get("reporting_date"))
    selected_group_by_field = _resolve_ecl_summary_group_by_field(request.GET.get("group_by_field"))
    group_by_fields = (selected_group_by_field,)
    if selected_date is None:
        messages.error(request, "Choose a Reporting Date before downloading the ECL Summary Report.")
        return redirect("scorecard:ifrs9_results_ecl_summary")

    latest_run_key = _get_latest_run_key_for_date(scope_queryset, selected_date)
    if latest_run_key is None:
        messages.error(request, "No IFRS9 summary rows were found for the selected Reporting Date.")
        query_string = _build_results_query_params(
            selected_date=selected_date,
            group_by_field=selected_group_by_field,
        )
        return redirect(f"{reverse('scorecard:ifrs9_results_ecl_summary')}?{query_string}")

    grouped_data, grouped_percentages, grand_totals = _build_ecl_summary_payload(
        scope_queryset,
        selected_date,
        latest_run_key,
        group_by_fields,
    )
    summary_rows, percentage_rows = _build_ecl_summary_pdf_rows(
        grouped_data,
        grouped_percentages,
        group_by_fields,
    )
    group_by_labels = _ecl_summary_group_by_label_map()
    pdf_html = render_to_string(
        "credit_scoreshifts/ifrs9_results_ecl_summary_pdf.html",
        {
            "reporting_currency_code": _get_reporting_currency_code(reporting_currency_model),
            "results_selected_date_value": selected_date.isoformat(),
            "results_latest_run_key": latest_run_key,
            "results_branch_scope_label": current_branch_display_name(request) or "No branch selected",
            "ecl_summary_group_by_fields": group_by_fields,
            "ecl_summary_group_by_headers": [group_by_labels.get(field, field) for field in group_by_fields],
            "ecl_summary_grouped_by_label": group_by_labels.get(selected_group_by_field, selected_group_by_field),
            "ecl_summary_rows": summary_rows,
            "ecl_summary_percentages": percentage_rows,
            "ecl_summary_grand_totals": grand_totals,
        },
    )

    pdf_bytes = BytesIO()
    pisa_status = pisa.CreatePDF(BytesIO(pdf_html.encode("utf-8")), dest=pdf_bytes, encoding="utf-8")
    if pisa_status.err:
        messages.error(request, "Unable to generate the IFRS9 ECL Summary PDF right now.")
        query_string = _build_results_query_params(
            selected_date=selected_date,
            group_by_field=selected_group_by_field,
        )
        return redirect(f"{reverse('scorecard:ifrs9_results_ecl_summary')}?{query_string}")

    pdf_bytes.seek(0)
    response = HttpResponse(pdf_bytes.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = (
        "attachment; filename="
        f"\"scorecard_ifrs9_ecl_summary_{selected_date.isoformat()}_run_{latest_run_key}.pdf\""
    )
    log_ifrs9_results_audit(
        request.user,
        "download_ecl_summary_pdf",
        details=(
            f"Downloaded IFRS9 ECL Summary PDF. Reporting Date: {selected_date.isoformat()}; "
            f"Run Key: {latest_run_key}; Group by: {group_by_labels.get(selected_group_by_field, selected_group_by_field)}; "
            f"Branch scope: {current_branch_display_name(request) or 'No branch selected'}; Rows: {len(grouped_data)}"
        ),
        object_id=f"{selected_date.isoformat()}:{latest_run_key}",
        branch_name=current_branch_display_name(request) or "",
    )
    return response
