from __future__ import annotations

from collections import Counter
import csv
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from functools import lru_cache
from pathlib import Path
import pickle
import re
import tempfile
import threading
import uuid

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.db import IntegrityError, close_old_connections, transaction
from django.http import FileResponse, Http404, HttpRequest, HttpResponse
from django.shortcuts import redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_http_methods

from scorecard.functions_view.audit import log_historical_score_audit, log_upload_audit
from scorecard.models import BankBranch, CreditEvaluation, HistoricalScore, IFRS9Evaluation

try:
    from openpyxl import Workbook, load_workbook
except ImportError:  # pragma: no cover
    Workbook = None
    load_workbook = None

try:
    from pyxlsb import open_workbook as open_xlsb_workbook
except ImportError:  # pragma: no cover
    open_xlsb_workbook = None

try:
    import xlsxwriter
except ImportError:  # pragma: no cover
    xlsxwriter = None


SUPPORTED_EXTENSIONS = {".xlsx", ".xlsb", ".csv"}
PREVIEW_ROW_LIMIT = 10
PREVIEW_SCREEN_ROW_LIMIT = 500
BULK_CREATE_BATCH_SIZE = 1000
EVALUATION_IMPORT_BATCH_SIZE = 100
EXCEL_SERIAL_DATE_START = date(1899, 12, 30)
QUERY_CHUNK_SIZE = 1000
IMPORT_SESSION_PREFIX = "scorecard_external_upload"
IMPORT_RESULT_SESSION_PREFIX = "scorecard_external_upload_result"
PREVIEW_CACHE_MODE_COMPACT = "compact_preview_v1"
_preview_scan_lock = threading.BoundedSemaphore(value=1)
_upload_job_threads: dict[str, threading.Thread] = {}
_upload_job_threads_lock = threading.Lock()


@dataclass(frozen=True)
class ImportField:
    key: str
    label: str
    required: bool
    help_text: str
    aliases: tuple[str, ...]


IMPORT_CONFIGS = {
    "basel": {
        "label": "Basel credit evaluation import",
        "short_label": "Basel",
        "table_name": "SCORECARD_CREDIT_EVALUATION",
        "model": CreditEvaluation,
        "supports_grade": True,
        "fields": [
            ImportField(
                key="branch_name",
                label="Branch",
                required=True,
                help_text="Branch name or branch code exactly as it appears in the file. The importer only uppercases it.",
                aliases=("branch", "branchname", "branchcode", "branch code", "branch name"),
            ),
            ImportField(
                key="customer_id",
                label="Customer Number",
                required=True,
                help_text="Customer identifier. Leading zeroes will be removed for numeric customer numbers.",
                aliases=("customerid", "customer id", "customer number", "customer no", "cusno", "cus no", "id"),
            ),
            ImportField(
                key="customer_name",
                label="Customer Name",
                required=True,
                help_text="Customer name as shown in the sheet.",
                aliases=("customername", "customer name", "name", "clientname", "client name"),
            ),
            ImportField(
                key="total_weighted_percent",
                label="Weighted Percent / Score",
                required=True,
                help_text="Main Basel or IFRS score to load into total_weighted_percent. Use the 0 to 100 scale, for example 76 instead of 0.76.",
                aliases=(
                    "weightedpercent",
                    "weighted score",
                    "weighted percent",
                    "ifrs9score",
                    "ifrs 9 score",
                    "score",
                    "totalscore",
                ),
            ),
            ImportField(
                key="total_raw_score",
                label="Raw Score",
                required=False,
                help_text="Optional Basel raw score.",
                aliases=("rawscore", "raw score", "baseliiscore", "basel ii score", "baselscore", "basel score"),
            ),
            ImportField(
                key="final_grade",
                label="Final Grade",
                required=True,
                help_text="Required Basel grade.",
                aliases=("grade", "finalgrade", "final grade", "baseliigrade", "basel ii grade"),
            ),
        ],
    },
    "ifrs9": {
        "label": "IFRS9 evaluation import",
        "short_label": "IFRS9",
        "table_name": "SCORECARD_IFRS9_EVALUATION",
        "model": IFRS9Evaluation,
        "supports_grade": False,
        "fields": [
            ImportField(
                key="branch_name",
                label="Branch",
                required=True,
                help_text="Branch name or branch code exactly as it appears in the file. The importer only uppercases it.",
                aliases=("branch", "branchname", "branchcode", "branch code", "branch name"),
            ),
            ImportField(
                key="customer_id",
                label="Customer Number",
                required=True,
                help_text="Customer identifier. Leading zeroes will be removed for numeric customer numbers.",
                aliases=("customerid", "customer id", "customer number", "customer no", "cusno", "cus no", "id"),
            ),
            ImportField(
                key="customer_name",
                label="Customer Name",
                required=True,
                help_text="Customer name as shown in the sheet.",
                aliases=("customername", "customer name", "name", "clientname", "client name"),
            ),
            ImportField(
                key="total_weighted_percent",
                label="Weighted Percent / Score",
                required=True,
                help_text="Main IFRS9 score to load into total_weighted_percent. Use the 0 to 100 scale, for example 76 instead of 0.76.",
                aliases=(
                    "weightedpercent",
                    "weighted score",
                    "weighted percent",
                    "ifrs9score",
                    "ifrs 9 score",
                    "score",
                    "totalscore",
                ),
            ),
            ImportField(
                key="total_raw_score",
                label="Raw Score",
                required=False,
                help_text="Optional raw score if the source file carries one.",
                aliases=("rawscore", "raw score", "score raw"),
            ),
        ],
    },
    "historical_scores": {
        "label": "Historical score snapshot import",
        "short_label": "History",
        "table_name": "SCORECARD_HISTORICAL_SCORES",
        "model": HistoricalScore,
        "supports_grade": True,
        "is_historical_scores": True,
        "fields": [
            ImportField(
                key="reporting_date",
                label="Reporting Date",
                required=True,
                help_text="Historical reporting date for the snapshot, for example 2025-12-31.",
                aliases=("reportingdate", "reporting date", "ficmisdate", "fic mis date", "fic_mis_date", "date"),
            ),
            ImportField(
                key="branch_name",
                label="Branch",
                required=True,
                help_text="Branch name or branch code exactly as it appears in the file. The importer only uppercases it.",
                aliases=("branch", "branchname", "branchcode", "branch code", "branch name"),
            ),
            ImportField(
                key="customer_id",
                label="Customer Number",
                required=True,
                help_text="Customer identifier. Leading zeroes will be removed for numeric customer numbers.",
                aliases=("customerid", "customer id", "customer number", "customer no", "cusno", "cus no", "id", "customer code"),
            ),
            ImportField(
                key="customer_name",
                label="Customer Name",
                required=True,
                help_text="Customer name as shown in the sheet.",
                aliases=("customername", "customer name", "name", "clientname", "client name"),
            ),
            ImportField(
                key="basel_ii_score",
                label="Basel II Score",
                required=False,
                help_text="Optional historical Basel II weighted score on the 0 to 100 scale.",
                aliases=("baseliiscore", "basel ii score", "basel score", "baselscore", "weighted score", "weighted percent"),
            ),
            ImportField(
                key="basel_ii_grade",
                label="Basel II Grade",
                required=False,
                help_text="Optional Basel II grade captured for the reporting date.",
                aliases=("baseliigrade", "basel ii grade", "basel grade", "grade", "final grade", "finalgrade"),
            ),
            ImportField(
                key="basel_override_grade",
                label="Basel Override Grade",
                required=False,
                help_text="Optional manual Basel override grade captured for the reporting date.",
                aliases=(
                    "baseloverridegrade",
                    "basel override grade",
                    "override grade",
                    "manual grade",
                    "manual override grade",
                ),
            ),
            ImportField(
                key="ifrs_9_score",
                label="IFRS9 Score",
                required=False,
                help_text="Optional historical IFRS9 score on the 0 to 100 scale.",
                aliases=("ifrs9score", "ifrs 9 score", "ifrs score", "ifrsscore"),
            ),
        ],
    },
}


def _get_import_config(upload_kind: str) -> dict:
    config = IMPORT_CONFIGS.get(upload_kind)
    if not config:
        raise Http404("Unknown upload type.")
    return config


def _session_key(upload_kind: str) -> str:
    return f"{IMPORT_SESSION_PREFIX}_{upload_kind}"


def _result_session_key(upload_kind: str) -> str:
    return f"{IMPORT_RESULT_SESSION_PREFIX}_{upload_kind}"


def _get_state(request: HttpRequest, upload_kind: str) -> dict:
    return request.session.get(_session_key(upload_kind), {})


def _get_result_state(request: HttpRequest, upload_kind: str) -> dict:
    return request.session.get(_result_session_key(upload_kind), {})


def _store_state(request: HttpRequest, upload_kind: str, state: dict) -> None:
    request.session[_session_key(upload_kind)] = state
    request.session.modified = True


def _store_result_state(request: HttpRequest, upload_kind: str, state: dict) -> None:
    request.session[_result_session_key(upload_kind)] = state
    request.session.modified = True


def _delete_temp_file(path_value: str | None) -> None:
    if not path_value:
        return
    try:
        Path(path_value).unlink(missing_ok=True)
    except OSError:
        pass


def _clear_state(request: HttpRequest, upload_kind: str) -> None:
    state = _get_state(request, upload_kind)
    _delete_temp_file(state.get("file_path"))
    _delete_temp_file(state.get("preview_cache_path"))
    _delete_temp_file(state.get("preview_export_path"))
    request.session.pop(_session_key(upload_kind), None)
    request.session.modified = True


def _clear_result_state(request: HttpRequest, upload_kind: str) -> None:
    request.session.pop(_result_session_key(upload_kind), None)
    request.session.modified = True


def _save_uploaded_file(uploaded_file) -> str:
    suffix = Path(uploaded_file.name).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError("Only .xlsx, .xlsb, and .csv files are supported.")

    temp_dir = Path(tempfile.gettempdir()) / "scorecard_external_imports"
    temp_dir.mkdir(parents=True, exist_ok=True)
    target_path = temp_dir / f"{timezone.now().strftime('%Y%m%d%H%M%S')}_{uuid.uuid4().hex}{suffix}"

    with target_path.open("wb") as output_file:
        for chunk in uploaded_file.chunks():
            output_file.write(chunk)

    return str(target_path)


def _save_preview_cache(analysis: dict) -> str:
    temp_dir = Path(tempfile.gettempdir()) / "scorecard_external_imports"
    temp_dir.mkdir(parents=True, exist_ok=True)
    target_path = temp_dir / f"preview_{uuid.uuid4().hex}.pickle"
    with target_path.open("wb") as cache_file:
        pickle.dump(analysis, cache_file, protocol=pickle.HIGHEST_PROTOCOL)
    return str(target_path)


def _load_preview_cache(path_value: str | None) -> dict | None:
    if not path_value:
        return None
    cache_path = Path(path_value)
    if not cache_path.exists():
        return None
    try:
        with cache_path.open("rb") as cache_file:
            return pickle.load(cache_file)
    except (OSError, pickle.PickleError, EOFError):
        _delete_temp_file(path_value)
        return None


def _save_upload_job_state(job_path: str, state: dict) -> None:
    """Persist background-import progress atomically for IIS worker hand-offs."""
    target = Path(job_path)
    temporary = target.with_suffix(".tmp")
    with temporary.open("wb") as job_file:
        pickle.dump(state, job_file, protocol=pickle.HIGHEST_PROTOCOL)
    temporary.replace(target)


def _load_upload_job_state(job_path: str | None) -> dict | None:
    if not job_path:
        return None
    target = Path(job_path)
    if not target.exists():
        return None
    try:
        with target.open("rb") as job_file:
            return pickle.load(job_file)
    except (OSError, pickle.PickleError, EOFError):
        return None


def _new_upload_job_path() -> str:
    temp_dir = Path(tempfile.gettempdir()) / "scorecard_external_imports"
    temp_dir.mkdir(parents=True, exist_ok=True)
    return str(temp_dir / f"job_{uuid.uuid4().hex}.pickle")


def _list_sheet_names(file_path: str) -> list[str]:
    suffix = Path(file_path).suffix.lower()
    if suffix == ".csv":
        return ["CSV file"]
    if suffix == ".xlsx":
        if load_workbook is None:
            raise ValueError("openpyxl is not installed for .xlsx imports.")
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()
    if suffix == ".xlsb":
        if open_xlsb_workbook is None:
            raise ValueError("pyxlsb is not installed for .xlsb imports.")
        with open_xlsb_workbook(file_path) as workbook:
            return list(workbook.sheets)
    raise ValueError("Unsupported file type.")


def _iter_sheet_rows(file_path: str, sheet_name: str):
    suffix = Path(file_path).suffix.lower()
    if suffix == ".csv":
        for encoding in ("utf-8-sig", "utf-8", "cp1252"):
            try:
                with open(file_path, newline="", encoding=encoding) as csv_file:
                    sample = csv_file.read(4096)
                    csv_file.seek(0)
                    try:
                        dialect = csv.Sniffer().sniff(sample)
                    except csv.Error:
                        dialect = csv.excel
                    for row in csv.reader(csv_file, dialect):
                        yield row
                return
            except UnicodeDecodeError:
                continue
        raise ValueError("The CSV file could not be decoded. Please save it as UTF-8 or Windows CSV and try again.")
    if suffix == ".xlsx":
        if load_workbook is None:
            raise ValueError("openpyxl is not installed for .xlsx imports.")
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            worksheet = workbook[sheet_name]
            meaningful_width: int | None = None
            for row in worksheet.iter_rows(values_only=True):
                values = list(row)
                if meaningful_width is None:
                    meaningful_width = max(
                        (
                            index
                            for index, value in enumerate(values, start=1)
                            if value not in (None, "") and str(value).strip()
                        ),
                        default=0,
                    )
                    if meaningful_width == 0:
                        continue
                yield values[:meaningful_width]
        finally:
            workbook.close()
        return

    if suffix == ".xlsb":
        if open_xlsb_workbook is None:
            raise ValueError("pyxlsb is not installed for .xlsb imports.")
        with open_xlsb_workbook(file_path) as workbook:
            with workbook.get_sheet(sheet_name) as worksheet:
                meaningful_width: int | None = None
                for row in worksheet.rows():
                    values = [cell.v for cell in row]
                    if meaningful_width is None:
                        meaningful_width = max(
                            (
                                index
                                for index, value in enumerate(values, start=1)
                                if value not in (None, "") and str(value).strip()
                            ),
                            default=0,
                        )
                        if meaningful_width == 0:
                            continue
                    yield values[:meaningful_width]
        return

    raise ValueError("Unsupported file type.")


def _is_empty_row(values: list) -> bool:
    for value in values:
        if value not in (None, "") and str(value).strip():
            return False
    return True


def _clean_header_label(value, fallback_index: int) -> str:
    if value is None:
        return f"Column {fallback_index}"
    text = str(value).strip()
    return text or f"Column {fallback_index}"


def _make_unique_headers(raw_headers: list) -> list[str]:
    counts: dict[str, int] = {}
    unique_headers: list[str] = []
    for index, header in enumerate(raw_headers, start=1):
        base_label = _clean_header_label(header, index)
        occurrence = counts.get(base_label, 0) + 1
        counts[base_label] = occurrence
        unique_headers.append(base_label if occurrence == 1 else f"{base_label} ({occurrence})")
    return unique_headers


def _normalise_text(value) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def _coerce_preview_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _extract_headers_and_samples(file_path: str, sheet_name: str) -> tuple[list[str], list[list[str]]]:
    headers: list[str] = []
    samples: list[list[str]] = []

    for row in _iter_sheet_rows(file_path, sheet_name):
        if not headers:
            headers = _make_unique_headers(row)
            continue
        if _is_empty_row(row):
            continue
        sample_row = [
            _coerce_preview_value(row[index] if index < len(row) else None)
            for index in range(len(headers))
        ]
        samples.append(sample_row)
        if len(samples) >= PREVIEW_ROW_LIMIT:
            break

    if not headers:
        raise ValueError("The selected sheet does not contain a header row.")

    return headers, samples


def _suggest_mapping(headers: list[str], fields: list[ImportField]) -> dict[str, str]:
    suggestions: dict[str, str] = {}
    header_index = {_normalise_text(header): header for header in headers}

    for field in fields:
        for alias in field.aliases:
            alias_key = _normalise_text(alias)
            exact_match = header_index.get(alias_key)
            if exact_match:
                suggestions[field.key] = exact_match
                break
        if field.key in suggestions:
            continue
        for header in headers:
            header_key = _normalise_text(header)
            for alias in field.aliases:
                alias_key = _normalise_text(alias)
                if alias_key and (alias_key in header_key or header_key in alias_key):
                    suggestions[field.key] = header
                    break
            if field.key in suggestions:
                break

    return suggestions


def _build_layout_guidance(headers: list[str], config: dict, suggestions: dict[str, str]) -> dict[str, object]:
    required_fields = [field for field in config["fields"] if field.required]
    matched_required = [field.label for field in required_fields if suggestions.get(field.key)]
    numeric_like_headers = 0

    for header in headers:
        header_text = str(header).strip()
        if re.fullmatch(r"\d+(\.\d+)?", header_text):
            numeric_like_headers += 1

    looks_like_matrix = len(matched_required) <= 1 and numeric_like_headers >= max(3, len(headers) // 3)

    if looks_like_matrix:
        message = (
            "This workbook looks like a matrix or summary table, not a customer-level score import. "
            "The columns shown on the right are taken directly from the first row of the uploaded sheet. "
            "For this importer we expect one row per customer with fields such as Branch, Customer Number, "
            "Customer Name and Score, with Basel Final Grade only where applicable."
        )
        return {
            "level": "warning",
            "message": message,
            "matched_required": matched_required,
        }

    if not matched_required:
        message = (
            "No required model fields were recognised automatically from the uploaded headers. "
            "If this is the correct workbook, you can still map the columns manually."
        )
        return {
            "level": "info",
            "message": message,
            "matched_required": matched_required,
        }

    return {
        "level": "",
        "message": "",
        "matched_required": matched_required,
    }


def _parse_decimal(value) -> Decimal | None:
    if value in (None, ""):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip().replace(",", "").replace("%", "")
    if not text:
        return None
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def _excel_serial_to_date(value) -> date | None:
    if isinstance(value, bool):
        return None
    try:
        serial = float(value)
    except (TypeError, ValueError):
        return None
    if serial < 1 or serial > 80000:
        return None
    converted = EXCEL_SERIAL_DATE_START + timedelta(days=int(serial))
    if date(1900, 1, 1) <= converted <= date(2200, 12, 31):
        return converted
    return None


def _parse_date(value) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float, Decimal)):
        return _excel_serial_to_date(value)
    text = str(value).strip()
    if not text:
        return None
    text = text.strip("'\"")
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".", 1)[0]
    if re.fullmatch(r"\d{8}", text):
        compact_candidates = (
            ("%Y%m%d", text),
            ("%d%m%Y", text),
        )
        for fmt, candidate in compact_candidates:
            try:
                return datetime.strptime(candidate, fmt).date()
            except ValueError:
                continue
    serial_date = _excel_serial_to_date(text)
    if serial_date:
        return serial_date
    normalised_text = text
    if re.search(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}[T ]", text) or re.search(r"^\d{1,2}[-/]\d{1,2}[-/]\d{2,4}[T ]", text):
        normalised_text = re.split(r"[T ]", text, maxsplit=1)[0]
    normalised_text = normalised_text.replace(".", "/")
    for fmt in (
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%m/%d/%Y",
        "%d-%m-%Y",
        "%Y/%m/%d",
        "%d/%m/%y",
        "%m/%d/%y",
        "%d-%b-%Y",
        "%d %b %Y",
        "%d-%B-%Y",
        "%d %B %Y",
        "%b %d %Y",
        "%B %d %Y",
    ):
        try:
            return datetime.strptime(normalised_text, fmt).date()
        except ValueError:
            continue
    return None


def _normalise_score_range(
    value: Decimal | None,
    source_value,
    label: str,
    errors: list[str],
    warnings: list[str],
) -> Decimal | None:
    if source_value in (None, ""):
        return value
    if value is None:
        errors.append(f"{label} is invalid.")
        return value
    elif value < Decimal("0"):
        errors.append(f"{label} cannot be negative.")
        return value
    elif value > Decimal("100"):
        errors.append(f"{label} cannot be greater than 100.")
        return value
    elif Decimal("0") < value <= Decimal("1"):
        converted_value = value * Decimal("100")
        warnings.append(f"{label} ratio values between 0 and 1 will be multiplied by 100 during import.")
        return converted_value
    return value


def _clean_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


@lru_cache(maxsize=1)
def _branch_lookup() -> dict[str, str]:
    lookup: dict[str, str] = {}
    for branch_name, branch_code in BankBranch.objects.values_list("branch_name", "branch_code"):
        canonical_name = _clean_text(branch_name).upper()
        code_value = _clean_text(branch_code).upper()
        if canonical_name:
            lookup[canonical_name] = canonical_name
        if code_value:
            lookup[code_value] = canonical_name
    return lookup


def _clean_customer_id(value) -> str:
    raw_text = _clean_text(value)
    if not raw_text:
        return ""
    numeric_pattern = re.fullmatch(r"0*(\d+)(?:\.0+)?", raw_text)
    if numeric_pattern:
        cleaned = numeric_pattern.group(1).lstrip("0")
        return cleaned or "0"
    return raw_text


def _row_to_dict(headers: list[str], raw_row: list) -> dict[str, object]:
    return {
        header: raw_row[index] if index < len(raw_row) else None
        for index, header in enumerate(headers)
    }


def _mapped_value(mapping: dict[str, str], row: dict[str, object], field_key: str):
    source_column = mapping.get(field_key) or ""
    if not source_column:
        return None
    return row.get(source_column)


def _build_clean_result(
    *,
    row_number: int,
    row_data: dict[str, object],
    mapping: dict[str, str],
    config: dict,
) -> tuple[dict | None, list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    field_requirements = {field.key: field.required for field in config["fields"]}
    is_historical_scores = bool(config.get("is_historical_scores"))

    reporting_date = None
    if is_historical_scores:
        reporting_date_value = _mapped_value(mapping, row_data, "reporting_date")
        reporting_date = _parse_date(reporting_date_value)
        if reporting_date_value in (None, ""):
            errors.append("Reporting date is missing.")
        elif reporting_date is None:
            errors.append("Reporting date is invalid.")

    branch_value = _clean_text(_mapped_value(mapping, row_data, "branch_name")).upper()
    branch_name = branch_value
    if not branch_value:
        errors.append("Branch is missing.")
    else:
        canonical_branch = _branch_lookup().get(branch_value)
        if canonical_branch:
            branch_name = canonical_branch
        else:
            errors.append("Branch does not exist in BankBranch.")

    customer_id = _clean_customer_id(_mapped_value(mapping, row_data, "customer_id"))
    if not customer_id:
        errors.append("Customer number is missing.")

    customer_name = _clean_text(_mapped_value(mapping, row_data, "customer_name"))
    if not customer_name:
        errors.append("Customer name is missing.")

    if is_historical_scores:
        basel_score_value = _mapped_value(mapping, row_data, "basel_ii_score")
        basel_ii_score = _parse_decimal(basel_score_value)
        basel_ii_score = _normalise_score_range(basel_ii_score, basel_score_value, "Basel II score", errors, warnings)

        ifrs_score_value = _mapped_value(mapping, row_data, "ifrs_9_score")
        ifrs_9_score = _parse_decimal(ifrs_score_value)
        ifrs_9_score = _normalise_score_range(ifrs_9_score, ifrs_score_value, "IFRS9 score", errors, warnings)

        basel_ii_grade = _clean_text(_mapped_value(mapping, row_data, "basel_ii_grade"))[:10]
        basel_override_grade = _clean_text(_mapped_value(mapping, row_data, "basel_override_grade"))[:10]
        if basel_score_value in (None, "") and ifrs_score_value in (None, ""):
            errors.append("At least one of Basel II score or IFRS9 score is required.")

        if errors:
            return None, errors
        return {
            "row_number": row_number,
            "reporting_date": reporting_date,
            "branch_name": branch_name,
            "customer_id": customer_id,
            "customer_name": customer_name,
            "basel_ii_score": basel_ii_score,
            "basel_ii_grade": basel_ii_grade,
            "basel_override_grade": basel_override_grade,
            "ifrs_9_score": ifrs_9_score,
            "_warnings": warnings,
        }, []

    weighted_value = _mapped_value(mapping, row_data, "total_weighted_percent")
    total_weighted_percent = _parse_decimal(weighted_value)
    if weighted_value in (None, ""):
        errors.append("Weighted percent / score is missing.")
    else:
        total_weighted_percent = _normalise_score_range(
            total_weighted_percent,
            weighted_value,
            "Weighted percent / score",
            errors,
            warnings,
        )

    raw_score_value = _mapped_value(mapping, row_data, "total_raw_score")
    total_raw_score = _parse_decimal(raw_score_value)
    if raw_score_value not in (None, "") and total_raw_score is None:
        errors.append("Raw score is invalid.")

    final_grade = ""
    if config["supports_grade"]:
        final_grade = _clean_text(_mapped_value(mapping, row_data, "final_grade"))[:10]
        if field_requirements.get("final_grade") and not final_grade:
            errors.append("Final grade is missing.")

    if errors:
        return None, errors

    return {
        "row_number": row_number,
        "branch_name": branch_name,
        "customer_id": customer_id,
        "customer_name": customer_name,
        "total_weighted_percent": total_weighted_percent,
        "total_raw_score": total_raw_score,
        "final_grade": final_grade,
        "_warnings": warnings,
    }, []


def _build_error_summary(counter: Counter[str]) -> list[dict[str, object]]:
    return [
        {"reason": reason, "count": count}
        for reason, count in counter.most_common()
    ]


def _append_skipped_preview(
    skipped_preview: list[dict[str, str]],
    *,
    row_number: int | str,
    branch_name: str,
    customer_id: str,
    customer_name: str,
    errors: list[str],
    reporting_date: str = "",
    limit: int | None = PREVIEW_SCREEN_ROW_LIMIT,
) -> None:
    if limit is not None and len(skipped_preview) >= limit:
        return
    skipped_preview.append(
        {
            "row_number": str(row_number),
            "reporting_date": reporting_date,
            "branch_name": branch_name,
            "customer_id": customer_id,
            "customer_name": customer_name,
            "errors": errors,
            "errors_display": " ".join(errors),
        }
    )


def _build_invalid_export_row(
    *,
    row_number: int,
    mapping: dict[str, str],
    row_data: dict[str, object] | None = None,
    cleaned_row: dict[str, object] | None = None,
    errors: list[str],
    config: dict,
) -> dict[str, str]:
    reporting_date = ""
    branch_name = ""
    customer_id = ""
    customer_name = ""
    weighted_score = ""
    raw_score = ""
    final_grade = ""
    basel_override_grade = ""
    ifrs9_score = ""

    if cleaned_row:
        reporting_date = _coerce_preview_value(cleaned_row.get("reporting_date"))
        branch_name = _coerce_preview_value(cleaned_row.get("branch_name"))
        customer_id = _coerce_preview_value(cleaned_row.get("customer_id"))
        customer_name = _coerce_preview_value(cleaned_row.get("customer_name"))
        score_value = cleaned_row.get("total_weighted_percent")
        if score_value is None and config.get("is_historical_scores"):
            score_value = cleaned_row.get("basel_ii_score")
        weighted_score = _coerce_preview_value(score_value)
        raw_score = _coerce_preview_value(cleaned_row.get("total_raw_score"))
        grade_value = cleaned_row.get("final_grade")
        if grade_value in (None, "") and config.get("is_historical_scores"):
            grade_value = cleaned_row.get("basel_ii_grade")
        final_grade = _coerce_preview_value(grade_value)
        basel_override_grade = _coerce_preview_value(cleaned_row.get("basel_override_grade"))
        ifrs9_score = _coerce_preview_value(cleaned_row.get("ifrs_9_score"))

    if row_data:
        reporting_date = reporting_date or _coerce_preview_value(_mapped_value(mapping, row_data, "reporting_date"))
        branch_name = branch_name or _coerce_preview_value(_mapped_value(mapping, row_data, "branch_name"))
        customer_id = customer_id or _coerce_preview_value(_mapped_value(mapping, row_data, "customer_id"))
        customer_name = customer_name or _coerce_preview_value(_mapped_value(mapping, row_data, "customer_name"))
        if config.get("is_historical_scores"):
            weighted_score = weighted_score or _coerce_preview_value(_mapped_value(mapping, row_data, "basel_ii_score"))
            final_grade = final_grade or _coerce_preview_value(_mapped_value(mapping, row_data, "basel_ii_grade"))
            basel_override_grade = basel_override_grade or _coerce_preview_value(_mapped_value(mapping, row_data, "basel_override_grade"))
            ifrs9_score = ifrs9_score or _coerce_preview_value(_mapped_value(mapping, row_data, "ifrs_9_score"))
        else:
            weighted_score = weighted_score or _coerce_preview_value(_mapped_value(mapping, row_data, "total_weighted_percent"))
            raw_score = raw_score or _coerce_preview_value(_mapped_value(mapping, row_data, "total_raw_score"))
            if config["supports_grade"]:
                final_grade = final_grade or _coerce_preview_value(_mapped_value(mapping, row_data, "final_grade"))

    export_row = {
        "row_number": str(row_number),
        "reporting_date": reporting_date,
        "branch_name": branch_name,
        "customer_id": customer_id,
        "customer_name": customer_name,
        "weighted_score": weighted_score,
        "raw_score": raw_score,
        "final_grade": final_grade,
        "basel_override_grade": basel_override_grade,
        "ifrs9_score": ifrs9_score,
        "errors": list(errors),
        "errors_display": " | ".join(errors),
    }
    return export_row


def _collect_row_entries(file_path: str, sheet_name: str, mapping: dict[str, str], config: dict, *, max_rows: int | None = None) -> dict:
    headers: list[str] = []
    row_entries: list[dict[str, object]] = []
    total_rows = 0

    for index, raw_row in enumerate(_iter_sheet_rows(file_path, sheet_name), start=1):
        if not headers:
            headers = _make_unique_headers(raw_row)
            continue
        if _is_empty_row(raw_row):
            continue

        total_rows += 1
        row_data = _row_to_dict(headers, raw_row)
        cleaned_row, errors = _build_clean_result(
            row_number=index,
            row_data=row_data,
            mapping=mapping,
            config=config,
        )
        reporting_date = cleaned_row.get("reporting_date") if cleaned_row and config.get("is_historical_scores") else None
        branch_name = cleaned_row.get("branch_name") if cleaned_row else _clean_text(_mapped_value(mapping, row_data, "branch_name")).upper()
        customer_id = cleaned_row.get("customer_id") if cleaned_row else _clean_customer_id(_mapped_value(mapping, row_data, "customer_id"))
        customer_name = cleaned_row.get("customer_name") if cleaned_row else _clean_text(_mapped_value(mapping, row_data, "customer_name"))
        row_entries.append(
            {
                "row_number": index,
                "row_data": row_data,
                "cleaned_row": cleaned_row,
                "base_errors": list(errors),
                "reporting_date": reporting_date,
                "branch_name": branch_name,
                "customer_id": customer_id,
                "customer_name": customer_name,
            }
        )
        if max_rows and total_rows >= max_rows:
            break

    return {
        "total_rows": total_rows,
        "row_entries": row_entries,
    }


def _chunked_values(values, chunk_size: int = QUERY_CHUNK_SIZE):
    values_list = list(values)
    for index in range(0, len(values_list), chunk_size):
        yield values_list[index : index + chunk_size]


def _find_existing_keys(model, keys: set[tuple], *, is_historical_scores: bool = False) -> set[tuple]:
    if not keys:
        return set()
    branches = {key[-2] for key in keys}
    customer_ids = {key[-1] for key in keys}
    existing_keys: set[tuple] = set()
    if is_historical_scores:
        reporting_dates = {key[0] for key in keys}
        for customer_chunk in _chunked_values(customer_ids):
            existing_rows = model.objects.filter(
                branch_name__in=branches,
                customer_id__in=customer_chunk,
                reporting_date__in=reporting_dates,
            ).values_list("reporting_date", "branch_name", "customer_id")
            existing_keys.update(
                (reporting_date, branch_name, customer_id)
                for reporting_date, branch_name, customer_id in existing_rows
                if (reporting_date, branch_name, customer_id) in keys
            )
        return existing_keys
    for customer_chunk in _chunked_values(customer_ids):
        existing_pairs = model.objects.filter(
            branch_name__in=branches,
            customer_id__in=customer_chunk,
        ).values_list("branch_name", "customer_id")
        existing_keys.update(
            (branch_name, customer_id)
            for branch_name, customer_id in existing_pairs
            if (branch_name, customer_id) in keys
        )
    return existing_keys


def _evaluate_row_entries(
    row_entries: list[dict[str, object]],
    mapping: dict[str, str],
    config: dict,
    *,
    check_existing: bool = False,
) -> dict:
    is_historical_scores = bool(config.get("is_historical_scores"))
    duplicate_in_file_message = (
        "This customer appears more than once for the same reporting date and branch in the uploaded file."
        if is_historical_scores
        else "This customer appears more than once for the same branch in the uploaded file."
    )
    key_counter = Counter(
        (
            (item["reporting_date"], item["branch_name"], item["customer_id"])
            if is_historical_scores
            else (item["branch_name"], item["customer_id"])
        )
        for item in row_entries
        if item["branch_name"] and item["customer_id"] and (not is_historical_scores or item["reporting_date"])
    )
    keys = set(key_counter.keys())
    existing_keys = _find_existing_keys(config["model"], keys, is_historical_scores=is_historical_scores) if check_existing else set()
    duplicate_in_system_message = (
        "This historical score already exists for this reporting date, branch, and customer."
        if is_historical_scores
        else "This customer already exists in this branch."
    )

    ready_rows: list[dict[str, object]] = []
    invalid_rows = 0
    invalid_preview: list[dict[str, str]] = []
    invalid_export_rows: list[dict[str, str]] = []
    error_counter: Counter[str] = Counter()
    warning_counter: Counter[str] = Counter()

    for item in row_entries:
        cleaned_row = item["cleaned_row"]
        row_number = item["row_number"]
        key = (
            (item["reporting_date"], item["branch_name"], item["customer_id"])
            if is_historical_scores
            else (item["branch_name"], item["customer_id"])
        )
        errors: list[str] = list(item["base_errors"])

        if item["branch_name"] and item["customer_id"] and key_counter[key] > 1 and duplicate_in_file_message not in errors:
            errors.append(duplicate_in_file_message)
        if check_existing and item["branch_name"] and item["customer_id"] and key in existing_keys and duplicate_in_system_message not in errors:
            errors.append(duplicate_in_system_message)

        if errors:
            invalid_rows += 1
            error_counter.update(errors)
            _append_skipped_preview(
                invalid_preview,
                row_number=row_number,
                branch_name=item["branch_name"],
                customer_id=item["customer_id"],
                customer_name=item["customer_name"],
                errors=errors,
                reporting_date=item["reporting_date"].isoformat() if item.get("reporting_date") else "",
            )
            invalid_export_rows.append(
                _build_invalid_export_row(
                    row_number=row_number,
                    mapping=mapping,
                    row_data=item["row_data"],
                    cleaned_row=cleaned_row,
                    errors=errors,
                    config=config,
                )
            )
            continue

        if cleaned_row:
            row_warnings = list(cleaned_row.get("_warnings", []))
            warning_counter.update(row_warnings)
            ready_rows.append({"cleaned_row": cleaned_row, "row_number": row_number, "warnings": row_warnings})

    return {
        "candidates": ready_rows,
        "invalid_rows": invalid_rows,
        "invalid_preview": invalid_preview,
        "invalid_export_rows": invalid_export_rows,
        "error_counter": error_counter,
        "warning_counter": warning_counter,
    }


def _prepare_import_dataset(
    file_path: str,
    sheet_name: str,
    mapping: dict[str, str],
    config: dict,
    *,
    max_rows: int | None = None,
    check_existing: bool = False,
    preview_row_limit: int | None = PREVIEW_SCREEN_ROW_LIMIT,
) -> dict:
    collected = _collect_row_entries(file_path, sheet_name, mapping, config, max_rows=max_rows)
    evaluated = _evaluate_row_entries(collected["row_entries"], mapping, config, check_existing=check_existing)
    error_counter: Counter[str] = evaluated["error_counter"]
    warning_counter: Counter[str] = evaluated["warning_counter"]
    invalid_preview = evaluated["invalid_preview"]
    ready_rows = evaluated["candidates"]

    valid_preview: list[dict[str, str]] = []
    ready_rows_for_preview = ready_rows if preview_row_limit is None else ready_rows[:preview_row_limit]
    for item in ready_rows_for_preview:
        cleaned_row = item["cleaned_row"]
        score_value = cleaned_row.get("total_weighted_percent")
        if score_value is None and config.get("is_historical_scores"):
            score_value = cleaned_row.get("basel_ii_score")
        grade_value = cleaned_row.get("final_grade")
        if grade_value in (None, "") and config.get("is_historical_scores"):
            grade_value = cleaned_row.get("basel_ii_grade")
        valid_preview.append(
            {
                "row_number": str(cleaned_row["row_number"]),
                "reporting_date": "" if cleaned_row.get("reporting_date") is None else cleaned_row["reporting_date"].isoformat(),
                "branch_name": cleaned_row["branch_name"],
                "customer_id": cleaned_row["customer_id"],
                "customer_name": cleaned_row["customer_name"],
                "total_weighted_percent": "" if score_value is None else str(score_value),
                "total_raw_score": "" if cleaned_row.get("total_raw_score") is None else str(cleaned_row.get("total_raw_score")),
                "final_grade": "" if grade_value is None else str(grade_value),
                "basel_override_grade": _coerce_preview_value(cleaned_row.get("basel_override_grade")),
                "ifrs_9_score": "" if cleaned_row.get("ifrs_9_score") is None else str(cleaned_row.get("ifrs_9_score")),
                "warnings_display": " ".join(item.get("warnings", [])),
            }
        )

    return {
        "total_rows": collected["total_rows"],
        "valid_rows": len(ready_rows),
        "invalid_rows": evaluated["invalid_rows"],
        "valid_preview": valid_preview,
        "invalid_preview": invalid_preview,
        "invalid_export_rows": evaluated["invalid_export_rows"],
        "valid_preview_truncated": preview_row_limit is not None and len(ready_rows) > preview_row_limit,
        "invalid_preview_truncated": preview_row_limit is not None and evaluated["invalid_rows"] > len(invalid_preview),
        "preview_display_limit": preview_row_limit or "",
        "error_summary": _build_error_summary(error_counter),
        "warning_summary": _build_error_summary(warning_counter),
        "ready_rows": ready_rows,
        "is_fast_preview": bool(max_rows),
        "preview_scan_limit": max_rows or "",
        "existing_records_checked": check_existing,
    }


def _preview_row_from_cleaned(cleaned_row: dict[str, object], config: dict) -> dict[str, str]:
    score_value = cleaned_row.get("total_weighted_percent")
    if score_value is None and config.get("is_historical_scores"):
        score_value = cleaned_row.get("basel_ii_score")
    grade_value = cleaned_row.get("final_grade")
    if grade_value in (None, "") and config.get("is_historical_scores"):
        grade_value = cleaned_row.get("basel_ii_grade")
    return {
        "row_number": str(cleaned_row["row_number"]),
        "reporting_date": (
            ""
            if cleaned_row.get("reporting_date") is None
            else cleaned_row["reporting_date"].isoformat()
        ),
        "branch_name": cleaned_row["branch_name"],
        "customer_id": cleaned_row["customer_id"],
        "customer_name": cleaned_row["customer_name"],
        "total_weighted_percent": "" if score_value is None else str(score_value),
        "total_raw_score": (
            ""
            if cleaned_row.get("total_raw_score") is None
            else str(cleaned_row.get("total_raw_score"))
        ),
        "final_grade": "" if grade_value is None else str(grade_value),
        "basel_override_grade": _coerce_preview_value(cleaned_row.get("basel_override_grade")),
        "ifrs_9_score": (
            ""
            if cleaned_row.get("ifrs_9_score") is None
            else str(cleaned_row.get("ifrs_9_score"))
        ),
        "warnings_display": " ".join(cleaned_row.get("_warnings", [])),
    }


def _prepare_compact_preview_dataset(
    file_path: str,
    sheet_name: str,
    mapping: dict[str, str],
    config: dict,
    *,
    preview_row_limit: int = PREVIEW_SCREEN_ROW_LIMIT,
) -> dict:
    """Validate the complete file without retaining the complete file in memory."""
    is_historical_scores = bool(config.get("is_historical_scores"))
    duplicate_message = (
        "This customer appears more than once for the same reporting date and branch in the uploaded file."
        if is_historical_scores
        else "This customer appears more than once for the same branch in the uploaded file."
    )

    compact_rows: list[tuple] = []
    key_counter: Counter[tuple] = Counter()
    total_rows = 0
    for cleaned_row, errors, row_data, row_number in _iter_clean_rows(
        file_path, sheet_name, mapping, config
    ):
        total_rows += 1
        reporting_date = (
            cleaned_row.get("reporting_date")
            if cleaned_row and is_historical_scores
            else _parse_date(_mapped_value(mapping, row_data, "reporting_date"))
            if is_historical_scores
            else None
        )
        branch_name = (
            cleaned_row.get("branch_name")
            if cleaned_row
            else _clean_text(_mapped_value(mapping, row_data, "branch_name")).upper()
        )
        customer_id = (
            cleaned_row.get("customer_id")
            if cleaned_row
            else _clean_customer_id(_mapped_value(mapping, row_data, "customer_id"))
        )
        if branch_name and customer_id and (not is_historical_scores or reporting_date):
            key = (
                (reporting_date, branch_name, customer_id)
                if is_historical_scores
                else (branch_name, customer_id)
            )
            key_counter[key] += 1
        compact_rows.append(
            (
                cleaned_row,
                tuple(errors),
                reporting_date,
                branch_name,
                customer_id,
                (
                    cleaned_row.get("customer_name")
                    if cleaned_row
                    else _clean_text(_mapped_value(mapping, row_data, "customer_name"))
                ),
                row_number,
            )
        )

    valid_rows = 0
    invalid_rows = 0
    valid_preview: list[dict[str, str]] = []
    invalid_preview: list[dict[str, str]] = []
    error_counter: Counter[str] = Counter()
    warning_counter: Counter[str] = Counter()

    for (
        cleaned_row,
        base_errors,
        reporting_date,
        branch_name,
        customer_id,
        customer_name,
        row_number,
    ) in compact_rows:
        errors = list(base_errors)
        key = (
            (reporting_date, branch_name, customer_id)
            if is_historical_scores
            else (branch_name, customer_id)
        )
        if branch_name and customer_id and key_counter[key] > 1:
            errors.append(duplicate_message)

        if errors:
            invalid_rows += 1
            error_counter.update(errors)
            _append_skipped_preview(
                invalid_preview,
                row_number=row_number,
                reporting_date=reporting_date.isoformat() if reporting_date else "",
                branch_name=branch_name,
                customer_id=customer_id,
                customer_name=customer_name,
                errors=errors,
                limit=preview_row_limit,
            )
            continue

        if cleaned_row:
            valid_rows += 1
            warning_counter.update(cleaned_row.get("_warnings", []))
            if len(valid_preview) < preview_row_limit:
                valid_preview.append(_preview_row_from_cleaned(cleaned_row, config))

    return {
        "cache_mode": PREVIEW_CACHE_MODE_COMPACT,
        "total_rows": total_rows,
        "valid_rows": valid_rows,
        "invalid_rows": invalid_rows,
        "valid_preview": valid_preview,
        "invalid_preview": invalid_preview,
        "invalid_export_rows": [],
        "ready_rows": [],
        "valid_preview_truncated": valid_rows > len(valid_preview),
        "invalid_preview_truncated": invalid_rows > len(invalid_preview),
        "preview_display_limit": preview_row_limit,
        "error_summary": _build_error_summary(error_counter),
        "warning_summary": _build_error_summary(warning_counter),
        "is_fast_preview": True,
        "preview_scan_limit": "",
        "existing_records_checked": False,
    }


def _collect_invalid_rows_for_export(file_path: str, sheet_name: str, mapping: dict[str, str], config: dict) -> list[dict[str, str]]:
    collected = _collect_row_entries(file_path, sheet_name, mapping, config)
    evaluated = _evaluate_row_entries(collected["row_entries"], mapping, config)
    return evaluated["invalid_export_rows"]


def _slugify_download_label(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", (value or "").strip().lower()).strip("-")
    return slug or "all-skipped-rows"


def _build_error_export_response(
    *,
    rows: list[dict[str, str]],
    config: dict,
    original_name: str,
    sheet_name: str,
    reason: str,
) -> HttpResponse:
    if Workbook is None:
        raise ValueError("openpyxl is not installed for Excel exports.")

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_rows = [
        ("Import type", config["label"]),
        ("Source file", original_name),
        ("Sheet", sheet_name),
        ("Downloaded rows", len(rows)),
        ("Filter", reason or "All skipped rows"),
        ("Generated at", timezone.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for row_index, (label, value) in enumerate(summary_rows, start=1):
        summary_sheet.cell(row=row_index, column=1, value=label)
        summary_sheet.cell(row=row_index, column=2, value=value)

    detail_sheet = workbook.create_sheet("Customers With Errors")
    if config.get("is_historical_scores"):
        headers = [
            "Excel Row",
            "Reporting Date",
            "Branch",
            "Customer Number",
            "Customer Name",
            "Basel II Score",
            "Basel II Grade",
            "Basel Override Grade",
            "IFRS9 Score",
            "Validation Errors",
        ]
    else:
        headers = [
            "Excel Row",
            "Branch",
            "Customer Number",
            "Customer Name",
            "Weighted Score",
            "Raw Score",
        ]
        if config["supports_grade"]:
            headers.append("Final Grade")
        headers.append("Validation Errors")

    for column_index, header in enumerate(headers, start=1):
        detail_sheet.cell(row=1, column=column_index, value=header)

    for row_index, row in enumerate(rows, start=2):
        if config.get("is_historical_scores"):
            values = [
                row["row_number"],
                row["reporting_date"],
                row["branch_name"],
                row["customer_id"],
                row["customer_name"],
                row["weighted_score"],
                row["final_grade"],
                row["basel_override_grade"],
                row["ifrs9_score"],
                row["errors_display"],
            ]
        else:
            values = [
                row["row_number"],
                row["branch_name"],
                row["customer_id"],
                row["customer_name"],
                row["weighted_score"],
                row["raw_score"],
            ]
            if config["supports_grade"]:
                values.append(row["final_grade"])
            values.append(row["errors_display"])
        for column_index, value in enumerate(values, start=1):
            detail_sheet.cell(row=row_index, column=column_index, value=value)

    for sheet in (summary_sheet, detail_sheet):
        for column_cells in sheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                if len(value) > max_length:
                    max_length = len(value)
            sheet.column_dimensions[column_letter].width = min(max_length + 3, 60)

    file_suffix = _slugify_download_label(reason)
    stem = Path(original_name or "import").stem
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        f'attachment; filename="{stem}_{file_suffix}_errors.xlsx"'
    )
    workbook.save(response)
    return response


def _build_full_preview_export_response(
    *,
    analysis: dict,
    config: dict,
    original_name: str,
    sheet_name: str,
) -> tuple[FileResponse, str]:
    if xlsxwriter is None:
        raise ValueError("XlsxWriter is not installed for fast Excel exports.")

    temp_dir = Path(tempfile.gettempdir()) / "scorecard_external_imports"
    temp_dir.mkdir(parents=True, exist_ok=True)
    export_path = temp_dir / f"full_preview_{uuid.uuid4().hex}.xlsx"
    workbook = xlsxwriter.Workbook(str(export_path), {"constant_memory": True})
    header_format = workbook.add_format({"bold": True, "bg_color": "#DCEBFA"})

    summary_sheet = workbook.add_worksheet("Summary")
    summary_rows = [
        ("Import type", config["label"]),
        ("Source file", original_name),
        ("Sheet", sheet_name),
        ("Total rows scanned", analysis["total_rows"]),
        ("Valid rows", analysis["valid_rows"]),
        ("Skipped rows", analysis["invalid_rows"]),
        ("Generated at", timezone.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for row_index, values in enumerate(summary_rows):
        summary_sheet.write_row(row_index, 0, values)

    cleaned_sheet = workbook.add_worksheet("Cleaned Rows")
    if config.get("is_historical_scores"):
        cleaned_headers = [
            "Excel Row",
            "Reporting Date",
            "Branch",
            "Customer Number",
            "Customer Name",
            "Basel II Score",
            "Basel II Grade",
            "Basel Override Grade",
            "IFRS9 Score",
        ]
    else:
        cleaned_headers = [
            "Excel Row",
            "Branch",
            "Customer Number",
            "Customer Name",
            "Weighted Score",
            "Raw Score",
        ]
        if config["supports_grade"]:
            cleaned_headers.append("Final Grade")
    cleaned_sheet.write_row(0, 0, cleaned_headers, header_format)

    for row_index, item in enumerate(analysis["ready_rows"], start=1):
        cleaned_row = item["cleaned_row"]
        if config.get("is_historical_scores"):
            values = [
                cleaned_row["row_number"],
                _coerce_preview_value(cleaned_row.get("reporting_date")),
                cleaned_row["branch_name"],
                cleaned_row["customer_id"],
                cleaned_row["customer_name"],
                _coerce_preview_value(cleaned_row.get("basel_ii_score")),
                _coerce_preview_value(cleaned_row.get("basel_ii_grade")),
                _coerce_preview_value(cleaned_row.get("basel_override_grade")),
                _coerce_preview_value(cleaned_row.get("ifrs_9_score")),
            ]
        else:
            values = [
                cleaned_row["row_number"],
                cleaned_row["branch_name"],
                cleaned_row["customer_id"],
                cleaned_row["customer_name"],
                _coerce_preview_value(cleaned_row.get("total_weighted_percent")),
                _coerce_preview_value(cleaned_row.get("total_raw_score")),
            ]
            if config["supports_grade"]:
                values.append(_coerce_preview_value(cleaned_row.get("final_grade")))
        cleaned_sheet.write_row(row_index, 0, values)

    skipped_sheet = workbook.add_worksheet("Skipped Rows")
    if config.get("is_historical_scores"):
        skipped_headers = [
            "Excel Row",
            "Reporting Date",
            "Branch",
            "Customer Number",
            "Customer Name",
            "Basel II Score",
            "Basel II Grade",
            "Basel Override Grade",
            "IFRS9 Score",
            "Validation Errors",
        ]
    else:
        skipped_headers = [
            "Excel Row",
            "Branch",
            "Customer Number",
            "Customer Name",
            "Weighted Score",
            "Raw Score",
        ]
        if config["supports_grade"]:
            skipped_headers.append("Final Grade")
        skipped_headers.append("Validation Errors")
    skipped_sheet.write_row(0, 0, skipped_headers, header_format)

    for row_index, row in enumerate(analysis["invalid_export_rows"], start=1):
        if config.get("is_historical_scores"):
            values = [
                row["row_number"],
                row["reporting_date"],
                row["branch_name"],
                row["customer_id"],
                row["customer_name"],
                row["weighted_score"],
                row["final_grade"],
                row["basel_override_grade"],
                row["ifrs9_score"],
                row["errors_display"],
            ]
        else:
            values = [
                row["row_number"],
                row["branch_name"],
                row["customer_id"],
                row["customer_name"],
                row["weighted_score"],
                row["raw_score"],
            ]
            if config["supports_grade"]:
                values.append(row["final_grade"])
            values.append(row["errors_display"])
        skipped_sheet.write_row(row_index, 0, values)

    workbook.close()
    stem = Path(original_name or "import").stem
    response = FileResponse(
        export_path.open("rb"),
        as_attachment=True,
        filename=f"{stem}_full_preview.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    return response, str(export_path)


def _iter_clean_rows(file_path: str, sheet_name: str, mapping: dict[str, str], config: dict):
    headers: list[str] = []

    for index, raw_row in enumerate(_iter_sheet_rows(file_path, sheet_name), start=1):
        if not headers:
            headers = _make_unique_headers(raw_row)
            continue
        if _is_empty_row(raw_row):
            continue
        row_data = _row_to_dict(headers, raw_row)
        cleaned_row, errors = _build_clean_result(
            row_number=index,
            row_data=row_data,
            mapping=mapping,
            config=config,
        )
        yield cleaned_row, errors, row_data, index


def _build_row_payload(cleaned_row: dict[str, object], user, metadata: dict, now, model):
    if model is HistoricalScore:
        return model(
            reporting_date=cleaned_row["reporting_date"],
            branch_name=cleaned_row["branch_name"],
            customer_name=cleaned_row["customer_name"],
            customer_id=cleaned_row["customer_id"],
            basel_ii_score=cleaned_row["basel_ii_score"],
            basel_ii_grade=cleaned_row["basel_ii_grade"],
            basel_override_grade=cleaned_row.get("basel_override_grade", ""),
            ifrs_9_score=cleaned_row["ifrs_9_score"],
            created_at=now,
            updated_at=now,
        )
    payload = {
        "template": None,
        "template_section_name": "",
        "branch_name": cleaned_row["branch_name"],
        "customer_name": cleaned_row["customer_name"],
        "customer_id": cleaned_row["customer_id"],
        "total_raw_score": cleaned_row["total_raw_score"],
        "total_weighted_percent": cleaned_row["total_weighted_percent"],
        "final_grade": cleaned_row["final_grade"],
        "status": "approved",
        "maker": user,
        "submitted_by": user,
        "submitted_at": now,
        "approved_by": user,
        "approved_at": now,
        "approved_weighted_percent": cleaned_row["total_weighted_percent"],
        "approved_grade": cleaned_row["final_grade"],
        "autofill_metadata": metadata,
        "created_at": now,
        "updated_at": now,
    }
    return model(**payload)


def _import_key(cleaned_row: dict[str, object], model) -> tuple:
    if model is HistoricalScore:
        return (cleaned_row["reporting_date"], cleaned_row["branch_name"], cleaned_row["customer_id"])
    return (cleaned_row["branch_name"], cleaned_row["customer_id"])


def _existing_import_objects(model, batch: list[tuple[object, dict[str, object]]]) -> dict[tuple, object]:
    if not batch:
        return {}
    keys = {_import_key(cleaned_row, model) for _, cleaned_row in batch}
    branches = {key[-2] for key in keys}
    customer_ids = {key[-1] for key in keys}
    existing: dict[tuple, object] = {}
    if model is HistoricalScore:
        reporting_dates = {key[0] for key in keys}
        queryset = model.objects.filter(
            reporting_date__in=reporting_dates,
            branch_name__in=branches,
            customer_id__in=customer_ids,
        )
        for obj in queryset:
            key = (obj.reporting_date, obj.branch_name, obj.customer_id)
            if key in keys:
                existing[key] = obj
        return existing
    queryset = model.objects.filter(branch_name__in=branches, customer_id__in=customer_ids)
    for obj in queryset:
        key = (obj.branch_name, obj.customer_id)
        if key in keys:
            existing[key] = obj
    return existing


def _apply_import_override(existing_obj, incoming_obj, cleaned_row: dict[str, object], model, now) -> None:
    if model is HistoricalScore:
        existing_obj.customer_name = cleaned_row["customer_name"]
        existing_obj.basel_ii_score = cleaned_row["basel_ii_score"]
        existing_obj.basel_ii_grade = cleaned_row["basel_ii_grade"]
        existing_obj.basel_override_grade = cleaned_row.get("basel_override_grade", "")
        existing_obj.ifrs_9_score = cleaned_row["ifrs_9_score"]
        existing_obj.updated_at = now
        return
    existing_obj.template = None
    existing_obj.template_section_name = ""
    existing_obj.customer_name = cleaned_row["customer_name"]
    existing_obj.total_raw_score = cleaned_row["total_raw_score"]
    existing_obj.total_weighted_percent = cleaned_row["total_weighted_percent"]
    existing_obj.final_grade = cleaned_row["final_grade"]
    existing_obj.status = "approved"
    existing_obj.maker = incoming_obj.maker
    existing_obj.submitted_by = incoming_obj.submitted_by
    existing_obj.submitted_at = incoming_obj.submitted_at
    existing_obj.approved_by = incoming_obj.approved_by
    existing_obj.approved_at = incoming_obj.approved_at
    existing_obj.approved_weighted_percent = cleaned_row["total_weighted_percent"]
    existing_obj.approved_grade = cleaned_row["final_grade"]
    existing_obj.autofill_metadata = incoming_obj.autofill_metadata
    existing_obj.updated_at = now


def _override_update_fields(model) -> list[str]:
    if model is HistoricalScore:
        return [
            "customer_name",
            "basel_ii_score",
            "basel_ii_grade",
            "basel_override_grade",
            "ifrs_9_score",
            "updated_at",
        ]
    return [
        "template",
        "template_section_name",
        "customer_name",
        "total_raw_score",
        "total_weighted_percent",
        "final_grade",
        "status",
        "maker",
        "submitted_by",
        "submitted_at",
        "approved_by",
        "approved_at",
        "approved_weighted_percent",
        "approved_grade",
        "autofill_metadata",
        "updated_at",
    ]


def _row_integrity_errors(exc: IntegrityError, cleaned_row: dict[str, object]) -> list[str]:
    message = str(exc)
    if "uq_hist_scores_reporting_branch_customer" in message:
        return [
            f"Historical score for customer '{cleaned_row['customer_id']}' already exists in branch '{cleaned_row['branch_name']}' for {cleaned_row['reporting_date']}.",
        ]
    if "uq_credit_eval_branch_customer" in message or "uq_ifrs9_eval_branch_customer" in message:
        return [
            f"Customer '{cleaned_row['customer_id']}' already exists in branch '{cleaned_row['branch_name']}'.",
        ]
    return ["This row could not be imported because it violates a database integrity rule."]


def _flush_batch(model, batch: list[tuple[object, dict[str, object]]], *, inserted_rows: int, skipped_rows: int, skipped_preview: list[dict[str, str]], error_counter: Counter[str]) -> tuple[int, int]:
    if not batch:
        return inserted_rows, skipped_rows

    now = timezone.now()
    existing_objects = _existing_import_objects(model, batch)
    create_objects = []
    update_objects = []
    for obj, cleaned_row in batch:
        existing_obj = existing_objects.get(_import_key(cleaned_row, model))
        if existing_obj is None:
            create_objects.append((obj, cleaned_row))
            continue
        _apply_import_override(existing_obj, obj, cleaned_row, model, now)
        update_objects.append(existing_obj)

    database_batch_size = BULK_CREATE_BATCH_SIZE if model is HistoricalScore else EVALUATION_IMPORT_BATCH_SIZE
    try:
        with transaction.atomic():
            if create_objects:
                model.objects.bulk_create([obj for obj, _ in create_objects], batch_size=database_batch_size)
            if update_objects:
                model.objects.bulk_update(update_objects, _override_update_fields(model), batch_size=database_batch_size)
        inserted_rows += len(create_objects) + len(update_objects)
        return inserted_rows, skipped_rows
    except IntegrityError:
        for obj, cleaned_row in create_objects:
            try:
                with transaction.atomic():
                    obj.save(force_insert=True)
                inserted_rows += 1
            except IntegrityError as exc:
                lookup = (
                    {
                        "reporting_date": cleaned_row["reporting_date"],
                        "branch_name": cleaned_row["branch_name"],
                        "customer_id": cleaned_row["customer_id"],
                    }
                    if model is HistoricalScore
                    else {"branch_name": cleaned_row["branch_name"], "customer_id": cleaned_row["customer_id"]}
                )
                existing_obj = model.objects.filter(**lookup).first()
                if existing_obj is not None:
                    _apply_import_override(existing_obj, obj, cleaned_row, model, timezone.now())
                    existing_obj.save(update_fields=_override_update_fields(model))
                    inserted_rows += 1
                    continue
                skipped_rows += 1
                errors = _row_integrity_errors(exc, cleaned_row)
                error_counter.update(errors)
                _append_skipped_preview(
                    skipped_preview,
                    row_number=cleaned_row["row_number"],
                    branch_name=cleaned_row["branch_name"],
                    customer_id=cleaned_row["customer_id"],
                    customer_name=cleaned_row["customer_name"],
                    errors=errors,
                )
        return inserted_rows, skipped_rows


def _build_import_metadata_for_user(user, upload_kind: str, state: dict, mapping: dict[str, str]) -> dict:
    return {
        "external_import": {
            "upload_kind": upload_kind,
            "uploaded_by_id": user.id,
            "uploaded_by_username": user.get_username(),
            "uploaded_at": timezone.now().isoformat(),
            "original_filename": state.get("original_name", ""),
            "sheet_name": state.get("sheet_name", ""),
            "column_mapping": mapping,
            "cleaning_rules": [
                "Trim whitespace",
                "Uppercase branch_name",
                "Remove leading zeroes from numeric customer_id values",
                "Coerce reporting dates and score columns to the expected data types",
            ],
        }
    }


def _build_import_metadata(request: HttpRequest, upload_kind: str, state: dict, mapping: dict[str, str]) -> dict:
    return _build_import_metadata_for_user(request.user, upload_kind, state, mapping)


def _perform_upload_import(*, upload_kind: str, state: dict, user, job_path: str) -> dict:
    """Run a validated external import without tying up the web request."""
    config = _get_import_config(upload_kind)
    model = config["model"]
    file_name = state.get("original_name", "the uploaded file")
    started_at = timezone.now()
    metadata = _build_import_metadata_for_user(user, upload_kind, state, state["mapping"])
    prepared = _load_preview_cache(state.get("preview_cache_path"))
    compact_import = bool(prepared and prepared.get("cache_mode") == PREVIEW_CACHE_MODE_COMPACT)
    if prepared is None:
        prepared = _prepare_import_dataset(
            file_path=state["file_path"],
            sheet_name=state["sheet_name"],
            mapping=state["mapping"],
            config=config,
        )
    if prepared["invalid_rows"] > 0:
        raise ValueError(
            "Import is disabled because the preview found validation issues. "
            "Correct the skipped rows before loading."
        )

    inserted_rows = 0
    skipped_rows = prepared["invalid_rows"]
    skipped_preview = list(prepared["invalid_preview"])
    warning_summary = list(prepared.get("warning_summary", []))
    error_counter = Counter()
    for item in prepared["error_summary"]:
        error_counter[item["reason"]] = item["count"]

    total_rows = int(prepared.get("valid_rows") or 0)
    processed_rows = 0
    batch: list[tuple[object, dict[str, object]]] = []
    now = timezone.now()
    if compact_import:
        clean_row_items = (
            cleaned_row
            for cleaned_row, errors, _row_data, _row_number in _iter_clean_rows(
                state["file_path"], state["sheet_name"], state["mapping"], config
            )
            if cleaned_row is not None and not errors
        )
    else:
        clean_row_items = (item["cleaned_row"] for item in prepared["ready_rows"])

    import_batch_size = BULK_CREATE_BATCH_SIZE if model is HistoricalScore else EVALUATION_IMPORT_BATCH_SIZE
    for cleaned_row in clean_row_items:
        obj = _build_row_payload(cleaned_row, user, metadata, now, model)
        batch.append((obj, cleaned_row))
        if len(batch) >= import_batch_size:
            inserted_rows, skipped_rows = _flush_batch(
                model,
                batch,
                inserted_rows=inserted_rows,
                skipped_rows=skipped_rows,
                skipped_preview=skipped_preview,
                error_counter=error_counter,
            )
            processed_rows += len(batch)
            batch = []
            _save_upload_job_state(
                job_path,
                {
                    "status": "running",
                    "file_name": file_name,
                    "sheet_name": state.get("sheet_name", ""),
                    "target_table": config["table_name"],
                    "inserted_rows": inserted_rows,
                    "skipped_rows": skipped_rows,
                    "processed_rows": processed_rows,
                    "total_rows": total_rows,
                    "message": "Importing validated rows in the background.",
                    "started_at": started_at.strftime("%Y-%m-%d %H:%M:%S"),
                },
            )

    if batch:
        inserted_rows, skipped_rows = _flush_batch(
            model,
            batch,
            inserted_rows=inserted_rows,
            skipped_rows=skipped_rows,
            skipped_preview=skipped_preview,
            error_counter=error_counter,
        )
        processed_rows += len(batch)

    return {
        "status": "completed",
        "file_name": file_name,
        "sheet_name": state.get("sheet_name", ""),
        "target_table": config["table_name"],
        "inserted_rows": inserted_rows,
        "skipped_rows": skipped_rows,
        "processed_rows": processed_rows,
        "total_rows": total_rows,
        "error_summary": _build_error_summary(error_counter),
        "warning_summary": warning_summary,
        "skipped_preview": skipped_preview,
        "imported_at": timezone.now().strftime("%Y-%m-%d %H:%M:%S"),
        "started_at": started_at.strftime("%Y-%m-%d %H:%M:%S"),
        "message": "Import completed.",
    }


def _run_upload_import_job(*, job_path: str, upload_kind: str, state: dict, user_id: int) -> None:
    config = _get_import_config(upload_kind)
    file_name = state.get("original_name", "the uploaded file")
    close_old_connections()
    try:
        user = get_user_model().objects.get(pk=user_id)
        result_state = _perform_upload_import(
            upload_kind=upload_kind,
            state=state,
            user=user,
            job_path=job_path,
        )
        import_details = (
            f"Loaded rows: {result_state['inserted_rows']}; "
            f"Skipped rows: {result_state['skipped_rows']}; Target table: {config['table_name']}"
        )
        log_upload_audit(
            user,
            "import_complete",
            upload_kind=upload_kind,
            file_name=file_name,
            details=import_details,
            object_id=config["table_name"],
        )
        if config.get("is_historical_scores"):
            log_historical_score_audit(
                user,
                "import_complete",
                details=f"File: {file_name}; {import_details}",
                object_id=config["table_name"],
            )
    except Exception as exc:
        error_message = str(exc) or exc.__class__.__name__
        result_state = {
            "status": "failed",
            "file_name": file_name,
            "sheet_name": state.get("sheet_name", ""),
            "target_table": config["table_name"],
            "message": error_message,
            "inserted_rows": 0,
            "skipped_rows": 0,
            "processed_rows": 0,
            "total_rows": 0,
            "error_summary": [{"reason": error_message, "count": 1}],
            "skipped_preview": [],
            "imported_at": timezone.now().strftime("%Y-%m-%d %H:%M:%S"),
        }
        try:
            user = get_user_model().objects.get(pk=user_id)
            log_upload_audit(
                user,
                "import_failed",
                upload_kind=upload_kind,
                file_name=file_name,
                details=error_message,
                object_id=config["table_name"],
            )
            if config.get("is_historical_scores"):
                log_historical_score_audit(
                    user,
                    "import_failed",
                    details=f"File: {file_name}; {error_message}",
                    object_id=config["table_name"],
                )
        except Exception:
            pass
    finally:
        _save_upload_job_state(job_path, result_state)
        _delete_temp_file(state.get("file_path"))
        _delete_temp_file(state.get("preview_cache_path"))
        _delete_temp_file(state.get("preview_export_path"))
        close_old_connections()
        with _upload_job_threads_lock:
            _upload_job_threads.pop(job_path, None)


@login_required
def upload_home_view(request: HttpRequest) -> HttpResponse:
    cards = []
    for key, config in IMPORT_CONFIGS.items():
        cards.append(
            {
                "key": key,
                "label": config["label"],
                "table_name": config["table_name"],
                "start_url": reverse("scorecard:upload_start", kwargs={"upload_kind": key}),
            }
        )
    return render(
        request,
        "uploads/home.html",
        {
            "cards": cards,
        },
    )


@login_required
@require_http_methods(["GET", "POST"])
def upload_start_view(request: HttpRequest, upload_kind: str) -> HttpResponse:
    config = _get_import_config(upload_kind)
    state = _get_state(request, upload_kind)

    if request.method == "POST":
        uploaded_file = request.FILES.get("source_file")
        if uploaded_file:
            _delete_temp_file(state.get("file_path"))
            _delete_temp_file(state.get("preview_cache_path"))
            _delete_temp_file(state.get("preview_export_path"))
            try:
                file_path = _save_uploaded_file(uploaded_file)
                sheet_names = _list_sheet_names(file_path)
            except ValueError as exc:
                messages.error(request, str(exc))
                return redirect("scorecard:upload_start", upload_kind=upload_kind)

            if not sheet_names:
                _delete_temp_file(file_path)
                messages.error(request, "The uploaded workbook does not contain any sheets.")
                return redirect("scorecard:upload_start", upload_kind=upload_kind)

            state = {
                "file_path": file_path,
                "original_name": uploaded_file.name,
                "sheet_names": sheet_names,
                "sheet_name": sheet_names[0],
                "mapping": {},
            }
            _store_state(request, upload_kind, state)
        else:
            state = _get_state(request, upload_kind)
            if not state.get("file_path") or not state.get("sheet_name"):
                messages.error(request, "Please upload a file before continuing.")
                return redirect("scorecard:upload_start", upload_kind=upload_kind)

        try:
            headers, sample_rows = _extract_headers_and_samples(state["file_path"], state["sheet_name"])
        except (KeyError, ValueError) as exc:
            messages.error(request, str(exc))
            return redirect("scorecard:upload_start", upload_kind=upload_kind)

        state.update(
            {
                "headers": headers,
                "sample_rows": sample_rows,
                "mapping": state.get("mapping", {}),
            }
        )
        _store_state(request, upload_kind, state)
        return redirect("scorecard:upload_mapping", upload_kind=upload_kind)

    return render(
        request,
        "uploads/start.html",
        {
            "config": config,
            "upload_kind": upload_kind,
            "state": state,
        },
    )


@login_required
@require_http_methods(["GET", "POST"])
def upload_mapping_view(request: HttpRequest, upload_kind: str) -> HttpResponse:
    config = _get_import_config(upload_kind)
    state = _get_state(request, upload_kind)
    if not state.get("file_path") or not state.get("sheet_name") or not state.get("headers"):
        messages.info(request, "Start by uploading a file and selecting a sheet.")
        return redirect("scorecard:upload_start", upload_kind=upload_kind)

    headers = state["headers"]
    suggestions = _suggest_mapping(headers, config["fields"])
    layout_guidance = _build_layout_guidance(headers, config, suggestions)
    current_mapping = state.get("mapping", {})

    if request.method == "POST":
        mapping = {}
        for field in config["fields"]:
            mapping[field.key] = request.POST.get(field.key, "").strip()

        missing_required = [field.label for field in config["fields"] if field.required and not mapping.get(field.key)]
        if missing_required:
            messages.error(request, "Map all required columns before continuing.")
            current_mapping = mapping
        else:
            _delete_temp_file(state.get("preview_cache_path"))
            _delete_temp_file(state.get("preview_export_path"))
            state["mapping"] = mapping
            state.pop("preview_cache_path", None)
            state.pop("preview_export_path", None)
            _store_state(request, upload_kind, state)
            return redirect("scorecard:upload_preview", upload_kind=upload_kind)

    field_rows = []
    for field in config["fields"]:
        field_rows.append(
            {
                "key": field.key,
                "label": field.label,
                "required": field.required,
                "help_text": field.help_text,
                "selected": current_mapping.get(field.key) or suggestions.get(field.key, ""),
            }
        )

    return render(
        request,
        "uploads/mapping.html",
        {
            "config": config,
            "upload_kind": upload_kind,
            "headers": headers,
            "sample_rows": state.get("sample_rows", []),
            "field_rows": field_rows,
            "sheet_name": state.get("sheet_name"),
            "original_name": state.get("original_name"),
            "layout_guidance": layout_guidance,
        },
    )


@login_required
@require_http_methods(["GET"])
def upload_preview_view(request: HttpRequest, upload_kind: str) -> HttpResponse:
    config = _get_import_config(upload_kind)
    state = _get_state(request, upload_kind)
    if not state.get("file_path") or not state.get("sheet_name") or not state.get("mapping"):
        messages.info(request, "Upload a file and finish the mapping step first.")
        return redirect("scorecard:upload_start", upload_kind=upload_kind)

    analysis = _load_preview_cache(state.get("preview_cache_path"))
    if analysis is None:
        if not _preview_scan_lock.acquire(blocking=False):
            messages.warning(
                request,
                "Another large preview is currently being prepared. Please wait a moment and try again.",
            )
            return redirect("scorecard:upload_mapping", upload_kind=upload_kind)
        try:
            analysis = _prepare_compact_preview_dataset(
                file_path=state["file_path"],
                sheet_name=state["sheet_name"],
                mapping=state["mapping"],
                config=config,
            )
        finally:
            _preview_scan_lock.release()
        _delete_temp_file(state.get("preview_cache_path"))
        state["preview_cache_path"] = _save_preview_cache(analysis)
        _store_state(request, upload_kind, state)

    mapping_summary = []
    field_lookup = {field.key: field for field in config["fields"]}
    for key, source in state["mapping"].items():
        field = field_lookup.get(key)
        if not field or not source:
            continue
        mapping_summary.append({"target": field.label, "source": source})

    return render(
        request,
        "uploads/preview.html",
        {
            "config": config,
            "upload_kind": upload_kind,
            "state": state,
            "analysis": analysis,
            "mapping_summary": mapping_summary,
        },
    )


@login_required
@require_http_methods(["GET"])
def upload_preview_errors_download_view(request: HttpRequest, upload_kind: str) -> HttpResponse:
    config = _get_import_config(upload_kind)
    state = _get_state(request, upload_kind)
    if not state.get("file_path") or not state.get("sheet_name") or not state.get("mapping"):
        messages.info(request, "Upload a file and finish the mapping step first.")
        return redirect("scorecard:upload_start", upload_kind=upload_kind)

    reason = (request.GET.get("reason") or "").strip()
    analysis = _load_preview_cache(state.get("preview_cache_path"))
    if analysis is None or analysis.get("cache_mode") == PREVIEW_CACHE_MODE_COMPACT:
        analysis = _prepare_import_dataset(
            file_path=state["file_path"],
            sheet_name=state["sheet_name"],
            mapping=state["mapping"],
            config=config,
            check_existing=False,
            preview_row_limit=0,
        )
        state["preview_cache_path"] = _save_preview_cache(analysis)
        _store_state(request, upload_kind, state)
    invalid_rows = analysis["invalid_export_rows"]
    if reason:
        invalid_rows = [row for row in invalid_rows if reason in row["errors"]]

    if not invalid_rows:
        messages.info(request, "There are no skipped customers matching that validation reason.")
        return redirect("scorecard:upload_preview", upload_kind=upload_kind)

    try:
        return _build_error_export_response(
            rows=invalid_rows,
            config=config,
            original_name=state.get("original_name", "import"),
            sheet_name=state.get("sheet_name", ""),
            reason=reason,
        )
    except ValueError as exc:
        messages.error(request, str(exc))
        return redirect("scorecard:upload_preview", upload_kind=upload_kind)


@login_required
@require_http_methods(["GET"])
def upload_preview_full_download_view(request: HttpRequest, upload_kind: str) -> HttpResponse:
    config = _get_import_config(upload_kind)
    state = _get_state(request, upload_kind)
    if not state.get("file_path") or not state.get("sheet_name") or not state.get("mapping"):
        messages.info(request, "Upload a file and finish the mapping step first.")
        return redirect("scorecard:upload_start", upload_kind=upload_kind)

    cached_export_path = state.get("preview_export_path")
    if cached_export_path and Path(cached_export_path).exists():
        stem = Path(state.get("original_name") or "import").stem
        return FileResponse(
            Path(cached_export_path).open("rb"),
            as_attachment=True,
            filename=f"{stem}_full_preview.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    analysis = _load_preview_cache(state.get("preview_cache_path"))
    if analysis is None or analysis.get("cache_mode") == PREVIEW_CACHE_MODE_COMPACT:
        analysis = _prepare_import_dataset(
            file_path=state["file_path"],
            sheet_name=state["sheet_name"],
            mapping=state["mapping"],
            config=config,
            check_existing=False,
            preview_row_limit=0,
        )
        state["preview_cache_path"] = _save_preview_cache(analysis)

    try:
        response, export_path = _build_full_preview_export_response(
            analysis=analysis,
            config=config,
            original_name=state.get("original_name", "import"),
            sheet_name=state.get("sheet_name", ""),
        )
        state["preview_export_path"] = export_path
        _store_state(request, upload_kind, state)
        return response
    except ValueError as exc:
        messages.error(request, str(exc))
        return redirect("scorecard:upload_preview", upload_kind=upload_kind)


@login_required
@require_http_methods(["GET"])
def upload_result_view(request: HttpRequest, upload_kind: str) -> HttpResponse:
    config = _get_import_config(upload_kind)
    result = _get_result_state(request, upload_kind)
    if not result:
        messages.info(request, "There is no recent import result to display.")
        return redirect("scorecard:upload_start", upload_kind=upload_kind)

    job_result = _load_upload_job_state(result.get("job_path"))
    if job_result:
        job_result["job_path"] = result.get("job_path")
        result = job_result
        _store_result_state(request, upload_kind, result)

    return render(
        request,
        "uploads/result.html",
        {
            "config": config,
            "upload_kind": upload_kind,
            "result": result,
        },
    )


@login_required
@require_http_methods(["POST"])
def upload_run_view(request: HttpRequest, upload_kind: str) -> HttpResponse:
    config = _get_import_config(upload_kind)
    model = config["model"]
    state = _get_state(request, upload_kind)
    if not state.get("file_path") or not state.get("sheet_name") or not state.get("mapping"):
        messages.info(request, "Upload a file and finish the mapping step first.")
        return redirect("scorecard:upload_start", upload_kind=upload_kind)

    prepared = _load_preview_cache(state.get("preview_cache_path"))
    if prepared and prepared.get("invalid_rows", 0) > 0:
        messages.error(
            request,
            "Import is disabled because the preview found validation issues. Correct the skipped rows before loading.",
        )
        return redirect("scorecard:upload_preview", upload_kind=upload_kind)

    file_name = state.get("original_name", "the uploaded file")
    job_path = _new_upload_job_path()
    running_state = {
        "status": "running",
        "job_path": job_path,
        "file_name": file_name,
        "sheet_name": state.get("sheet_name", ""),
        "target_table": config["table_name"],
        "inserted_rows": 0,
        "skipped_rows": 0,
        "processed_rows": 0,
        "total_rows": int((prepared or {}).get("valid_rows") or 0),
        "message": "The validated rows are being imported in the background.",
        "started_at": timezone.now().strftime("%Y-%m-%d %H:%M:%S"),
    }
    _save_upload_job_state(job_path, running_state)
    _store_result_state(request, upload_kind, running_state)
    request.session.pop(_session_key(upload_kind), None)
    request.session.modified = True

    log_upload_audit(
        request.user,
        "import_started",
        upload_kind=upload_kind,
        file_name=file_name,
        details=f"Background import started for target table: {config['table_name']}",
        object_id=config["table_name"],
    )
    if config.get("is_historical_scores"):
        log_historical_score_audit(
            request.user,
            "import_started",
            details=f"File: {file_name}; Background import started for target table: {config['table_name']}",
            object_id=config["table_name"],
        )

    worker = threading.Thread(
        target=_run_upload_import_job,
        kwargs={
            "job_path": job_path,
            "upload_kind": upload_kind,
            "state": dict(state),
            "user_id": request.user.pk,
        },
        name=f"scorecard-external-import-{upload_kind}-{request.user.pk}",
        daemon=True,
    )
    with _upload_job_threads_lock:
        _upload_job_threads[job_path] = worker
    worker.start()
    messages.info(request, "Import started in the background. This page will update automatically.")
    return redirect("scorecard:upload_result", upload_kind=upload_kind)

    file_name = state.get("original_name", "the uploaded file")
    started_at = timezone.now()
    log_upload_audit(
        request.user,
        "import_started",
        upload_kind=upload_kind,
        file_name=file_name,
        details=f"Immediate import started for target table: {config['table_name']}",
        object_id=config["table_name"],
    )
    if config.get("is_historical_scores"):
        log_historical_score_audit(
            request.user,
            "import_started",
            details=f"File: {file_name}; Immediate import started for target table: {config['table_name']}",
            object_id=config["table_name"],
        )

    try:
        metadata = _build_import_metadata_for_user(request.user, upload_kind, state, state["mapping"])
        prepared = _load_preview_cache(state.get("preview_cache_path"))
        compact_import = bool(
            prepared
            and prepared.get("cache_mode") == PREVIEW_CACHE_MODE_COMPACT
        )
        if prepared is None:
            prepared = _prepare_import_dataset(
                file_path=state["file_path"],
                sheet_name=state["sheet_name"],
                mapping=state["mapping"],
                config=config,
            )
        if prepared["invalid_rows"] > 0:
            messages.error(request, "Import is disabled because the preview found validation issues. Correct the skipped rows before loading.")
            return redirect("scorecard:upload_preview", upload_kind=upload_kind)

        inserted_rows = 0
        skipped_rows = prepared["invalid_rows"]
        skipped_preview = list(prepared["invalid_preview"])
        warning_summary = list(prepared.get("warning_summary", []))
        error_counter = Counter()
        for item in prepared["error_summary"]:
            error_counter[item["reason"]] = item["count"]

        batch: list[tuple[object, dict[str, object]]] = []
        now = timezone.now()
        if compact_import:
            clean_row_items = (
                cleaned_row
                for cleaned_row, errors, _row_data, _row_number in _iter_clean_rows(
                    state["file_path"],
                    state["sheet_name"],
                    state["mapping"],
                    config,
                )
                if cleaned_row is not None and not errors
            )
        else:
            clean_row_items = (
                item["cleaned_row"] for item in prepared["ready_rows"]
            )

        for cleaned_row in clean_row_items:
            obj = _build_row_payload(cleaned_row, request.user, metadata, now, model)
            batch.append((obj, cleaned_row))
            if len(batch) >= BULK_CREATE_BATCH_SIZE:
                inserted_rows, skipped_rows = _flush_batch(
                    model,
                    batch,
                    inserted_rows=inserted_rows,
                    skipped_rows=skipped_rows,
                    skipped_preview=skipped_preview,
                    error_counter=error_counter,
                )
                batch = []

        if batch:
            inserted_rows, skipped_rows = _flush_batch(
                model,
                batch,
                inserted_rows=inserted_rows,
                skipped_rows=skipped_rows,
                skipped_preview=skipped_preview,
                error_counter=error_counter,
            )

        result_state = {
            "status": "completed",
            "file_name": file_name,
            "sheet_name": state.get("sheet_name", ""),
            "target_table": config["table_name"],
            "inserted_rows": inserted_rows,
            "skipped_rows": skipped_rows,
            "error_summary": _build_error_summary(error_counter),
            "warning_summary": warning_summary,
            "skipped_preview": skipped_preview,
            "imported_at": timezone.now().strftime("%Y-%m-%d %H:%M:%S"),
            "started_at": started_at.strftime("%Y-%m-%d %H:%M:%S"),
            "message": "Import completed.",
        }
        import_details = f"Loaded rows: {inserted_rows}; Skipped rows: {skipped_rows}; Target table: {config['table_name']}"
        log_upload_audit(
            request.user,
            "import_complete",
            upload_kind=upload_kind,
            file_name=file_name,
            details=import_details,
            object_id=config["table_name"],
        )
        if config.get("is_historical_scores"):
            log_historical_score_audit(
                request.user,
                "import_complete",
                details=f"File: {file_name}; {import_details}",
                object_id=config["table_name"],
            )
        messages.success(request, "Import completed.")
    except Exception as exc:
        error_message = str(exc) or exc.__class__.__name__
        result_state = {
            "status": "failed",
            "file_name": file_name,
            "sheet_name": state.get("sheet_name", ""),
            "target_table": config["table_name"],
            "message": error_message,
            "inserted_rows": 0,
            "skipped_rows": 0,
            "error_summary": [{"reason": error_message, "count": 1}],
            "skipped_preview": [],
            "imported_at": timezone.now().strftime("%Y-%m-%d %H:%M:%S"),
            "started_at": started_at.strftime("%Y-%m-%d %H:%M:%S"),
        }
        log_upload_audit(
            request.user,
            "import_failed",
            upload_kind=upload_kind,
            file_name=file_name,
            details=error_message,
            object_id=config["table_name"],
        )
        if config.get("is_historical_scores"):
            log_historical_score_audit(
                request.user,
                "import_failed",
                details=f"File: {file_name}; {error_message}",
                object_id=config["table_name"],
            )
        messages.error(request, "Import could not be completed. Review the result details below.")
    finally:
        _delete_temp_file(state.get("file_path"))
        _clear_state(request, upload_kind)

    _store_result_state(request, upload_kind, result_state)
    return redirect("scorecard:upload_result", upload_kind=upload_kind)


@login_required
@require_http_methods(["POST"])
def upload_reset_view(request: HttpRequest, upload_kind: str) -> HttpResponse:
    _get_import_config(upload_kind)
    current_state = _get_state(request, upload_kind)
    current_result = _get_result_state(request, upload_kind)
    log_upload_audit(
        request.user,
        "reset_session",
        upload_kind=upload_kind,
        file_name=current_state.get("original_name", ""),
        details="Upload session and recent result were cleared.",
    )
    _clear_state(request, upload_kind)
    _clear_result_state(request, upload_kind)
    messages.info(request, "The current import session was cleared.")
    return redirect("scorecard:upload_start", upload_kind=upload_kind)
