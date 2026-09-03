from __future__ import annotations

import csv
from datetime import datetime
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.db.models import Avg, Count, Exists, OuterRef, Q
from django.http import HttpResponse, HttpResponseRedirect
from django.shortcuts import render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from scorecard.functions_view.audit import log_historical_score_audit
from scorecard.functions_view.main_customer_lookup import (
    current_branch_display_name,
    get_request_branch_names,
    is_all_branches_selected,
)
from scorecard.functions_view.historical_scores import build_historical_score_seed_rows, capture_historical_scores
from scorecard.models import CustomerLoan, CustomerOverdraft, HistoricalScore

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - dependency is installed in the project env
    Workbook = None
    Alignment = None
    Font = None
    PatternFill = None
    get_column_letter = None


HISTORICAL_SCORE_EXPORT_HEADERS = [
    "Reporting Date",
    "Branch",
    "Customer Code",
    "Customer Name",
    "Basel II Score",
    "Basel II Grade",
    "Basel Override Grade",
    "IFRS9 Score",
    "Captured",
]


def _clean(value: object) -> str:
    return str(value or "").strip()


def _historical_scores_base_queryset(request):
    branch_names = get_request_branch_names(request)
    scores = HistoricalScore.objects.all()
    if branch_names:
        scores = scores.filter(branch_name__in=branch_names)
    else:
        scores = scores.none()
    return scores, branch_names


def _annotate_historical_exposure_source(scores):
    loan_match = CustomerLoan.objects.filter(
        reporting_date=OuterRef("reporting_date"),
        customer_code=OuterRef("customer_id"),
    )
    overdraft_match = CustomerOverdraft.objects.filter(
        reporting_date=OuterRef("reporting_date"),
        customer_code=OuterRef("customer_id"),
    )
    return scores.annotate(
        has_active_loan=Exists(loan_match),
        has_active_overdraft=Exists(overdraft_match),
    )


def _apply_historical_score_filters(scores, params):
    selected_reporting_date = _clean(params.get("reporting_date"))
    selected_branch = _clean(params.get("branch"))
    selected_exposure_source = _clean(params.get("exposure_source"))
    search = _clean(params.get("search"))

    if selected_reporting_date:
        scores = scores.filter(reporting_date=selected_reporting_date)
    if selected_branch:
        scores = scores.filter(branch_name=selected_branch)
    if selected_exposure_source:
        scores = _annotate_historical_exposure_source(scores)
        if selected_exposure_source == "loan":
            scores = scores.filter(has_active_loan=True)
        elif selected_exposure_source == "overdraft":
            scores = scores.filter(has_active_overdraft=True)
        elif selected_exposure_source == "loan_or_overdraft":
            scores = scores.filter(Q(has_active_loan=True) | Q(has_active_overdraft=True))
        elif selected_exposure_source == "loan_only":
            scores = scores.filter(has_active_loan=True, has_active_overdraft=False)
        elif selected_exposure_source == "overdraft_only":
            scores = scores.filter(has_active_loan=False, has_active_overdraft=True)
    if search:
        scores = scores.filter(
            Q(customer_id__icontains=search)
            | Q(customer_name__icontains=search)
            | Q(branch_name__icontains=search)
            | Q(basel_ii_grade__icontains=search)
            | Q(basel_override_grade__icontains=search)
        )
    return scores, selected_reporting_date, selected_branch, selected_exposure_source, search


def _can_manage_historical_scores(user) -> bool:
    return (
        user.is_superuser
        or user.has_perm("scorecard.manage_scorecard_historical_scores")
    )


def _can_delete_historical_scores(user) -> bool:
    return _can_manage_historical_scores(user)


def _can_refresh_historical_scores(user) -> bool:
    return _can_manage_historical_scores(user)


def _coverage_totals_for_seed_rows(seed_rows):
    totals = {"both": 0, "basel_only": 0, "ifrs9_only": 0}
    for row in seed_rows:
        has_basel = (
            row.get("basel_ii_score") is not None
            or bool(row.get("basel_ii_grade"))
            or bool(row.get("basel_override_grade"))
        )
        has_ifrs9 = row.get("ifrs_9_score") is not None
        if has_basel and has_ifrs9:
            totals["both"] += 1
        elif has_basel:
            totals["basel_only"] += 1
        elif has_ifrs9:
            totals["ifrs9_only"] += 1
    return totals


def _format_score_percent(value):
    if value is None:
        return ""
    return f"{float(value):.2f}%"


def _historical_score_export_rows(scores):
    for score in scores.iterator(chunk_size=2000):
        yield [
            score.reporting_date.isoformat() if score.reporting_date else "",
            score.branch_name or "",
            score.customer_id or "",
            score.customer_name or "",
            _format_score_percent(score.basel_ii_score),
            score.basel_ii_grade or "",
            score.basel_override_grade or "",
            _format_score_percent(score.ifrs_9_score),
            timezone.localtime(score.created_at).strftime("%Y-%m-%d %H:%M") if score.created_at else "",
        ]


def _build_historical_scores_excel_response(filename, scores):
    if Workbook is None:
        return HttpResponse("Excel export requires openpyxl library. Please install it.", status=500)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Historical Scores"
    worksheet.append(HISTORICAL_SCORE_EXPORT_HEADERS)

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    header_font = Font(bold=True, color="0B2D52")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in _historical_score_export_rows(scores):
        worksheet.append(row)

    worksheet.freeze_panes = "A2"
    column_widths = {
        1: 16,
        2: 24,
        3: 18,
        4: 34,
        5: 18,
        6: 18,
        7: 22,
        8: 16,
        9: 18,
    }
    for index, width in column_widths.items():
        worksheet.column_dimensions[get_column_letter(index)].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@login_required
def historical_scores_list_view(request):
    """Read-only browser for scheduled historical score snapshots."""

    base_scores, branch_names = _historical_scores_base_queryset(request)

    reporting_dates = list(
        base_scores.order_by("-reporting_date")
        .values_list("reporting_date", flat=True)
        .distinct()
    )
    branch_options = list(
        base_scores.order_by("branch_name")
        .values_list("branch_name", flat=True)
        .distinct()
    )

    scores, selected_reporting_date, selected_branch, selected_exposure_source, search = _apply_historical_score_filters(base_scores, request.GET)

    summary = scores.aggregate(
        total_rows=Count("id"),
        total_customers=Count("customer_id", distinct=True),
        avg_basel_score=Avg("basel_ii_score"),
        avg_ifrs9_score=Avg("ifrs_9_score"),
    )

    latest_reporting_date = reporting_dates[0] if reporting_dates else None
    latest_count = (
        HistoricalScore.objects.filter(reporting_date=latest_reporting_date, branch_name__in=branch_names).count()
        if latest_reporting_date and branch_names
        else 0
    )

    scores = scores.order_by("-reporting_date", "branch_name", "customer_id")
    page_size = request.GET.get("page_size") or 25
    try:
        page_size = int(page_size)
    except (TypeError, ValueError):
        page_size = 25
    page_size = min(max(page_size, 10), 100)

    paginator = Paginator(scores, page_size)
    page_obj = paginator.get_page(request.GET.get("page"))

    query_params = request.GET.copy()
    query_params.pop("page", None)

    context = {
        "page_obj": page_obj,
        "scores": page_obj.object_list,
        "reporting_dates": reporting_dates,
        "branch_options": branch_options,
        "filters": {
            "reporting_date": selected_reporting_date,
            "branch": selected_branch,
            "exposure_source": selected_exposure_source,
            "search": search,
            "page_size": page_size,
        },
        "summary": {
            "total_rows": summary.get("total_rows") or 0,
            "total_customers": summary.get("total_customers") or 0,
            "avg_basel_score": summary.get("avg_basel_score"),
            "avg_ifrs9_score": summary.get("avg_ifrs9_score"),
            "latest_reporting_date": latest_reporting_date,
            "latest_count": latest_count,
        },
        "branch_scope_label": current_branch_display_name(request),
        "all_branches_selected": is_all_branches_selected(request),
        "query_string": query_params.urlencode(),
        "can_delete_historical_scores": _can_delete_historical_scores(request.user),
        "can_refresh_historical_scores": _can_refresh_historical_scores(request.user),
    }
    return render(request, "scorecard_historical_scores/historical_scores_list.html", context)


@login_required
@require_POST
def historical_scores_bulk_delete_view(request):
    """Delete all historical scores currently selected by the submitted filters."""

    if not _can_delete_historical_scores(request.user):
        raise PermissionDenied("You do not have permission to delete historical score snapshots.")

    scores, _branch_names = _historical_scores_base_queryset(request)
    scores, selected_reporting_date, selected_branch, selected_exposure_source, search = _apply_historical_score_filters(scores, request.POST)
    delete_count = scores.count()

    if delete_count:
        scores.delete()
        log_historical_score_audit(
            request.user,
            "bulk_delete",
            details=(
                f"Deleted {delete_count} historical score rows. "
                f"Reporting Date: {selected_reporting_date or 'All'}; "
                f"Branch: {selected_branch or current_branch_display_name(request) or 'Visible branches'}; "
                f"Exposure: {selected_exposure_source or 'All'}; "
                f"Search: {search or 'None'}"
            ),
            object_id=f"historical_scores_bulk_delete:{timezone.now().strftime('%Y%m%d_%H%M%S')}",
            branch_name=selected_branch or current_branch_display_name(request) or "",
        )
        messages.success(request, f"Deleted {delete_count} historical score row(s) matching the current filters.")
    else:
        messages.info(request, "No historical score rows matched the selected filters.")

    redirect_params = request.POST.copy()
    redirect_params.pop("csrfmiddlewaretoken", None)
    redirect_params.pop("page", None)
    url = reverse("scorecard:historical_scores_list")
    query_string = redirect_params.urlencode()
    return HttpResponseRedirect(f"{url}?{query_string}" if query_string else url)


@login_required
@require_POST
def historical_scores_refresh_view(request):
    """Rebuild one selected reporting date from current scores and active loan/overdraft exposure."""

    if not _can_refresh_historical_scores(request.user):
        raise PermissionDenied("You do not have permission to refresh historical score snapshots.")

    selected_reporting_date = _clean(request.POST.get("refresh_reporting_date") or request.POST.get("reporting_date"))
    selected_branch = _clean(request.POST.get("branch"))
    if not selected_reporting_date:
        messages.error(request, "Choose or enter a reporting date before refreshing historical scores.")
        return HttpResponseRedirect(reverse("scorecard:historical_scores_list"))

    try:
        reporting_date = datetime.strptime(selected_reporting_date, "%Y-%m-%d").date()
    except ValueError:
        messages.error(request, "The selected reporting date is invalid.")
        return HttpResponseRedirect(reverse("scorecard:historical_scores_list"))

    _base_scores, branch_names = _historical_scores_base_queryset(request)
    refresh_branch_names = list(branch_names)
    if selected_branch:
        if selected_branch not in refresh_branch_names:
            raise PermissionDenied("You do not have permission to refresh this branch.")
        refresh_branch_names = [selected_branch]

    seed_rows, _coverage_totals = build_historical_score_seed_rows(reporting_date)
    seed_rows = [row for row in seed_rows if row.get("branch_name") in refresh_branch_names]
    coverage_totals = _coverage_totals_for_seed_rows(seed_rows)

    delete_query = HistoricalScore.objects.filter(reporting_date=reporting_date, branch_name__in=refresh_branch_names)
    deleted_count = delete_query.count()
    delete_query.delete()

    result = capture_historical_scores(
        reporting_date,
        seed_rows=seed_rows,
        coverage_totals=coverage_totals,
    )

    log_historical_score_audit(
        request.user,
        "refresh",
        details=(
            f"Refreshed historical scores for {selected_reporting_date}. "
            f"Branch: {selected_branch or current_branch_display_name(request) or 'Visible branches'}; "
            f"Deleted={deleted_count}; Created={result.created}; Updated={result.updated}."
        ),
        object_id=f"historical_scores_refresh:{selected_reporting_date}:{timezone.now().strftime('%Y%m%d_%H%M%S')}",
        branch_name=selected_branch or current_branch_display_name(request) or "",
    )

    if result.total_rows:
        messages.success(
            request,
            f"Refreshed {result.total_rows} historical score row(s) for {selected_reporting_date}.",
        )
    else:
        messages.warning(
            request,
            f"No active loan or overdraft score rows were found for {selected_reporting_date}; existing rows in scope were removed.",
        )

    redirect_params = request.POST.copy()
    redirect_params.pop("csrfmiddlewaretoken", None)
    redirect_params.pop("page", None)
    redirect_params.pop("refresh_reporting_date", None)
    redirect_params["reporting_date"] = selected_reporting_date
    url = reverse("scorecard:historical_scores_list")
    query_string = redirect_params.urlencode()
    return HttpResponseRedirect(f"{url}?{query_string}" if query_string else url)


@login_required
def historical_scores_download_view(request):
    """Download the filtered historical score snapshots as CSV or Excel."""

    scores, _branch_names = _historical_scores_base_queryset(request)
    scores, selected_reporting_date, selected_branch, _selected_exposure_source, _search = _apply_historical_score_filters(scores, request.GET)
    scores = scores.order_by("-reporting_date", "branch_name", "customer_id")
    row_count = scores.count()
    export_format = _clean(request.GET.get("format")).lower()
    is_excel = export_format in {"excel", "xlsx"}

    timestamp = timezone.now().strftime("%Y%m%d_%H%M%S")
    date_part = selected_reporting_date or "all_dates"
    branch_part = (selected_branch or "visible_branches").replace(" ", "_").replace("/", "_")
    extension = "xlsx" if is_excel else "csv"
    filename = f"scorecard_historical_scores_{date_part}_{branch_part}_{timestamp}.{extension}"

    log_historical_score_audit(
        request.user,
        "download",
        details=(
            f"Downloaded historical score {extension.upper()} extract. Reporting Date: {selected_reporting_date or 'All'}; "
            f"Branch: {selected_branch or current_branch_display_name(request) or 'Visible branches'}; "
            f"Rows: {row_count}"
        ),
        object_id=f"historical_scores_download:{timestamp}",
        branch_name=selected_branch or current_branch_display_name(request) or "",
    )

    if is_excel:
        return _build_historical_scores_excel_response(filename, scores)

    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    writer = csv.writer(response)
    writer.writerow(HISTORICAL_SCORE_EXPORT_HEADERS)
    for row in _historical_score_export_rows(scores):
        writer.writerow(row)
    return response
