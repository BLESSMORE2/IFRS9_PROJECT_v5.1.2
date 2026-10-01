from __future__ import annotations

from collections import Counter, defaultdict, deque
from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any
from urllib.parse import quote_plus

import openpyxl
from django.contrib import messages
from django.core.paginator import Paginator
from django.http import HttpRequest, HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from scorecard.functions_view.audit import log_ifrs9_results_audit
from scorecard.functions_view.main_customer_lookup import (
    current_branch_display_name,
    get_request_branch_names,
    get_request_branch_scope,
    is_all_branches_selected,
)
from scorecard.functions_view.validation_export_cache import (
    cache_validation_payload,
    cache_validation_export,
    get_cached_validation_payload,
    get_cached_validation_export,
    validation_payload_cache_key,
    validation_export_cache_key,
)
from scorecard.models import HistoricalScore


MOVEMENT_OPTIONS = (
    ("", "All movements"),
    ("increased", "Score increased"),
    ("decreased", "Score decreased"),
    ("unchanged", "Score unchanged"),
    ("new", "New this month"),
    ("exited", "Exited since prior month"),
)
VALID_MOVEMENT_FILTERS = {value for value, _label in MOVEMENT_OPTIONS}
PAGE_SIZE_OPTIONS = (25, 50, 100)
SCORE_BANDS = (
    ("0% - 20%", Decimal("0"), Decimal("20")),
    (">20% - 40%", Decimal("20"), Decimal("40")),
    (">40% - 60%", Decimal("40"), Decimal("60")),
    (">60% - 80%", Decimal("60"), Decimal("80")),
    (">80% - 100%", Decimal("80"), Decimal("100")),
)


def _clean(value: Any) -> str:
    return str(value or "").strip()


def _parse_date(value: str | None) -> date | None:
    try:
        return date.fromisoformat(_clean(value)) if _clean(value) else None
    except ValueError:
        return None


def _optional_decimal(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value))
    except Exception:
        return None


def _average(values: list[Decimal]) -> float | None:
    return float(sum(values) / len(values)) if values else None


def _historical_scope_queryset(request: HttpRequest, selected_branch: str = ""):
    branch_names = get_request_branch_names(request)
    queryset = (
        HistoricalScore.objects.filter(branch_name__in=branch_names, ifrs_9_score__isnull=False)
        if branch_names
        else HistoricalScore.objects.none()
    )
    if selected_branch:
        queryset = queryset.filter(branch_name__iexact=selected_branch)
    return queryset, branch_names


def _available_historical_dates(queryset) -> list[date]:
    return list(queryset.order_by("-reporting_date").values_list("reporting_date", flat=True).distinct())


def _resolve_comparison_dates(
    available_dates: list[date],
    current_value: str | None,
    previous_value: str | None,
) -> tuple[date | None, date | None]:
    current_date = _parse_date(current_value)
    if current_date not in available_dates:
        current_date = available_dates[0] if available_dates else None

    previous_date = _parse_date(previous_value)
    if previous_date not in available_dates or current_date is None or previous_date >= current_date:
        previous_date = next((item for item in available_dates if current_date and item < current_date), None)
    return current_date, previous_date


def _historical_values_for_dates(
    queryset,
    previous_date: date,
    current_date: date,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    rows_by_date = {previous_date: [], current_date: []}
    rows = (
        queryset.filter(reporting_date__in=(previous_date, current_date))
        .values("reporting_date", "branch_name", "customer_id", "customer_name", "ifrs_9_score")
        .order_by("reporting_date", "branch_name", "customer_id")
    )
    for row in rows:
        reporting_date = row.pop("reporting_date")
        rows_by_date[reporting_date].append(row)
    return rows_by_date[previous_date], rows_by_date[current_date]


def _build_detail_row(
    previous_row: dict[str, Any] | None,
    current_row: dict[str, Any] | None,
) -> dict[str, Any]:
    previous_score = _optional_decimal(previous_row.get("ifrs_9_score")) if previous_row else None
    current_score = _optional_decimal(current_row.get("ifrs_9_score")) if current_row else None
    if previous_row is None:
        movement_key, movement_label = "new", "New this month"
    elif current_row is None:
        movement_key, movement_label = "exited", "Exited since prior month"
    elif current_score > previous_score:
        movement_key, movement_label = "increased", "Score increased"
    elif current_score < previous_score:
        movement_key, movement_label = "decreased", "Score decreased"
    else:
        movement_key, movement_label = "unchanged", "Score unchanged"

    score_delta = None
    if previous_score is not None and current_score is not None:
        score_delta = current_score - previous_score

    previous_branch = _clean(previous_row.get("branch_name")) if previous_row else ""
    current_branch = _clean(current_row.get("branch_name")) if current_row else ""
    source_row = current_row or previous_row or {}
    return {
        "customer_id": _clean(source_row.get("customer_id")),
        "customer_name": _clean(source_row.get("customer_name")),
        "branch_name": current_branch or previous_branch,
        "previous_branch": previous_branch,
        "current_branch": current_branch,
        "branch_transfer": bool(previous_branch and current_branch and previous_branch.casefold() != current_branch.casefold()),
        "previous_score": previous_score,
        "current_score": current_score,
        "score_delta": score_delta,
        "movement_key": movement_key,
        "movement_label": movement_label,
        "has_previous": previous_row is not None,
        "has_current": current_row is not None,
    }


def _pair_historical_rows(
    previous_rows: list[dict[str, Any]],
    current_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    previous_by_customer: dict[str, list[dict[str, Any]]] = defaultdict(list)
    current_by_customer: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in previous_rows:
        previous_by_customer[_clean(row.get("customer_id")).casefold()].append(row)
    for row in current_rows:
        current_by_customer[_clean(row.get("customer_id")).casefold()].append(row)

    details = []
    for customer_key in sorted(set(previous_by_customer) | set(current_by_customer)):
        previous_remaining = list(previous_by_customer.get(customer_key, []))
        current_remaining = list(current_by_customer.get(customer_key, []))
        previous_by_branch: dict[str, deque] = defaultdict(deque)
        for row in previous_remaining:
            previous_by_branch[_clean(row.get("branch_name")).casefold()].append(row)

        matched_previous_ids: set[int] = set()
        unmatched_current = []
        for current_row in current_remaining:
            branch_key = _clean(current_row.get("branch_name")).casefold()
            if previous_by_branch[branch_key]:
                previous_row = previous_by_branch[branch_key].popleft()
                matched_previous_ids.add(id(previous_row))
                details.append(_build_detail_row(previous_row, current_row))
            else:
                unmatched_current.append(current_row)

        unmatched_previous = [row for row in previous_remaining if id(row) not in matched_previous_ids]
        pair_count = min(len(unmatched_previous), len(unmatched_current))
        for index in range(pair_count):
            details.append(_build_detail_row(unmatched_previous[index], unmatched_current[index]))
        for current_row in unmatched_current[pair_count:]:
            details.append(_build_detail_row(None, current_row))
        for previous_row in unmatched_previous[pair_count:]:
            details.append(_build_detail_row(previous_row, None))
    return details


def _movement_summary(details: list[dict[str, Any]]) -> dict[str, Any]:
    movement_counts = Counter(row["movement_key"] for row in details)
    previous_scores = [row["previous_score"] for row in details if row["previous_score"] is not None]
    current_scores = [row["current_score"] for row in details if row["current_score"] is not None]
    score_deltas = [row["score_delta"] for row in details if row["score_delta"] is not None]
    return {
        "previous_customers": sum(1 for row in details if row["has_previous"]),
        "current_customers": sum(1 for row in details if row["has_current"]),
        "matched_customers": sum(1 for row in details if row["has_previous"] and row["has_current"]),
        "increased": movement_counts["increased"],
        "decreased": movement_counts["decreased"],
        "unchanged": movement_counts["unchanged"],
        "new": movement_counts["new"],
        "exited": movement_counts["exited"],
        "branch_transfers": sum(1 for row in details if row["branch_transfer"]),
        "average_previous_score": _average(previous_scores),
        "average_current_score": _average(current_scores),
        "average_score_delta": _average(score_deltas),
    }


def _branch_summary(details: list[dict[str, Any]]) -> list[dict[str, Any]]:
    buckets: dict[str, dict[str, Any]] = {}

    def bucket_for(branch_name: str) -> dict[str, Any]:
        key = _clean(branch_name).casefold()
        if key not in buckets:
            buckets[key] = {
                "branch_name": _clean(branch_name) or "Unassigned",
                "previous_customers": 0,
                "current_customers": 0,
                "matched_customers": 0,
                "increased": 0,
                "decreased": 0,
                "unchanged": 0,
                "new": 0,
                "exited": 0,
                "score_deltas": [],
            }
        return buckets[key]

    for detail in details:
        if detail["has_previous"]:
            bucket_for(detail["previous_branch"])["previous_customers"] += 1
        if detail["has_current"]:
            bucket_for(detail["current_branch"])["current_customers"] += 1
        bucket = bucket_for(detail["current_branch"] or detail["previous_branch"])
        if detail["has_previous"] and detail["has_current"]:
            bucket["matched_customers"] += 1
        bucket[detail["movement_key"]] += 1
        if detail["score_delta"] is not None:
            bucket["score_deltas"].append(detail["score_delta"])

    rows = []
    for bucket in sorted(buckets.values(), key=lambda row: row["branch_name"].casefold()):
        score_deltas = bucket.pop("score_deltas")
        bucket["average_score_delta"] = _average(score_deltas)
        rows.append(bucket)
    return rows


def _score_band(score: Decimal) -> str:
    if score < 0 or score > 100:
        return "Outside 0% - 100%"
    for index, (label, lower, upper) in enumerate(SCORE_BANDS):
        if index == 0 and lower <= score <= upper:
            return label
        if lower < score <= upper:
            return label
    return "Outside 0% - 100%"


def _score_distribution(
    previous_rows: list[dict[str, Any]],
    current_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    previous_counts = Counter(_score_band(_optional_decimal(row["ifrs_9_score"])) for row in previous_rows)
    current_counts = Counter(_score_band(_optional_decimal(row["ifrs_9_score"])) for row in current_rows)
    labels = [item[0] for item in SCORE_BANDS] + ["Outside 0% - 100%"]
    return [
        {
            "band": label,
            "previous_count": previous_counts[label],
            "current_count": current_counts[label],
            "change": current_counts[label] - previous_counts[label],
        }
        for label in labels
        if previous_counts[label] or current_counts[label]
    ]


def _duplicate_customer_rows(
    historical_rows: list[dict[str, Any]],
    reporting_date: date,
    period_label: str,
) -> list[dict[str, Any]]:
    rows_by_customer: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in historical_rows:
        customer_id = _clean(row.get("customer_id"))
        if customer_id:
            rows_by_customer[customer_id.casefold()].append(row)

    duplicate_rows = []
    for customer_rows in rows_by_customer.values():
        if len(customer_rows) <= 1:
            continue
        ordered_rows = sorted(customer_rows, key=lambda row: _clean(row.get("branch_name")).casefold())
        occurrence_count = len(ordered_rows)
        for occurrence_number, row in enumerate(ordered_rows, start=1):
            duplicate_rows.append(
                {
                    "period_label": period_label,
                    "reporting_date": reporting_date,
                    "customer_id": _clean(row.get("customer_id")),
                    "customer_name": _clean(row.get("customer_name")),
                    "branch_name": _clean(row.get("branch_name")) or "Unassigned",
                    "occurrence_number": occurrence_number,
                    "occurrence_count": occurrence_count,
                    "ifrs9_score": _optional_decimal(row.get("ifrs_9_score")),
                }
            )
    return sorted(
        duplicate_rows,
        key=lambda row: (row["reporting_date"], row["customer_id"].casefold(), row["branch_name"].casefold()),
    )


def _filter_details(details: list[dict[str, Any]], search: str, movement_filter: str) -> list[dict[str, Any]]:
    filtered = details
    if movement_filter:
        filtered = [row for row in filtered if row["movement_key"] == movement_filter]
    search_key = _clean(search).casefold()
    if search_key:
        filtered = [
            row
            for row in filtered
            if search_key
            in " ".join(
                [row["customer_id"], row["customer_name"], row["branch_name"], row["movement_label"]]
            ).casefold()
        ]
    return filtered


def _build_validation_payload(historical_queryset, previous_date: date, current_date: date) -> dict[str, Any]:
    previous_rows, current_rows = _historical_values_for_dates(
        historical_queryset,
        previous_date,
        current_date,
    )
    details = _pair_historical_rows(previous_rows, current_rows)
    previous_duplicates = _duplicate_customer_rows(previous_rows, previous_date, "Previous")
    current_duplicates = _duplicate_customer_rows(current_rows, current_date, "Current")
    score_distribution = _score_distribution(previous_rows, current_rows)
    return {
        "previous_date": previous_date,
        "current_date": current_date,
        "details": details,
        "summary": _movement_summary(details),
        "branch_summary": _branch_summary(details),
        "score_distribution": score_distribution,
        "score_distribution_totals": {
            "previous_count": sum(row["previous_count"] for row in score_distribution),
            "current_count": sum(row["current_count"] for row in score_distribution),
            "change": sum(row["change"] for row in score_distribution),
        },
        "duplicate_customers": previous_duplicates + current_duplicates,
        "data_quality": {
            "previous_duplicate_customer_codes": len({row["customer_id"].casefold() for row in previous_duplicates}),
            "current_duplicate_customer_codes": len({row["customer_id"].casefold() for row in current_duplicates}),
            "branch_transfers": sum(1 for row in details if row["branch_transfer"]),
            "new_customers": sum(1 for row in details if row["movement_key"] == "new"),
            "exited_customers": sum(1 for row in details if row["movement_key"] == "exited"),
        },
    }


def _get_validation_payload(
    historical_queryset,
    previous_date: date,
    current_date: date,
    branch_names: list[str],
    selected_branch: str,
) -> dict[str, Any]:
    cache_key = validation_payload_cache_key(
        report_name="ifrs9",
        queryset=historical_queryset,
        previous_date=previous_date,
        current_date=current_date,
        branch_names=branch_names,
        selected_branch=selected_branch,
    )
    cached_payload = get_cached_validation_payload(cache_key)
    if cached_payload is not None:
        return cached_payload

    payload = _build_validation_payload(historical_queryset, previous_date, current_date)
    cache_validation_payload(cache_key, payload)
    return payload


def _style_title(worksheet, title: str, total_columns: int) -> None:
    dark_fill = PatternFill("solid", fgColor="0B2D52")
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(total_columns, 8))
    worksheet["A1"] = title
    worksheet["A1"].fill = dark_fill
    worksheet["A1"].font = Font(name="Arial", color="FFFFFF", bold=True, size=15)
    worksheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[1].height = 27
    for column in range(2, max(total_columns, 8) + 1):
        worksheet.cell(1, column).fill = dark_fill


def _style_header_row(worksheet, row_number: int, column_count: int) -> None:
    fill = PatternFill("solid", fgColor="0E4377")
    border = Border(bottom=Side(style="thin", color="CAD8E6"))
    for column in range(1, column_count + 1):
        cell = worksheet.cell(row_number, column)
        cell.fill = fill
        cell.font = Font(name="Arial", color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    worksheet.row_dimensions[row_number].height = 28


def _fit_columns(worksheet, max_width: int = 34) -> None:
    for column in range(1, worksheet.max_column + 1):
        width = 12
        for row in range(1, min(worksheet.max_row, 250) + 1):
            value = worksheet.cell(row, column).value
            if value not in (None, ""):
                width = max(width, min(len(str(value)) + 2, max_width))
        worksheet.column_dimensions[get_column_letter(column)].width = width


def _add_table(worksheet, name: str, header_row: int) -> None:
    if worksheet.max_row <= header_row:
        return
    last_column = max(
        column
        for column in range(1, worksheet.max_column + 1)
        if worksheet.cell(header_row, column).value not in (None, "")
    )
    table = Table(
        displayName=name,
        ref=f"A{header_row}:{get_column_letter(last_column)}{worksheet.max_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def _build_validation_workbook(
    payload: dict[str, Any],
    detail_rows: list[dict[str, Any]],
    branch_scope_label: str,
    selected_branch: str,
    movement_filter: str,
    search: str,
):
    workbook = openpyxl.Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Validation Summary"
    detail_sheet = workbook.create_sheet("Individual Customers")
    duplicate_sheet = workbook.create_sheet("Duplicate Customers")
    distribution_sheet = workbook.create_sheet("Score Population")
    definitions_sheet = workbook.create_sheet("Definitions")
    for worksheet in workbook.worksheets:
        worksheet.sheet_view.showGridLines = False

    summary = payload["summary"]
    _style_title(summary_sheet, "IFRS9 Scores Validation", 11)
    summary_sheet.append([
        "Previous Month-End", payload["previous_date"], "Current Month-End", payload["current_date"],
        "Branch Scope", selected_branch or branch_scope_label, "Generated", timezone.localtime().replace(tzinfo=None),
    ])
    summary_sheet.append([])
    summary_sheet.append(["Validation Metric", "Value"])
    _style_header_row(summary_sheet, 4, 2)
    metrics = [
        ("Previous customers", summary["previous_customers"]),
        ("Current customers", summary["current_customers"]),
        ("Matched customers", summary["matched_customers"]),
        ("Scores increased", summary["increased"]),
        ("Scores decreased", summary["decreased"]),
        ("Scores unchanged", summary["unchanged"]),
        ("New this month", summary["new"]),
        ("Exited since prior month", summary["exited"]),
        ("Branch transfers", summary["branch_transfers"]),
        ("Average previous IFRS9 score", summary["average_previous_score"]),
        ("Average current IFRS9 score", summary["average_current_score"]),
        ("Average score change (pp)", summary["average_score_delta"]),
    ]
    for metric, value in metrics:
        summary_sheet.append([metric, value])
        if "score" in metric.casefold():
            summary_sheet.cell(summary_sheet.max_row, 2).number_format = (
                '+0.00" pp";-0.00" pp";0.00" pp"' if "(pp)" in metric else '0.00"%"'
            )

    branch_header_row = summary_sheet.max_row + 2
    summary_sheet.cell(branch_header_row, 1, "Branch Validation Summary")
    summary_sheet.cell(branch_header_row, 1).font = Font(name="Arial", bold=True, size=12, color="0B2D52")
    branch_columns = [
        "Branch",
        f"Previous Population ({payload['previous_date']})",
        f"Current Population ({payload['current_date']})",
        "Matched Customers", "Increased", "Decreased",
        "Unchanged", "New", "Exited", "Average Score Change (pp)",
    ]
    summary_sheet.append(branch_columns)
    branch_table_header = summary_sheet.max_row
    _style_header_row(summary_sheet, branch_table_header, len(branch_columns))
    for row in payload["branch_summary"]:
        summary_sheet.append([
            row["branch_name"], row["previous_customers"], row["current_customers"],
            row["matched_customers"], row["increased"], row["decreased"], row["unchanged"],
            row["new"], row["exited"], row["average_score_delta"],
        ])
    for row in summary_sheet.iter_rows(min_row=branch_table_header + 1, min_col=10, max_col=10):
        row[0].number_format = '+0.00" pp";-0.00" pp";0.00" pp"'
    summary_sheet.freeze_panes = "A5"
    _add_table(summary_sheet, "IFRS9ValidationBranchSummary", branch_table_header)

    detail_columns = [
        "Branch", "Customer Code", "Customer Name",
        f"Previous IFRS9 Score ({payload['previous_date']})",
        f"Current IFRS9 Score ({payload['current_date']})",
        "Score Change (pp)", "Movement", "Previous Branch", "Branch Transfer",
    ]
    _style_title(detail_sheet, "Individual Customer IFRS9 Score Changes", len(detail_columns))
    detail_sheet.append([
        "Previous Month-End", payload["previous_date"], "Current Month-End", payload["current_date"],
        "Branch Scope", selected_branch or branch_scope_label, "Movement Filter", movement_filter or "All",
        "Search", search or "All customers",
    ])
    detail_sheet.append([])
    detail_sheet.append(detail_columns)
    _style_header_row(detail_sheet, 4, len(detail_columns))
    for row in detail_rows:
        detail_sheet.append([
            row["branch_name"], row["customer_id"], row["customer_name"], row["previous_score"],
            row["current_score"], row["score_delta"], row["movement_label"], row["previous_branch"],
            "Yes" if row["branch_transfer"] else "No",
        ])
    for row in detail_sheet.iter_rows(min_row=5, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = '0.00"%"' if cell.column in (4, 5) else '+0.00" pp";-0.00" pp";0.00" pp"'
    if detail_sheet.max_row >= 5:
        movement_range = f"G5:G{detail_sheet.max_row}"
        for movement, fill_color, font_color in (
            ("Score increased", "E6F7EF", "087443"),
            ("Score decreased", "FFF0F1", "B4232C"),
            ("New this month", "EAF4FF", "145EA8"),
            ("Exited since prior month", "FFF4DF", "8C4B00"),
        ):
            detail_sheet.conditional_formatting.add(
                movement_range,
                FormulaRule(
                    formula=[f'$G5="{movement}"'],
                    fill=PatternFill("solid", fgColor=fill_color),
                    font=Font(color=font_color, bold=True),
                ),
            )
    detail_sheet.freeze_panes = "A5"
    detail_sheet.sheet_properties.tabColor = "176BB3"
    _add_table(detail_sheet, "IFRS9IndividualCustomers", 4)

    duplicate_columns = [
        "Comparison Period", "Reporting Date", "Customer Code", "Customer Name", "Branch",
        "Occurrence", "Duplicate Count", "IFRS9 Score",
    ]
    _style_title(duplicate_sheet, "Historical IFRS9 Duplicate Customer Details", len(duplicate_columns))
    duplicate_sheet.append([
        "Definition", "A duplicate is a customer code appearing in more than one IFRS9 historical row for the same reporting date.",
        "Previous Month-End", payload["previous_date"], "Current Month-End", payload["current_date"],
    ])
    duplicate_sheet.append([])
    duplicate_sheet.append(duplicate_columns)
    _style_header_row(duplicate_sheet, 4, len(duplicate_columns))
    for row in payload["duplicate_customers"]:
        duplicate_sheet.append([
            row["period_label"], row["reporting_date"], row["customer_id"], row["customer_name"],
            row["branch_name"], f"{row['occurrence_number']} of {row['occurrence_count']}",
            row["occurrence_count"], row["ifrs9_score"],
        ])
    if payload["duplicate_customers"]:
        for row in duplicate_sheet.iter_rows(min_row=5, min_col=8, max_col=8):
            row[0].number_format = '0.00"%"'
        duplicate_sheet.conditional_formatting.add(
            f"G5:G{duplicate_sheet.max_row}",
            FormulaRule(
                formula=["$G5>1"],
                fill=PatternFill("solid", fgColor="FFF0F1"),
                font=Font(color="B4232C", bold=True),
            ),
        )
        _add_table(duplicate_sheet, "IFRS9DuplicateCustomers", 4)
    else:
        duplicate_sheet.append(["No duplicate IFRS9 customer codes were found for the selected comparison dates."])
    duplicate_sheet.freeze_panes = "A5"
    duplicate_sheet.sheet_properties.tabColor = "E5484D"

    distribution_columns = [
        "Analytical Score Range",
        f"Previous Population ({payload['previous_date']})",
        f"Current Population ({payload['current_date']})",
        "Population Change",
    ]
    _style_title(distribution_sheet, "IFRS9 Score Population by Date", len(distribution_columns))
    distribution_sheet.append([
        "Previous Month-End", payload["previous_date"], "Current Month-End", payload["current_date"],
        "Note", "These are analytical ranges only and are not IFRS9 stages or regulatory grade bands.",
    ])
    distribution_sheet.append([])
    distribution_sheet.append(distribution_columns)
    _style_header_row(distribution_sheet, 4, len(distribution_columns))
    for row in payload["score_distribution"]:
        distribution_sheet.append([row["band"], row["previous_count"], row["current_count"], row["change"]])
    distribution_totals = payload["score_distribution_totals"]
    distribution_sheet.append([
        "Total score population",
        distribution_totals["previous_count"],
        distribution_totals["current_count"],
        distribution_totals["change"],
    ])
    for cell in distribution_sheet[distribution_sheet.max_row]:
        cell.fill = PatternFill("solid", fgColor="EAF3FB")
        cell.font = Font(name="Arial", color="0B2D52", bold=True)
    distribution_sheet.freeze_panes = "A5"
    _add_table(distribution_sheet, "IFRS9ScoreDistribution", 4)

    _style_title(definitions_sheet, "IFRS9 Scores Validation Definitions and Sources", 4)
    definitions_sheet.append(["Item", "Definition", "Source", "Notes"])
    _style_header_row(definitions_sheet, 2, 4)
    definitions = [
        ("IFRS9 score", "The IFRS9 score frozen for a customer at the selected historical month-end.", "SCORECARD_HISTORICAL_SCORES", "Only rows with a historical IFRS9 score are included."),
        ("Score movement", "Current historical IFRS9 score minus the previous historical IFRS9 score.", "SCORECARD_HISTORICAL_SCORES", "Reported in percentage points without interpreting an increase or decrease as a regulatory stage movement."),
        ("New / exited", "Present only in the current month-end / present only in the previous month-end.", "SCORECARD_HISTORICAL_SCORES", "Matching uses customer code and prefers the same branch."),
        ("Duplicate customer", "The same customer code appears in more than one historical IFRS9 row for a reporting date.", "SCORECARD_HISTORICAL_SCORES", "Every occurrence is listed in the Duplicate Customers sheet."),
        ("Score population by date", "Each selected date's complete IFRS9-scored population grouped into equal 20-point analytical ranges.", "SCORECARD_HISTORICAL_SCORES", "These ranges are not IFRS9 stages or regulatory grade bands; movement metrics use matched customers only."),
        ("Branch scope", "Only branches assigned to the user and selected in the scorecard workspace.", "Scorecard branch access", "A selected branch narrows every workbook sheet."),
    ]
    for row in definitions:
        definitions_sheet.append(row)
    definitions_sheet.freeze_panes = "A3"
    _add_table(definitions_sheet, "IFRS9ValidationDefinitions", 2)

    for worksheet in workbook.worksheets:
        _fit_columns(worksheet)
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_margins.left = 0.25
        worksheet.page_margins.right = 0.25
        worksheet.page_margins.top = 0.5
        worksheet.page_margins.bottom = 0.5
    return workbook


def _validation_query_string(request: HttpRequest) -> str:
    keys = ("current_date", "previous_date", "branch", "movement", "search", "page_size")
    return "&".join(
        f"{key}={quote_plus(_clean(request.GET.get(key)))}"
        for key in keys
        if _clean(request.GET.get(key))
    )


def ifrs9_validations_view(request: HttpRequest):
    branch_scope = get_request_branch_scope(request)
    branch_names = [name for name in get_request_branch_names(request) if name]
    selected_branch = _clean(request.GET.get("branch"))
    if selected_branch and selected_branch.casefold() not in {name.casefold() for name in branch_names}:
        messages.error(request, "Choose a branch within your assigned branch scope.")
        selected_branch = ""

    historical_queryset, _branch_names = _historical_scope_queryset(request, selected_branch)
    available_dates = _available_historical_dates(historical_queryset)
    current_date, previous_date = _resolve_comparison_dates(
        available_dates,
        request.GET.get("current_date"),
        request.GET.get("previous_date"),
    )
    movement_filter = _clean(request.GET.get("movement"))
    if movement_filter not in VALID_MOVEMENT_FILTERS:
        movement_filter = ""
    search = _clean(request.GET.get("search"))
    try:
        page_size = int(request.GET.get("page_size") or 50)
    except (TypeError, ValueError):
        page_size = 50
    if page_size not in PAGE_SIZE_OPTIONS:
        page_size = 50

    payload = None
    filtered_details = []
    page_obj = None
    if current_date and previous_date:
        payload = _get_validation_payload(
            historical_queryset,
            previous_date,
            current_date,
            branch_names,
            selected_branch,
        )
        filtered_details = _filter_details(payload["details"], search, movement_filter)
        page_obj = Paginator(filtered_details, page_size).get_page(request.GET.get("page") or 1)

    return render(
        request,
        "credit_scoreshifts/ifrs9_validations.html",
        {
            "results_page_key": "ifrs9_validations",
            "results_branch_scope_label": current_branch_display_name(request) or "No branch selected",
            "results_branch_mode_is_all": is_all_branches_selected(request),
            "results_branch_scope_count": len(branch_scope),
            "validation_branch_names": branch_names,
            "validation_selected_branch": selected_branch,
            "validation_available_dates": available_dates,
            "validation_current_date": current_date,
            "validation_previous_date": previous_date,
            "validation_payload": payload,
            "validation_movement_options": MOVEMENT_OPTIONS,
            "validation_movement_filter": movement_filter,
            "validation_search": search,
            "validation_page_size": page_size,
            "validation_page_size_options": PAGE_SIZE_OPTIONS,
            "validation_filtered_count": len(filtered_details),
            "validation_page_obj": page_obj,
            "validation_query_string": _validation_query_string(request),
        },
    )


def ifrs9_validations_download_view(request: HttpRequest):
    selected_branch = _clean(request.GET.get("branch"))
    branch_names = get_request_branch_names(request)
    if selected_branch and selected_branch.casefold() not in {name.casefold() for name in branch_names}:
        messages.error(request, "Choose a branch within your assigned branch scope before downloading.")
        return redirect("scorecard:ifrs9_results_ifrs9_validations")

    historical_queryset, _branch_names = _historical_scope_queryset(request, selected_branch)
    available_dates = _available_historical_dates(historical_queryset)
    current_date, previous_date = _resolve_comparison_dates(
        available_dates,
        request.GET.get("current_date"),
        request.GET.get("previous_date"),
    )
    if current_date is None or previous_date is None:
        messages.error(request, "At least two historical month-end dates with IFRS9 scores are required for IFRS9 scores validation.")
        return redirect("scorecard:ifrs9_results_ifrs9_validations")

    movement_filter = _clean(request.GET.get("movement"))
    if movement_filter not in VALID_MOVEMENT_FILTERS:
        movement_filter = ""
    search = _clean(request.GET.get("search"))
    branch_scope_label = current_branch_display_name(request) or "No branch selected"
    export_cache_key = validation_export_cache_key(
        report_name="ifrs9",
        queryset=historical_queryset,
        previous_date=previous_date,
        current_date=current_date,
        branch_names=branch_names,
        selected_branch=selected_branch,
        branch_scope_label=branch_scope_label,
        movement_filter=movement_filter,
        search=search,
    )
    cached_export = get_cached_validation_export(export_cache_key)
    cache_status = "HIT" if cached_export else "MISS"
    if cached_export:
        workbook_bytes, customer_row_count = cached_export
    else:
        payload = _get_validation_payload(
            historical_queryset,
            previous_date,
            current_date,
            branch_names,
            selected_branch,
        )
        detail_rows = _filter_details(payload["details"], search, movement_filter)
        workbook = _build_validation_workbook(
            payload,
            detail_rows,
            branch_scope_label,
            selected_branch,
            movement_filter,
            search,
        )
        output = BytesIO()
        workbook.save(output)
        workbook_bytes = output.getvalue()
        customer_row_count = len(detail_rows)
        cache_validation_export(export_cache_key, workbook_bytes, customer_row_count)

    log_ifrs9_results_audit(
        request.user,
        "download_ifrs9_validations",
        details=(
            f"Downloaded IFRS9 scores validation. Previous month-end: {previous_date.isoformat()}; "
            f"Current month-end: {current_date.isoformat()}; Branch scope: "
            f"{selected_branch or current_branch_display_name(request) or 'No branch selected'}; "
            f"Customer rows: {customer_row_count}"
        ),
        object_id=f"{previous_date.isoformat()}:{current_date.isoformat()}",
        branch_name=selected_branch or current_branch_display_name(request) or "",
    )
    response = HttpResponse(
        workbook_bytes,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="ifrs9_scores_validation_{previous_date.isoformat()}_to_{current_date.isoformat()}.xlsx"'
    )
    response["X-Validation-Export-Cache"] = cache_status
    return response
